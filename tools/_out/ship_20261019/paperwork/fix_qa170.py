"""Batch 16 paperwork fix (2026-10-19): QA_Log #170 'Result' said '19 measure scripts added'; the deploy added 17
(checks/shipped_files_sha256.txt). Change that one cell; prove nothing else changed (snapshot() of update_masterdata_16.py)."""
import hashlib, os, shutil, sys
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
src = open(os.path.join(HERE, 'update_masterdata_16.py'), encoding='utf-8').read()
exec(src.split('def main(apply)')[0].replace("exec(open(", "pass  # exec(open(", 1))
SHA0 = '262392f98fdd99b39b1b73b537dc8e84be5a4a591520654bfe42edf2e525dc8c'
assert sha(LIVE) == SHA0
m0, s0 = snapshot(openpyxl.load_workbook(LIVE))
wb = openpyxl.load_workbook(LIVE); ws = wb['QA_Log']
row = next(r for r in range(2, ws.max_row + 1) if str(ws.cell(r, 1).value) == '170')
c = ws.cell(row, 5); assert c.value.count('19 measure scripts added') == 1
c.value = c.value.replace('19 measure scripts added', '17 measure scripts added')
out = os.path.join(HERE, 'MasterData_fix170.xlsx'); wb.save(out)
m1, s1 = snapshot(openpyxl.load_workbook(out))
assert m0 == m1
for n in m0['sheets']:
    a, b = s0[n], s1[n]
    diff = [k for k in set(a['cells']) | set(b['cells']) if a['cells'].get(k) != b['cells'].get(k)]
    assert diff == ([c.coordinate] if n == 'QA_Log' else []), (n, diff)
    assert a['styles'] == b['styles'] and all(a[x] == b[x] for x in ('freeze', 'filter', 'widths', 'merged', 'comments', 'max_row', 'max_col')), n
shutil.copyfile(out, LIVE)
print('QA_Log', c.coordinate, 'fixed; only that cell changed; live sha256', sha(LIVE))
