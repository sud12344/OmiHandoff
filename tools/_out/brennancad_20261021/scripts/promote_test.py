"""promote_test.py (2026-10-21): decide MISSING -> ASSUMED by the generator itself, one row at a time.
apply_cells.py promotes a touched row only when every required column is filled (a cautious guess at what the generator reads). Here every
touched row that is still MISSING and whose root has a generator SPEC is tried alone: a scratch copy of its table with only that row set to
ASSUMED is generated (same tools as regen.py); the promotion is kept when the generator exits 0, the row appears in the SW2 table with no
'GEOMETRY CHECK FAILED', and no other configuration of the root changes. Kept promotions are written into the stage tables and logged in
promote.csv (row, result, reason). Run apply_cells.py first, then this, then regen.py and the checks."""
import os, sys, re, csv, shutil, subprocess, collections, tempfile
import openpyxl

H = os.path.dirname(os.path.abspath(__file__)); A = os.path.dirname(H)
ROOT = os.path.normpath(os.path.join(A, '..', '..', '..'))
J = lambda *a: os.path.join(ROOT, *a)
AF = J('tools', '_out', 'airway_first_20261020'); OR = J('tools', '_out', 'orfs_gen_20261020', 'tools')
GEN_AF = os.path.join(AF, 'tools_af', 'generators'); GEN_OR = os.path.join(OR, 'generators')
FOLD = {'2404L': '2404', '2501LL': '2501', '6801L': '6801'}
norm = lambda s: re.sub(r'\s+', '-', str(s).strip()).upper()


def spec(g):
    code = open(g, encoding='utf-8').read(); m = re.search(r'^SPEC\s*=\s*\{', code, re.M)
    k = set(re.findall(r"^\s*'([0-9A-Z]{3,8})'\s*:\s*dict\(", code[m.end():], re.M)) if m else set()
    return k | set(re.findall(r"^SPEC\['([0-9A-Z]{3,8})'\]\s*=", code, re.M))


gs, gt = spec(os.path.join(GEN_AF, 'gen_straight2.py')), spec(os.path.join(GEN_AF, 'gen_tee.py'))
ogs, ogt = spec(os.path.join(GEN_OR, 'gen_straight2.py')), spec(os.path.join(GEN_OR, 'gen_tee.py'))


def gen_for(root):
    m = FOLD.get(root, root)
    if root.startswith('FS'):
        return m, (GEN_OR, 'gen_straight2.py', 'straight') if m in ogs else (GEN_OR, 'gen_tee.py', 'tee') if m in ogt else None
    return m, (GEN_AF, 'gen_straight2.py', 'straight') if m in gs else (GEN_AF, 'gen_tee.py', 'tee') if m in gt else None


def table_path(stage_dir, root):
    return os.path.join(stage_dir, 'orfs_tables' if root.startswith('FS') else '04_DESIGN_TABLES', f'DesignTable_{root}.xlsx')


def set_status(path, part, value):
    wb = openpyxl.load_workbook(path); ws = wb.worksheets[0]
    hr = next(i for i in range(1, 7) if any(c.value == '$PRP@PartNo' for c in ws[i]))
    head = {c.value: c.column for c in ws[hr] if c.value}
    for i in range(hr + 1, ws.max_row + 1):
        pn = ws.cell(row=i, column=head['$PRP@PartNo']).value
        sku = ws.cell(row=i, column=head['$PRP@OMiSKU']).value if '$PRP@OMiSKU' in head else None
        if pn and norm(sku or pn) == part:
            ws.cell(row=i, column=head['$PRP@DataStatus'], value=value); wb.save(path); return True
    return False


def generate(stage_dir, root, out):
    m, g = gen_for(root)
    gen, script, sub = g
    if root.startswith('FS'):
        env = dict(os.environ, OMI_ROOT=ROOT, OMI_REV_EXTRA=os.path.join(stage_dir, 'orfs_tables'), GS_OUT=out, GT_OUT=out)
    else:
        env = dict(os.environ, OMI_ROOT=stage_dir, GS_OUT=out, GT_OUT=out)
    p = subprocess.run([sys.executable, '-I', '-B', os.path.join(gen, script), m], env=env, capture_output=True, text=True)
    sw2 = os.path.join(out, f'DesignTable_{m}_SW2.xlsx')
    rows = {}
    if os.path.exists(sw2):
        rs = list(openpyxl.load_workbook(sw2, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)); h = rs[0]
        rows = {r[0]: dict(zip(h, r)) for r in rs[1:] if r and r[0]}
    return p.returncode, rows, p.stdout[-400:] + p.stderr[-400:]


STAGE = os.path.join(A, 'stage')
log = []
cand = [r for r in csv.DictReader(open(os.path.join(A, 'rows.csv'), encoding='utf-8')) if r['status_after'] == 'MISSING']
base_cache = {}
for r in cand:
    m, g = gen_for(r['root'])
    if g is None:
        log.append(dict(root=r['root'], part=r['part'], result='no SPEC', reason='')); continue
    tmp = tempfile.mkdtemp(prefix='promote_')
    shutil.copytree(os.path.join(STAGE, '04_DESIGN_TABLES'), os.path.join(tmp, '04_DESIGN_TABLES'))
    shutil.copytree(os.path.join(STAGE, 'orfs_tables'), os.path.join(tmp, 'orfs_tables'))
    if r['root'] not in base_cache:
        base_cache[r['root']] = generate(tmp, r['root'], os.path.join(tmp, 'base'))[1]
    base = base_cache[r['root']]
    set_status(table_path(tmp, r['root']), r['part'], 'ASSUMED')
    rc, rows, tail = generate(tmp, r['root'], os.path.join(tmp, 'try'))
    mine = [k for k, d in rows.items() if norm(d.get('$PRP@OMiSKU') or d.get('$PRP@PartNo') or k) == r['part']]
    others_changed = [k for k in set(base) | set(rows) if k not in mine and base.get(k) != rows.get(k)]
    failed = [k for k in mine if 'GEOMETRY CHECK FAILED' in str(rows[k].get('$COMMENT') or '')]
    if rc == 0 and mine and not failed and not others_changed:
        set_status(table_path(STAGE, r['root']), r['part'], 'ASSUMED')
        base_cache.pop(r['root'], None)      # the next row of this root is compared with the stage that now includes this promotion
        log.append(dict(root=r['root'], part=r['part'], result='PROMOTED', reason=f'generates {mine}'))
    else:
        why = (f'rc {rc}; ' if rc else '') + ('not generated (blocked/skipped); ' if not mine else '') + \
              (f'geometry check failed; ' if failed else '') + (f'other configs change {others_changed[:3]}; ' if others_changed else '')
        log.append(dict(root=r['root'], part=r['part'], result='stays MISSING', reason=why + tail.replace('\n', ' ')[-300:]))
    shutil.rmtree(tmp, ignore_errors=True)
    print(log[-1]['root'], log[-1]['part'], log[-1]['result'], flush=True)
with open(os.path.join(A, 'promote.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=list(log[0])); w.writeheader(); w.writerows(log)
print(collections.Counter(x['result'] for x in log))
