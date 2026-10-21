"""triveni_package.py (2026-10-21, cloud): the parts no source prints, as a package for Triveni (OMi).
Inputs: the PC's BLOCKED_NEEDS.md (2026-10-20) category (b) "Dimensions missing everywhere - ask OMi", checked against this folder's
sources.csv (source_lookup.py), and OMi's own parts file for the descriptions.
Changes to (b), from the source check:
- 6505-04-02 is left out: OMi's parts file prints Overall Length, L 1, Hex, Side 1/2 ID and Nut Hex 1 for it (a source; the review
  table's "no source" predates that check).
- the 2700-LN rows are asked as "same as the plain part?" (A10: -LN parts are modelled as the plain part, no locknut).
- FS0319-24 is a question (three printed values disagree), not a size request.
Writes 10_DRAWINGS/for_OMi_2026-10-21/: the PDF, the fill-in workbook and the email text. Nothing is sent."""
import os, re, csv, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, '..', '..', '..'))
J = lambda *a: os.path.join(ROOT, *a)
OUT = J('10_DRAWINGS', 'for_OMi_2026-10-21'); os.makedirs(OUT, exist_ok=True)
DATE = '2026-10-21'

# --- (b) groups from the PC list
B = open(J('tools', '_out', 'blocked_needs_20261020', 'BLOCKED_NEEDS.md'), encoding='utf-8').read().split('## Appendix')[1]
grp = {}
for m in re.finditer(r'^\*\*([a-e]) / (.+?)\*\* \((\d+)\): (.+)$', B, re.M):
    if m.group(1) == 'b':
        ps = [p.strip().upper() for p in m.group(4).split(',')]; assert len(ps) == int(m.group(3)); grp[m.group(2)] = ps
G_COMP = next(k for k in grp if k.startswith('Compression'))
G_TG = next(k for k in grp if k.startswith('Test gauge'))
G_SINGLE = next(k for k in grp if k.startswith('Parts with no Air-Way row'))
G_ORFS = next(k for k in grp if k.startswith('ORFS parts with no usable'))
G_4005 = next(k for k in grp if k.startswith('4005'))
G_FL = next(k for k in grp if k.startswith('Flareless'))
assert sum(map(len, grp.values())) == 291

src = {r['part']: r for r in csv.DictReader(open(os.path.join(HERE, 'sources.csv'), encoding='utf-8'))}
for g in grp.values():
    for p in g:
        assert p in src, p
LEFT_OUT = {p: f"OMi's parts file prints {src[p]['omi_dims']}" for g in grp.values() for p in g if src[p]['omi_file'].startswith('=')}
assert set(LEFT_OUT) == {'6505-04-02'}, LEFT_OUT   # the only (b) part a source covers - re-check by hand if this changes

ws = openpyxl.load_workbook(J('07_RAW_EXTRACTS', 'omi_dimensions', 'WebsitePartsInfo2026.xlsx'), read_only=True, data_only=True)['Items']
rows = list(ws.iter_rows(values_only=True)); H = rows[0]
desc = {}
for r in rows[3:]:
    d = dict(zip(H, r))
    if d.get('Part Number'):
        desc[str(d['Part Number']).strip().upper()] = str(d.get('Part Description') or '').strip()


def needs(p, ds):
    """What we ask for, in the column names of OMi's own parts file."""
    swivel = bool(re.search(r'\b(FJS|FFSS|FPS)\b', ds))
    if p.startswith('2700-LN-'):
        return 'Same as the plain part ' + p.replace('-LN', '') + '? (we model the -LN part without the lock nut)'
    if p.startswith('2700-'):
        return 'Overall Length, Hex Flat, Side 1 ID, Side 2 ID, bulkhead thread length'
    if p.startswith('5404-R-'):
        return 'Overall Length, Hex Flat, orifice length'
    if p.startswith('5404-L-'):
        return 'Hex Flat, hex length and position, Side 1 ID'
    if p.startswith('FS0304-C'):
        return 'Overall Length, Hex Flat, insert sizes (a drawing is best)'
    if '1803-' in p:
        return 'M 1 Size, L 1 Size, Side 1 ID, Side 2 ID, flange head diameter and thickness'
    if 'Tee' in ds:
        return 'M 1 Size, L 1 Size, branch leg length, Hex Flat, Side 1-3 ID' + (', Nut Hex 1' if swivel else '')
    if 'Elbow' in ds:
        return 'M 1 Size, L 1 Size, Hex Flat, Side 1 ID, Side 2 ID' + (', Nut Hex 1' if swivel else '')
    return 'Overall Length, Hex Flat, Side 1 ID, Side 2 ID' + (', Nut Hex 1' if swivel else '')


singles = [p for p in grp[G_SINGLE] + grp[G_ORFS] if p not in LEFT_OUT and p != 'FS0319-24']
single_rows = [(p, desc.get(p, ''), needs(p, desc.get(p, ''))) for p in singles]
tg_rows = [(p, desc.get(p, ''), 'M 1 Size, L 1 Size, test-port leg length, Hex Flat, Nut Hex 1, Side 1-3 ID') for p in grp[G_TG]]

# compression: one body per series; "-M1" = the same body sold without nut and sleeve
comp = collections.OrderedDict()
for p in sorted(grp[G_COMP]):
    s = p.split('-')[0]; comp.setdefault(s, {'parts': [], 'm1': []})
    comp[s]['m1' if p.endswith('-M1') else 'parts'].append(p)
SER_DESC = {}
for s, v in comp.items():
    d = desc.get((v['parts'] or v['m1'])[0], '')
    SER_DESC[s] = re.sub(r'^\d\d\s+', '', re.sub(r'\b\d\d\s+(?=(BT|MJ|MP|FP|MAORB)\b)', '', d)).replace(' (Body)', '')
n_out = 230 + len(tg_rows) + len(single_rows) + 1 + len(grp[G_4005]) + len(grp[G_FL])
assert n_out == 291 - len(LEFT_OUT), n_out


def dashes(ps):
    return ', '.join(re.sub(r'^[A-Z]*\d{4}-', '', p).replace('-NWO', '').replace('-FG', '').replace('-M1', '') for p in ps)


# ---------------------------------------------------------------- PDF
ss = getSampleStyleSheet()
INK, MUTED, RULE, HEAD = colors.HexColor('#1f2933'), colors.HexColor('#52606d'), colors.HexColor('#cbd2d9'), colors.HexColor('#e4e7eb')
st = {
    'title': ParagraphStyle('t', parent=ss['Title'], fontName='Helvetica-Bold', fontSize=20, leading=24, textColor=INK, alignment=0, spaceAfter=4),
    'sub': ParagraphStyle('s', parent=ss['Normal'], fontSize=10, leading=13, textColor=MUTED, spaceAfter=14),
    'h1': ParagraphStyle('h1', parent=ss['Heading2'], fontName='Helvetica-Bold', fontSize=13.5, leading=17, textColor=INK, spaceBefore=14, spaceAfter=6),
    'h2': ParagraphStyle('h2', parent=ss['Heading3'], fontName='Helvetica-Bold', fontSize=10.5, leading=13, textColor=INK, spaceBefore=8, spaceAfter=3),
    'p': ParagraphStyle('p', parent=ss['Normal'], fontSize=9.5, leading=13, textColor=INK, spaceAfter=6),
    'cell': ParagraphStyle('c', parent=ss['Normal'], fontSize=8.3, leading=10.4, textColor=INK),
    'cellb': ParagraphStyle('cb', parent=ss['Normal'], fontName='Helvetica-Bold', fontSize=8.3, leading=10.4, textColor=INK),
    'li': ParagraphStyle('li', parent=ss['Normal'], fontSize=9.5, leading=13, textColor=INK, leftIndent=12, bulletIndent=2, spaceAfter=3),
}
C = lambda t, b=False: Paragraph(str(t), st['cellb' if b else 'cell'])


def table(head, body, widths, zebra=True):
    t = Table([[C(h, True) for h in head]] + [[C(x) for x in r] for r in body], colWidths=widths, repeatRows=1)
    sty = [('BACKGROUND', (0, 0), (-1, 0), HEAD), ('LINEBELOW', (0, 0), (-1, 0), 0.8, MUTED), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
           ('LINEBELOW', (0, 1), (-1, -1), 0.25, RULE), ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
           ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]
    if zebra:
        sty += [('BACKGROUND', (0, i), (-1, i), colors.HexColor('#f5f7fa')) for i in range(2, len(body) + 1, 2)]
    t.setStyle(TableStyle(sty)); return t


def footer(c, doc):
    c.saveState(); c.setFont('Helvetica', 7.5); c.setFillColor(MUTED)
    c.drawString(0.75 * inch, 0.5 * inch, f'OM Industrial carbon-steel fitting library - parts without dimensions - {DATE}')
    c.drawRightString(letter[0] - 0.75 * inch, 0.5 * inch, f'Page {doc.page}'); c.restoreState()


W = letter[0] - 1.5 * inch
S = [Paragraph('Parts we have no dimensions for', st['title']),
     Paragraph(f'OM Industrial carbon-steel fitting library &nbsp;|&nbsp; {DATE} &nbsp;|&nbsp; {n_out} part numbers', st['sub']),
     Paragraph('We checked every part number on your website that is not modelled yet against Air-Way\'s catalog, Brennan\'s catalogs and '
               '3D files, and your website parts file (WebsitePartsInfo2026.xlsx). The parts below appear in none of them with '
               'dimensions, so we cannot model them without your help. All sizes in inches.', st['p']),
     Paragraph('<b>The quickest answer is usually your supplier\'s drawing:</b> one drawing per series with its size table is plenty, and '
               'then nothing needs to be filled in. Otherwise the attached workbook has a yellow cell for each value we need, using the '
               'same column names as your parts file.', st['p']),
     Spacer(1, 4),
     table(['', 'What', 'Parts', 'What we need'], [
         ['1', 'Compression (bite-type) fittings, nuts and sleeves', '230', f'A drawing per series ({len(comp)} series); which standard the bite-type port follows'],
         ['2', 'Test-gauge tees TGT-6504 / TGT-FS2406', str(len(tg_rows)), 'A drawing per series (2 series)'],
         ['3', 'Single sizes nobody prints', str(len(single_rows)), 'The values listed per part (or a drawing)'],
         ['4', 'Flange O-rings 4005', str(len(grp[G_4005])), 'Code 61 or 62; AS568 size or I.D. x cross-section; material, hardness'],
         ['5', 'Flareless nut 0618 / sleeve 0619', str(len(grp[G_FL])), 'Are they Brennan\'s V0318 / V0319? If not: sizes and shape'],
         ['6', 'FS0319-24 sleeve: three printed values disagree', '1', 'Which outside diameter is right'],
     ], [0.25 * inch, 2.6 * inch, 0.5 * inch, W - 3.35 * inch]),
     Spacer(1, 8),
     Paragraph('Not on this list: parts that Air-Way, Brennan or your parts file already print, and parts we are waiting to download '
               'from Brennan\'s 3D files. We are treating the kits 1901 / 1902 / 1927 (30 parts) as out of scope; tell us if you need them '
               'modelled.', st['p'])]

S += [PageBreak(), Paragraph('1. Compression (bite-type) fittings - 230 parts', st['h1']),
      Paragraph('No Air-Way series, and Brennan prints only its own V0318 / V0319 nut and sleeve. A "(Body)" part number (-M1) is, we assume, the '
                'same body sold without the nut and sleeve, so each series needs one set of sizes. Please also tell us which standard the '
                'bite-type port follows (thread, tube bore, the cone the sleeve bites into, port depth).', st['p']),
      table(['Series', 'Description (your file)', 'Sizes (dash numbers)', 'Parts'],
            [[s, SER_DESC[s], dashes(v['parts']) + (f'<br/><font color="#52606d">Body only (-M1): {dashes(v["m1"])}</font>' if v['m1'] else ''),
              str(len(v['parts']) + len(v['m1']))] for s, v in comp.items()],
            [0.6 * inch, 1.75 * inch, W - 2.85 * inch, 0.5 * inch]),
      Paragraph('What we need per series: Overall Length (or M 1 / L 1 Size for elbows and tees), Hex Flat, Side IDs, Nut Hex, and for the '
                'sleeve: length, outside diameter and the shape of its front edge.', st['p'])]

S += [Paragraph('2. Test-gauge tees - 12 parts', st['h1']),
      Paragraph('Swivel unions with a -04 pipe test port. Air-Way\'s test points are a different (metric) part, and Brennan does not print these.', st['p']),
      table(['Part number', 'Description (your file)', 'What we need'], tg_rows, [1.45 * inch, 2.35 * inch, W - 3.8 * inch])]

S += [PageBreak(), Paragraph(f'3. Single sizes nobody prints - {len(single_rows)} parts', st['h1']),
      Paragraph('The rest of each series is modelled; only these sizes are missing from Air-Way, Brennan and your parts file. The values we '
                'need use your parts file\'s column names (M 1 / L 1 Size = the two leg lengths of an elbow).', st['p']),
      table(['Part number', 'Description (your file)', 'What we need'], single_rows, [1.45 * inch, 2.35 * inch, W - 3.8 * inch])]

S += [Paragraph('4. Flange O-rings 4005 - 9 parts', st['h1']),
      Paragraph(f'{", ".join(grp[G_4005])}. Are they for code 61 or code 62 flanges? Which AS568 size is each one (or its inside diameter and '
                'cross-section)? Material and hardness. (Also on the O-ring tab of the earlier size request, attached.)', st['p']),
      Paragraph('5. Flareless nut 0618 and sleeve 0619 - 4 parts', st['h1']),
      Paragraph(f'{", ".join(grp[G_FL])}. Brennan sells a flareless nut V0318 and sleeve V0319 in the same sizes. Are yours the same parts? '
                'If yes, we use Brennan\'s 3D files; if not, please send length, hex, inside and outside diameters and the shape (a drawing or '
                'a photo with sizes).', st['p']),
      Paragraph('6. FS0319-24 face-seal sleeve - which outside diameter?', st['h1']),
      Paragraph('Air-Way prints 1.71 for the sleeve, but its own -24 female face-seal end (FF2407-24) is 2.00; Brennan prints 1.910. Which '
                'is right for your part?', st['p'])]


def build_pdf(path):
    SimpleDocTemplate(path, pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch, topMargin=0.7 * inch, bottomMargin=0.8 * inch,
                      title='Parts we have no dimensions for', author='OMi fitting library').build(S, onFirstPage=footer, onLaterPages=footer)


PDF = os.path.join(OUT, f'OMi_parts_without_dimensions_{DATE}.pdf'); build_pdf(PDF)

# ---------------------------------------------------------------- workbook (fill-in)
wb = openpyxl.Workbook(); YEL = PatternFill('solid', fgColor='FFF2A8'); GRY = PatternFill('solid', fgColor='E4E7EB')
thin = Side(style='thin', color='CBD2D9'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
COLS = ['Overall Length', 'M 1 Size', 'L 1 Size', 'Hex Flat', 'Side 1 ID', 'Side 2 ID', 'Side 3 ID', 'Nut Hex 1', 'Other (see "What we need")']


def sheet(ws, title, note, items):
    ws.title = title; ws['A1'] = note; ws['A1'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells('A1:M1')
    ws.row_dimensions[1].height = 48
    head = ['#', 'Part number', 'Description (your file)', 'What we need'] + COLS + ['Your notes']
    for j, h in enumerate(head, 1):
        c = ws.cell(row=2, column=j, value=h); c.font = Font(bold=True); c.fill = GRY; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, (p, d, n) in enumerate(items, 1):
        vals = [i, p, d, n]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i + 2, column=j, value=v); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top')
        for j, col in enumerate(COLS, 5):
            c = ws.cell(row=i + 2, column=j); c.border = BOX
            key = col.split(' (')[0]
            wanted = key in n or (key == 'Other' and re.search(r'length|orifice|insert|flange|Same as|thread|shape|branch|test-port', n))
            c.fill = YEL if wanted else GRY
        ws.cell(row=i + 2, column=len(head)).border = BOX
    for j, w in enumerate([4, 22, 34, 42] + [11] * len(COLS) + [28], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = 'C3'


comp_items = []
for s, v in comp.items():
    for p in v['parts'] + v['m1']:
        n = ('Body only: same as ' + p[:-3] + '? ' if p.endswith('-M1') else '') + (
            'Overall Length, Hex Flat, Side 1 ID, shape of the inside (a drawing)' if s == 'C0318' else
            'Overall Length, outside diameter, Side 1 ID, shape of the front edge (a drawing)' if s == 'C0319' else
            'M 1 Size, L 1 Size, Hex Flat, Side 1 ID, Side 2 ID' if re.match(r'C(25|26|68)', s) else 'Overall Length, Hex Flat, Side 1 ID, Side 2 ID')
        comp_items.append((p, desc.get(p, ''), n))
sheet(wb.active, 'Compression', 'Compression (bite-type) fittings. One supplier drawing per series with its size table is enough - then '
      'nothing needs to be filled in. Yellow = we need it, grey = not needed. Inches. Also: which standard does the bite-type port follow?', comp_items)
sheet(wb.create_sheet(), 'Test gauge', 'Test-gauge tees (-04 pipe test port). A drawing per series is enough. Yellow = we need it. Inches.', tg_rows)
sheet(wb.create_sheet(), 'Single sizes', 'Single sizes no catalog prints (the rest of each series is modelled). Yellow = we need it. Inches.',
      single_rows)
XLSX = os.path.join(OUT, f'OMi_parts_without_dimensions_{DATE}.xlsx'); wb.save(XLSX)

# ---------------------------------------------------------------- email
EMAIL = f"""DRAFT - not sent.

To: Triveni (OM Industrial)
Subject: Parts we have no dimensions for ({n_out} part numbers)

Hi Triveni,

We have now checked every part number on your website that is not modelled yet against Air-Way's catalog, Brennan's catalogs and 3D
files, and your website parts file. {n_out} of them appear in none of these with dimensions, so we cannot model them without your help.
The attached PDF lists them; the attached workbook has a yellow cell for each value we need, using your parts file's column names.

In short:
1. Compression (bite-type) fittings, nuts and sleeves - 230 part numbers in {len(comp)} series. One supplier drawing per series with its size
   table would cover them. Which standard does the bite-type port follow? And is each "(Body)" -M1 part the same body without the nut
   and sleeve?
2. Test-gauge tees TGT-6504 / TGT-FS2406 - 12 parts, two series.
3. {len(single_rows)} single sizes that no catalog prints (the rest of each series is modelled), e.g. 1405-20-24, 2403-06-02, the
   2700 bulkhead reducers and the 6500 / 6502 forged reducers. For the four 2700-LN parts: are they the same as the plain 2700 part?
4. 4005 flange O-rings (9): code 61 or 62, AS568 size (or I.D. x cross-section), material and hardness.
5. Flareless nut 0618 and sleeve 0619 (4): are they the same parts as Brennan's V0318 / V0319?
6. FS0319-24 sleeve: Air-Way prints 1.71, Brennan 1.910, Air-Way's own -24 female face-seal end 2.00 - which is right?

Drawings are the easiest for you and for us; if you send drawings, the cells do not need filling in. Also, we are treating the kits
1901 / 1902 / 1927 as out of scope - please tell us if you need them modelled.

Best regards,
Sudhanva

Attachments:
- OMi_parts_without_dimensions_{DATE}.pdf
- OMi_parts_without_dimensions_{DATE}.xlsx
"""
open(os.path.join(OUT, f'Email_to_Triveni_{DATE}.txt'), 'w', encoding='utf-8').write(EMAIL)
print('parts on the list:', n_out, '| left out:', LEFT_OUT, '| singles', len(single_rows), '| TG', len(tg_rows), '| compression series', len(comp))
print(PDF); print(XLSX)
