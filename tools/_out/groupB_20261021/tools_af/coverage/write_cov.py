import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
import omi_paths as P    # portable paths (2026-09-25)
_WORK = P.COVERAGE_WORK; _os.makedirs(_WORK, exist_ok=True); _HERE = _os.path.dirname(_os.path.abspath(__file__))
import json, os, collections, re, datetime, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
R = P.ROOT
exec(open(os.path.join(_HERE, 'coverage.py')).read().split("rows = []")[0].split("res = json.load")[0])   # imports only
src = open(os.path.join(_HERE, 'coverage.py')).read()
Gtxt = src[src.index('G = collections.OrderedDict(['):src.index('])\n', src.index('G = collections.OrderedDict(['))+2]
G = eval(Gtxt.split('=', 1)[1])
rows = json.load(open(os.path.join(_WORK, 'cov_rows.json')))
cnt = collections.Counter(r[4] for r in rows)
wb = openpyxl.Workbook()
B = Font(bold=True); H = PatternFill('solid', fgColor='DDEBF7'); W = Alignment(wrap_text=True, vertical='top')
def sheet(name, header, data, widths):
    ws = wb.create_sheet(name) if wb.worksheets[0].title != 'Sheet' else wb.worksheets[0]
    ws.title = name; ws.append(header)
    for c in ws[1]: c.font = B; c.fill = H
    for d in data: ws.append(list(d))
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = W
    ws.freeze_panes = 'A2'
    return ws
# ---- Summary
summ = []
for k, (lab, unl) in G.items(): summ.append((k, lab, cnt.get(k, 0), unl))
summ.append(('N1', 'Not captured yet - Tube Fittings & Adapters SKUs the OMi feed dropped (1,672 published - 1,595 captured)', 1672 - sum(1 for r in rows if r[1] == 'cs-adapters'), 'me: re-crawl the OMi feed'))
summ.append(('N2', 'Not started - O-Ring Face Seal (ORFS) collection', 385, 'SKU list from the OMi feed + ORFS dimensions (Air-Way ORFS pages / SAE J1453) + a face-seal end module'))
summ.append(('N3', 'Not started - Compression collection', 230, 'SKU list from the OMi feed + compression dimensions (SAE J514/2 or maker catalog) + nut/sleeve module'))
summ.append(('N4', 'Not started - Test Gauge collection', 12, 'SKU list + dimensions'))
tot = sum(x[2] for x in summ)
summ.append(('', 'TOTAL (OMi carbon-steel catalog)', tot, ''))
ws = sheet('Summary', ['Code', 'Status', 'OMi parts', 'What unlocks it'], summ, [6, 70, 10, 90])
ws.insert_rows(1, 2); ws['A1'] = f'OMi carbon-steel library - coverage by OMi part number, enumerated {datetime.date.today().isoformat()} from the OMi feed extract, every design table and 11_MASTERS'; ws['A1'].font = Font(bold=True, size=12)
# ---- STEP shopping list
steps = collections.OrderedDict()
for r in rows:
    g, note = r[4], r[6]
    if g in ('D1', 'D2', 'D3', 'D4', 'D5'):
        m = re.search(r'Brennan (\S+)', note)
        key = (g, note.split(':')[0].strip(), m.group(1) if m else note)
        steps.setdefault(key, []).append(r[0])
    elif g == 'D7' and 'Brennan' in note:
        key = (g, note.split(':')[0].strip(), re.search(r'Brennan (\S+)', note).group(1)); steps.setdefault(key, []).append(r[0])
    elif g == 'D6':
        steps.setdefault((g, 'adjustable ORB stud + locknut (' + r[3] + ')', '6801-XX-XX-NWO-FG per ORB size'), []).append(r[0])
sl = [(G[k[0]][0].replace('Needs a Brennan STEP - ', '').replace('Needs Brennan STEPs - ', ''), k[1], k[2], len(v), ' '.join(v)) for k, v in steps.items()]
sl.sort(key=lambda x: (-x[3]))
sheet('STEP shopping list', ['Group', 'What is missing', 'Brennan STEP that unlocks it', 'OMi parts unlocked', 'OMi parts'], sl, [40, 45, 34, 10, 120])
# ---- Decisions
dec = collections.OrderedDict()
for r in rows:
    if r[4] == 'E': dec.setdefault(r[6], []).append(r[0])
f1 = [r[0] for r in rows if r[4] == 'F1']
dl = [(q, len(v), ' '.join(v)) for q, v in dec.items()]
dl.append(('OMi sells forged (-FG) parts where Air-Way lists only the non-forged size: use the non-forged dimensions (ASSUMED, flagged)? '
           'Evidence from the 399 Air-Way sizes listed both ways: only 30 identical; the forged hex is usually 1/16" smaller (354 of 598 hex values differ, mostly -0.06/-0.07), '
           'lengths mostly within 0.03 (271 of 2154 differ, a few by 0.1+), bores almost never differ.', len(f1), ' '.join(f1)))
sheet('Decisions', ['Question', 'OMi parts', 'OMi parts'], dl, [100, 10, 120])
# ---- By series
bs = collections.defaultdict(collections.Counter)
for r in rows: bs[r[3] or r[2] or '?'][r[4]] += 1
keys = list(G.keys())
sheet('By series', ['Series / table'] + keys + ['Total'], [[k] + [bs[k].get(x, 0) or None for x in keys] + [sum(bs[k].values())] for k in sorted(bs)], [14] + [5] * len(keys) + [7])
# ---- SKUs
sheet('OMi parts', ['OMi SKU', 'Collection', 'Root', 'Design table / master', 'Code', 'Status', 'Note'], rows, [22, 34, 10, 14, 6, 60, 90])
p = f'{R}/00_DOCS/OMi_Coverage_2026-09-24.xlsx'
wb.save(p)
print('saved', p, 'rows', len(rows), 'total', tot)
for x in summ: print(x[0], x[2], x[1][:80])
print('--- top STEPs'); [print(' ', x[3], x[2], '|', x[1]) for x in sl[:40]]
