"""Q2 (user 2026-10-20: option C): the doubtful Air-Way values (checks/look_at.csv) as a list for Triveni -
10_DRAWINGS/for_OMi_2026-10-20/AirWay_values_to_confirm_2026-10-20.xlsx. Rows settled by the review PDFs get removed before it is sent
(column 'PDF check'). Air-Way letters as printed on the Air-Way pages (Triveni works in AW / BI)."""
import csv, os, re, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
H = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..'))
A = os.path.join(H, 'tools', '_out', 'airway_first_20261020')
OUT = os.path.join(H, '10_DRAWINGS', 'for_OMi_2026-10-20')
def what(col):
    b = col.split('@')[0]
    if re.fullmatch(r'L(13|15|17|24|25|26|29|31|32|33)', b): return 'overall length'
    if b in ('L16', 'L18', 'L14'): return 'length to the swivel nut'
    if b.startswith('D'): return 'bore (drill)'
    if b in ('N', 'N1', 'M', 'M1', 'M2', 'N2'): return 'leg length'
    if 'HEX' in b: return 'hex'
    if b == 'W': return 'hex (W)'
    return b
def why(r):
    f = r['flagged']; b = r['col'].split('@')[0]
    if r['root'] == '6405' and b == 'L31' and abs(float(r['airway']) - 1.0) < 1e-9:
        return 'Air-Way prints exactly 1.00 - looks like a filler value'
    if r['root'] == '6405' and b == 'L31':
        return f"Air-Way's 6405 lengths were flagged as suspect in our first build (2026-09); differs from Brennan / OMi website by {abs(float(r['delta'])):.3f} in"
    if 'tap drill' in f and b.startswith('D'): return "Air-Way's bore equals the pipe tap drill - looks like the tap-drill column printed in the bore column"
    if 'unusual' in f or 'non-FG' in f: return "Air-Way's forged (-FG) length is much longer than its plain part of the same size"
    if 'likely typo' in f and b in f: return 'Air-Way value out of step with the neighbouring sizes (likely typo)'
    return f"differs from Brennan / OMi website by {abs(float(r['delta'])):.3f} in"
def keep(r):
    f = r['flagged']; b = r['col'].split('@')[0]; d = abs(float(r['delta']))
    return (d > 0.15 or ('tap drill' in f and b.startswith('D')) or 'non-FG' in f or ('likely typo' in f and b in f)
            or (r['root'] == '6405' and b == 'L31' and abs(float(r['airway']) - 1.0) < 1e-9))   # Air-Way 1.00 on 6405 looks like a filler
rows = [r for r in csv.DictReader(open(os.path.join(A, 'checks', 'look_at.csv'), encoding='utf-8')) if keep(r)]
rows.sort(key=lambda r: (r['root'], r['col'].split('@')[0], r['cfg']))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Air-Way values to confirm'
hdr = ['#', 'Part number', 'Air-Way letter', 'Dimension', 'Air-Way (AW)', 'Brennan / OMi website (BI)', 'Why we ask', 'Use AW or BI?', 'PDF check (internal - delete before sending)']
ws.append(hdr)
for i, r in enumerate(rows, 1):
    bi = r['now_omi']
    if r['cfg'] == '6405-24-16-O' and r['col'].startswith('D3'): bi = '1.080 (OMi website) / 1.312 (Brennan catalog)'
    ws.append([i, r['cfg'], r['col'].split('@')[0], what(r['col']), float(r['airway']), bi, why(r), '', ''])
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDDDDD'); c.alignment = Alignment(wrap_text=True)
for col, w in zip('ABCDEFGHI', (4, 16, 9, 22, 11, 22, 60, 12, 30)): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
os.makedirs(OUT, exist_ok=True)
wb.save(os.path.join(OUT, 'AirWay_values_to_confirm_2026-10-20.xlsx'))
print(len(rows), 'rows')
