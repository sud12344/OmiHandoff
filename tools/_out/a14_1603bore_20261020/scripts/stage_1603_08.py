"""stage_1603_08.py <stage_dir>: A14 (user 2026-10-18 "Measure the STEP") for 1603-08-08-08-FG, the one 1603 -FG row the 2026-10-18 stage left
empty because Brennan's TFA catalog prints its two drills as a range (A1 / A2 DRILL "0.44/0.46", never averaged).
- D1@BORE / D2@BORE = 0.450, measured on 09_STEP_FILES/BRENNAN_1603-08-08-08-FG.stp (cyl_survey.txt: run bore cylinder dia 0.4500 from stem
  tip to stem tip, x = -1.11 .. +1.11; branch bore cylinder dia 0.4500 from the crossing to the branch stem tip, y = 0 .. 1.11);
- the other 15 cells = the row's prepared Brennan-catalog / end-library cells of wave2/tee_elbow/cells.csv (the same source the -02 / -04 /
  -06 / -12 / -16 siblings were staged from by apply_genfeat.py), cross-checked on the same STEP (plane_survey.txt);
- marking as apply_genfeat.py: $PRP@DimSource BRENNAN, AxialSource prefix [BRENNAN-ONLY 2026-10-18], MISSING -> ASSUMED (Brennan alone =
  one source: catalog and STEP are the same supplier).
Only EMPTY cells are filled (asserted). Appends to <stage>/CHANGES.csv."""
import os, sys, csv, re, shutil
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__)); TASK = os.path.dirname(HERE)
BY = os.path.normpath(os.path.join(TASK, '..', 'brennan_yes_20261018'))
STAGE = sys.argv[1]
T, PN = '1603', '1603-08-08-08-FG'
TAG = '[BRENNAN-ONLY 2026-10-18] '
P = os.path.join(STAGE, '04_DESIGN_TABLES', f'DesignTable_{T}.xlsx')
STEP_SRC = ('Brennan STEP 09_STEP_FILES/BRENNAN_1603-08-08-08-FG.stp measured 2026-10-20 (A14, user 2026-10-18 "Measure the STEP"; '
            'tools/_out/a14_1603bore_20261020/cyl_survey.txt): {what} = cylinder dia 0.4500; the TFA catalog p.48 prints "0.44/0.46" (a range, '
            'never averaged) - the STEP value lies inside it [Brennan = yes 2026-10-18]')
MEAS = [dict(table=T, config=PN, column='D1@BORE', value='0.45', kind='brennan-dim',
             source=STEP_SRC.format(what='branch (A1) bore, crossing to the branch stem tip y = 0 .. 1.11')),
        dict(table=T, config=PN, column='D2@BORE', value='0.45', kind='brennan-dim',
             source=STEP_SRC.format(what='run (A2 / A3) bore, stem tip to stem tip x = -1.11 .. +1.11'))]
NOTE = ('2026-10-20 A14: D1 / D2 bores 0.450 measured on the Brennan STEP (catalog prints the range 0.44/0.46); the other sizes = the Brennan '
        'TFA catalog p.48, all confirmed on the same STEP (nut faces 1.39, stem tips 1.11, nut hex 1.000, body flats 0.875) - staged, not shipped')

bak = os.path.join(TASK, 'old_files', f'DesignTable_{T}.xlsx')
if not os.path.exists(bak):
    shutil.copy2(P, bak)

cells = [c for c in csv.DictReader(open(os.path.join(BY, 'wave2', 'tee_elbow', 'cells.csv'), encoding='utf-8-sig'))
         if c['table'] == T and c['config'] == PN]
assert len(cells) == 15, len(cells)
assert not any(c['column'] in ('D1@BORE', 'D2@BORE') for c in cells)
cells += MEAS

wb = openpyxl.load_workbook(P); ws = wb.worksheets[0]
hdr = {c.value: c.column for c in ws[2] if c.value is not None}
rr = [c.row for c in ws['A'] if c.value == PN]; assert len(rr) == 1, rr
r = rr[0]
num = lambda s: float(s) if re.fullmatch(r'-?\d*\.\d+|-?\d+', s.strip()) else s
log = []
for c in cells:
    h = c['column']; cell = ws.cell(r, hdr[h])
    v = c['value'].split(' | before: ')[0]
    if h == '$COMMENT':
        old = cell.value or ''
        for add in (v, NOTE):
            if add not in (cell.value or ''):
                cell.value = (cell.value + ' | ' if cell.value else '') + add
        log.append((f'DesignTable_{T}.xlsx', PN, r, h, cell.coordinate, old, cell.value, 'note (tee_elbow) + A14 note'))
        continue
    assert cell.value in (None, ''), (h, cell.value, 'cell not empty')
    cell.value = num(v) if not h.startswith('$') else v
    log.append((f'DesignTable_{T}.xlsx', PN, r, h, cell.coordinate, '', cell.value, f"[{c['kind']}] {c['source'][:300]}"))

s = ws.cell(r, hdr['$PRP@DataStatus']); assert s.value == 'MISSING', s.value
s.value = 'ASSUMED'
log.append((f'DesignTable_{T}.xlsx', PN, r, '$PRP@DataStatus', s.coordinate, 'MISSING', 'ASSUMED',
            'A14 2026-10-20: every driving cell sourced (Brennan catalog + Brennan STEP bores = one source)'))
d = ws.cell(r, hdr['$PRP@DimSource']); assert d.value in (None, ''), d.value
d.value = 'BRENNAN'
ax = ws.cell(r, hdr['$PRP@AxialSource']); old = ax.value or ''
assert not old.startswith(TAG)
ax.value = (TAG + old + ' -> 2026-10-18 Brennan allowed (user: "yes use brennan but still differentiate"); D1 / D2 bores measured on '
            'BRENNAN_1603-08-08-08-FG.stp 2026-10-20 (A14); sources per cell in CHANGES.csv')
log.append((f'DesignTable_{T}.xlsx', PN, r, '$PRP@AxialSource', ax.coordinate, old, ax.value, 'marking + A14'))
log.append((f'DesignTable_{T}.xlsx', PN, r, '$PRP@DimSource', d.coordinate, '', 'BRENNAN', 'row takes sizes from Brennan'))
wb.save(P)
with open(os.path.join(STAGE, 'CHANGES.csv'), 'a', newline='', encoding='utf-8') as f:
    csv.writer(f).writerows(log)
print(len(log), 'log lines for', PN)
