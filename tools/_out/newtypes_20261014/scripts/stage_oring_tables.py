"""stage_oring_tables.py (2026-10-14): staged review tables DesignTable_4000.xlsx / DesignTable_4001.xlsx filled from Air-Way
series 218 ("O-rings For SAE J-1926/2/3 Inch Stud Ends", parts.air-way.com/api/series/218, fetched twice 2026-10-14 - the two
parsed passes are byte-identical; the rendered page showed the same values for its first 25 rows).

User 2026-10-14: "use airway sizes". Air-Way prints OMi's own part numbers (4000-8 = OMi 4000-08), so rows match by number.
ScopeAction: BUILD = on OMi's website (07_RAW_EXTRACTS/omi_feed/omi_skus_all.txt); an Air-Way row with no website SKU = DROP
(the convention of every review table; those parts are in OMi's parts file and asked about in the 2026-10-14 request).
Air-Way 4002 (CH Hi-Temp) is not an OMi series and is not carried.

python stage_oring_tables.py  ->  tools/_out/newtypes_20261014/review_tables/DesignTable_4000.xlsx, DesignTable_4001.xlsx"""
import os, re, sys
import openpyxl

R = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
AW = os.path.join(R, 'tools', '_out', 'airway_20261014')
FEED = os.path.join(R, '07_RAW_EXTRACTS', 'omi_feed', 'omi_skus_all.txt')
OMI = os.path.join(R, '07_RAW_EXTRACTS', 'omi_dimensions', 'WebsitePartsInfo2026.xlsx')
OUT = os.path.join(R, 'tools', '_out', 'newtypes_20261014', 'review_tables')
TODAY = '2026-10-14'

p1 = open(os.path.join(AW, 'series_218_pass1.txt'), encoding='utf-8').read()
p2 = open(os.path.join(AW, 'series_218_pass2.txt'), encoding='utf-8').read()
assert p1 == p2, 'Air-Way s218 passes differ - do not stage'
lines = [ln.split(' | ') for ln in p1.strip().split('\n')]
head, rows = lines[0], lines[1:]
col = {h: i for i, h in enumerate(head)}
assert [h for h in head[:10]] == ['Part #', 'Type', 'Tube Fitting Size', 'TUBE OD', 'I.D.', 'W', 'Temperature', 'Shore Hardness',
                                  'Elongation', 'Tensile'], head

web = {ln.strip() for ln in open(FEED, encoding='utf-8') if ln.strip()}
wb = openpyxl.load_workbook(OMI, read_only=True, data_only=True)
omi = {}
for i, r in enumerate(wb['Items'].iter_rows(values_only=True), start=1):
    if i > 1 and r[0]:
        omi[str(r[0]).strip()] = (i, str(r[3] or '').strip())

NUMTOL = re.compile(r'^\s*([0-9]*\.[0-9]+)\s*\u00b1\s*([0-9]*\.[0-9]+)\s*$')
MAT = {'CH (Buna-N)': 'Buna-N (nitrile), Air-Way compound CH', 'HK (Viton)': 'Viton (FKM), Air-Way compound HK'}


def num_tol(s):
    m = NUMTOL.match(s)
    assert m, ('not "<value> +/- <tol>"', s)
    return float(m.group(1)), m.group(2)


SRC = (f'Air-Way s218 only (O-rings For SAE J-1926/2/3 Inch Stud Ends; double-pass {len(rows)} rows, parts.air-way.com/api/series/218 '
       f'fetched twice {TODAY}, identical; rendered page cross-check 25 rows); Air-Way part numbers printed as 4000-n / 4001-n; config '
       f'names zero-padded to 4000-nn / 4001-nn [user {TODAY}: "use airway sizes"]')
HDR = [None, 'ID@DRIVE', 'W@DRIVE', '$PRP@TubeOD', '$PRP@PartNo', '$PRP@Material', '$PRP@Hardness', '$PRP@Temperature',
       '$PRP@Tolerance', '$PRP@DataStatus', '$PRP@AxialSource', '$COMMENT', '$PRP@ScopeAction', '$PRP@OMiSKU']

for ser in ('4000', '4001'):
    out = []
    for r in rows:
        m = re.match(rf'^{ser}-(\d+)$', r[0])
        if not m:
            continue
        pn = f'{ser}-{int(m.group(1)):02d}'
        idv, idt = num_tol(r[col['I.D.']])
        wv, wt = num_tol(r[col['W']])
        assert r[col['Tube Fitting Size']] == f'-{int(m.group(1))}', (pn, r[col['Tube Fitting Size']])
        mat = MAT[r[col['Type']]]
        on_web = pn in web
        o = omi.get(pn)
        cmt = (f'OMi parts file row {o[0]}: "{o[1]}" (no sizes)' if o else 'not in OMi parts file')
        if not on_web:
            cmt = ('DROP: Air-Way row with no matching OMi SKU in catalog feed (' + TODAY + ') | ' + cmt +
                   ' - asked in the 2026-10-14 request ("do you sell it?")')
        out.append([pn, idv, wv, r[col['TUBE OD']], pn, mat, r[col['Shore Hardness']].replace('  ', ' '),
                    r[col['Temperature']], f'I.D. +/-{idt}, W +/-{wt} (Air-Way)', 'ASSUMED', SRC, cmt,
                    'BUILD' if on_web else 'DROP', pn if on_web else None])
    missing = sorted(s for s in web if re.match(rf'^{ser}-\d\d$', s) and s not in {x[0] for x in out})
    assert not missing, ('website SKU with no Air-Way row', missing)
    nb = sum(1 for x in out if x[12] == 'BUILD')
    wbo = openpyxl.Workbook()
    ws = wbo.active
    ws.title = 'Sheet1'
    ws.append([f'OMi {ser} O-Ring ({"Buna-N" if ser == "4000" else "Viton"}) - Design Table v1 | AW s218 double-pass {len(out)} rows | '
               f'single-source | GENERATED {TODAY} (staged, newtypes_20261014)'])
    ws.append(HDR)
    for x in out:
        ws.append(x)
    os.makedirs(OUT, exist_ok=True)
    p = os.path.join(OUT, f'DesignTable_{ser}.xlsx')
    wbo.save(p)
    print(p, f'{len(out)} rows: {nb} BUILD, {len(out) - nb} DROP', [x[0] for x in out if x[12] == 'DROP'])
