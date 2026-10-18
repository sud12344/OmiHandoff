"""Stage OMi's 2026-10-18 answers into review tables (staging copy only; live 04_DESIGN_TABLES and the batch-16 staging untouched).

Inputs (read-only):
  07_RAW_EXTRACTS/omi_answers_20261018/zip_a, zip_b = the two zips the user's employer (OMi) sent, unpacked unchanged:
    - OMi_missing_sizes_and_typos_2026-10-09.xlsx = our 2026-10-09 request, returned with sheet "Lengths to check" column E
      ("Real length") and sheet "Possible typos" column F filled in by OMi (file modified 2026-10-14, by "Administrator").
    - 14 one-page PDF drawings (raster, Brennan-format sheets) = OMi's answer to "Sizes we need" rows 1, 2, 4, 5, 6, 8, 9, 10, 11,
      14, 15, 16, 17, 18 ("Your supplier's drawings or catalog pages are fine" - the request, question 3).
  ../source/drawings_transcribed.csv = every value read off those drawings by eye (one line per value, with what it measures).
  OMi WebsitePartsInfo2026.xlsx (sheet Items) for the hex / Side 1 ID of the 6410 rows whose length OMi now gave.
Base tables: tools/_out/wave1_ship_20261009/scratch_root/04_DESIGN_TABLES (= batch 16 staged; for these 7 roots identical to the live
04_DESIGN_TABLES except DesignTable_1502.xlsx, which batch 16 changes) -> copied whole to ../scratch_root/04_DESIGN_TABLES, edited there.

User 2026-10-18: "i attached everything my employer sent this is verified information, use it to fill out more tables".
DataStatus stays by the rules in CLAUDE.md: OMi alone (= Brennan data, evidence 2026-09-26: one source) -> ASSUMED; the user's
word "verified" is reported back as a question (VERIFIED needs 2 independent sources by the project rule).
Only empty cells are filled (asserted), except DataStatus and the appended notes.
"""
import os, re, csv, shutil, hashlib
import openpyxl

ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
WORK = os.path.join(ROOT, r'tools\_out\omi_answers_20261018')
BASE = os.path.join(ROOT, r'tools\_out\wave1_ship_20261009\scratch_root\04_DESIGN_TABLES')
DT = os.path.join(WORK, r'scratch_root\04_DESIGN_TABLES')
ANS = os.path.join(ROOT, r'07_RAW_EXTRACTS\omi_answers_20261018')
OMI = os.path.join(ROOT, r'07_RAW_EXTRACTS\omi_dimensions\WebsitePartsInfo2026.xlsx')
TAG = 'OMi answer 2026-10-18 (staged, not shipped)'
USER = 'user 2026-10-18: "this is verified information, use it"'

if os.path.exists(DT):
    raise SystemExit(f'{DT} exists - stage into a fresh folder')
shutil.copytree(BASE, DT)

# ---- the transcribed drawing values -------------------------------------------------------------------------------------------
TR = {}
with open(os.path.join(WORK, r'source\drawings_transcribed.csv'), newline='', encoding='utf-8') as f:
    for r in csv.DictReader(f):
        TR[(r['part_number'], r['feature'])] = r


def dv(pn, feat):
    r = TR[(pn, feat)]
    return float(r['value']), f"OMi drawing {r['pdf']} (sent 2026-10-18): {r['printed']} = {r['meaning']}"


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()[:16]


# ---- end-library values per dash (user 2026-10-08 Q-TRANSFER-ENDLIB = yes: reuse end details of the built masters) -----------
NPT_OD = {2: 0.405, 4: 0.54, 6: 0.675, 8: 0.84, 12: 1.05, 16: 1.315, 20: 1.66, 24: 1.9, 32: 2.375}    # ASME B1.20.1 pipe OD (ms51500 NPTF table; the 1405/1502 Air-Way rows carry the same for A and A1)
TAPDRILL = {2: 0.33, 4: 0.43, 6: 0.57, 8: 0.7, 12: 0.91, 16: 1.14, 20: 1.48, 24: 1.72, 32: 2.19}     # end-library female-pipe tap drill (every built port of the size; = Air-Way per-size value)
PIPE = {2: '1/8-27', 4: '1/4-18', 6: '3/8-18', 8: '1/2-14', 12: '3/4-14', 16: '1-11 1/2', 20: '1 1/4-11 1/2', 24: '1 1/2-11 1/2', 32: '2-11 1/2'}

CH = []   # (table, pn, header, new, source)


def put(t, pn, h, v, src):
    CH.append((t, pn, h, v, src))


def ends(t, pn, a, a1, cols):
    """tap drill + thread majors + thread text for a female NPTF (dash a) x NPSM swivel (dash a1) part"""
    put(t, pn, 'TAPDRILL@BORE', TAPDRILL[a], f'end-library tap drill of the -{a:02d} female pipe port = {TAPDRILL[a]} (as every built port of that size; Q-TRANSFER-ENDLIB yes 2026-10-08)')
    put(t, pn, 'THDMAJ_A@DRIVE', NPT_OD[a], f'pipe OD of the -{a:02d} NPTF port ({PIPE[a]} NPTF on the OMi drawing; ASME B1.20.1, tools/generators/ms51500.py NPTF[{a}])')
    for c in cols:
        put(t, pn, c, NPT_OD[a1], f'pipe OD of the -{a1:02d} NPSM swivel thread ({PIPE[a1]} NPSM on the OMi drawing; the value the 1405 / 1502 Air-Way rows carry for the NPSM end)')


# ---- 1405 straight swivel (gen_straight2 SPEC: W1 = port face to nut face, L1 = port face to stem tip, C1 body hex, C2 nut hex, D bore)
for pn, a, a1 in (('1405-04-06', 4, 6), ('1405-08-12', 8, 12)):
    for h, feat in (('L1@DRIVE', 'face_to_stemtip'), ('W1@HEX', 'overall'), ('C1HEX@HEX', 'hex'), ('C2HEX@HEX', 'nut'), ('D@BORE', 'bore')):
        v, s = dv(pn, feat); put('1405', pn, h, v, s)
    ends('1405', pn, a, a1, ['THDMAJ_A1@DRIVE'])
    put('1405', pn, '$PRP@Threads', f'{PIPE[a]} (A NPTF) x {PIPE[a1]} (A1 NPSM)', 'OMi drawing thread callouts')

# ---- 1502 90 deg forged swivel elbow (gen_tee SPEC 1502: N3 = crossing to port face, M1 = crossing to stem tip, W3 = crossing to nut face,
#      CHEX nut hex, C1HEX body hex, D swivel bore = Brennan "A2 DRILL" (column comment of 1502-02-02-FG))
for pn, a, a1 in (('1502-06-06-FG', 6, 6), ('1502-12-08-FG', 12, 8), ('1502-12-12-FG', 12, 12), ('1502-20-20-FG', 20, 20),
                  ('1502-24-24-FG', 24, 24), ('1502-32-32-FG', 32, 32)):
    for h, feat in (('N3@DRIVE', 'leg_port'), ('M1@DRIVE', 'leg_stemtip'), ('W3@HEX', 'leg_nutface'), ('CHEX@HEX', 'nut'), ('C1HEX@HEX', 'hex'), ('D@BORE', 'bore')):
        v, s = dv(pn, feat); put('1502', pn, h, v, s)
    ends('1502', pn, a, a1, ['THDMAJ_A1@DRIVE'])
    put('1502', pn, '$PRP@Threads', f'{PIPE[a]} (A NPTF) x {PIPE[a1]} (A1 NPSM)', 'OMi drawing thread callouts')

# ---- 1504 45 deg forged swivel elbow (no generator SPEC yet; Air-Way letters N2 / M / W2 read as 1502's N3 / M1 / W3: port leg /
#      stem tip / nut face - the drawing measures all three from the axis crossing; CHEX nut, C1HEX body as in the Air-Way 1504 rows)
pn = '1504-20-20-FG'
for h, feat in (('N2@DRIVE', 'leg_port'), ('M@DRIVE', 'leg_stemtip'), ('W2@HEX', 'leg_nutface'), ('CHEX@HEX', 'nut'), ('C1HEX@HEX', 'hex'), ('D@BORE', 'bore')):
    v, s = dv(pn, feat); put('1504', pn, h, v, s)
ends('1504', pn, 20, 20, ['THDMAJ_A1@DRIVE'])
put('1504', pn, '$PRP@Threads', f'{PIPE[20]} (A NPTF) x {PIPE[20]} (A1 NPSM)', 'OMi drawing thread callouts')

# ---- 1602 forged swivel branch tee (no generator SPEC yet; letters as 1502: N3 branch port leg, M1_2 / M1_3 stem tips, W3_2 / W3_3 nut faces;
#      C2HEX / C3HEX nuts, C1HEX body (Air-Way 1602-12-08-08: C1HEX follows the -12 NPTF), D2 / D3 swivel bores)
for pn, d in (('1602-02-02-02-FG', 2), ('1602-04-04-04-FG', 4), ('1602-08-08-08-FG', 8), ('1602-16-16-16-FG', 16), ('1602-20-20-20-FG', 20)):
    for h, feat in (('N3@DRIVE', 'leg_port'), ('M1_2@DRIVE', 'leg_stemtip'), ('M1_3@DRIVE', 'leg_stemtip'), ('W3_2@HEX', 'leg_nutface'),
                    ('W3_3@HEX', 'leg_nutface'), ('C2HEX@HEX', 'nut'), ('C3HEX@HEX', 'nut2'), ('C1HEX@HEX', 'hex'), ('D2@BORE', 'bore'), ('D3@BORE', 'bore2')):
        v, s = dv(pn, feat); put('1602', pn, h, v, s)
    ends('1602', pn, d, d, ['THDMAJ_A1_2@DRIVE', 'THDMAJ_A1_3@DRIVE'])
    put('1602', pn, '$PRP@Threads', f'{PIPE[d]} (A NPTF) x {PIPE[d]} (A1 (2) NPSM) x {PIPE[d]} (A1 (3) NPSM)', 'OMi drawing thread callouts')

DRAWN = sorted({(t, pn) for t, pn, *_ in CH})

# ---- "Lengths to check" (sheet column E) ----------------------------------------------------------------------------------------
lw = openpyxl.load_workbook(os.path.join(ANS, r'zip_a\OMi_missing_sizes_and_typos_2026-10-09.xlsx'))['Lengths to check']
REAL = {}
for r in range(3, lw.max_row + 1):
    REAL[lw[f'A{r}'].value] = (float(lw[f'E{r}'].value), lw[f'B{r}'].value, lw[f'C{r}'].value, f'E{r}')
assert len(REAL) == 14, REAL

omi = openpyxl.load_workbook(OMI, data_only=True)['Items']


def omi_row(row, pn):
    assert str(omi[f'A{row}'].value).strip() == pn, (row, omi[f'A{row}'].value, pn)
    return {k: omi[f'{k}{row}'].value for k in ('L', 'R', 'S', 'U', 'AA')}


def lsrc(pn):
    v, row, old, cell = REAL[pn]
    return v, (f"OMi answer: OMi_missing_sizes_and_typos_2026-10-09.xlsx (returned 2026-10-18) 'Lengths to check'!{cell} Real length = {v} "
               f"(website file Items row {row} said {old}, a filler value)")


# 6410 rows with no Air-Way row: length now from OMi; hex + Side 1 ID from the website file; ORB end values per dash from the table's own
# built rows (end library). D@BORE only where it fits inside the female port drill D3 (Q-BOREWIDE, question 2 of 2026-09-26, still open).
NEW6410 = {'6410-16-04-O': (1862, 16, 4), '6410-20-06-O': (1870, 20, 6), '6410-20-08-O': (1871, 20, 8),
           '6410-32-12-O': (1879, 32, 12), '6410-32-16-O': (1880, 32, 16)}
w10 = openpyxl.load_workbook(os.path.join(DT, 'DesignTable_6410.xlsx')).active
h10 = {c.value: c.column for c in w10[2] if c.value is not None}


def per_dash(male, female):
    """end values of male ORB dash `male` and female ORB port dash `female` from BUILD rows that carry them; assert one value each"""
    mv, fv = {}, {}
    for r in range(3, w10.max_row + 1):
        pn = w10.cell(r, 1).value
        if not pn or w10.cell(r, h10['$PRP@ScopeAction']).value != 'BUILD' or w10.cell(r, h10['L33@DRIVE']).value is None:
            continue
        m, fdash = int(pn.split('-')[1]), int(pn.split('-')[2])
        if m == male:
            for h in ('THDMAJ_B1@DRIVE', 'ORB_THDL@DRIVE', 'ORB_GRVDIA@DRIVE', 'ORB_GRVW@DRIVE'):
                mv.setdefault(h, set()).add(w10.cell(r, h10[h]).value)
        if fdash == female:
            for h in ('THDMAJ_B2@DRIVE', 'D3@BORE'):
                fv.setdefault(h, set()).add(w10.cell(r, h10[h]).value)
    out = {}
    for h, s in list(mv.items()) + list(fv.items()):
        s2 = {x for x in s if x is not None}
        assert len(s2) <= 1, (male, female, h, s)
        out[h] = s2.pop() if s2 else None
    return out


def thr(s):
    """'1-5/16-12 UN' -> '1 5/16-12', '7/16-20 UNF' -> '7/16-20' (the table's thread form)"""
    s = re.sub(r'\s+UNF?$', '', s.strip())
    return re.sub(r'^(\d)-(\d+/\d+)-(\d+)$', r'\1 \2-\3', s)


HELD = []
for pn, (row, m, fd) in NEW6410.items():
    o = omi_row(row, pn)
    v, s = lsrc(pn)
    put('6410', pn, 'L33@DRIVE', v, s)
    put('6410', pn, 'C1HEX@HEX', float(o['R']), f'OMi WebsitePartsInfo2026.xlsx Items!R{row} Hex Flat = {o["R"]} (mapping certain: body hex, as Task A on 6410)')
    pdv = per_dash(m, fd)
    for h, val in pdv.items():
        if val is None:
            continue
        put('6410', pn, h, val, f'end-library value of the {"male ORB -%02d end" % m if h in ("THDMAJ_B1@DRIVE", "ORB_THDL@DRIVE", "ORB_GRVDIA@DRIVE", "ORB_GRVW@DRIVE") else "female ORB -%02d port" % fd}'
                                f' = {val}, the value every built 6410 row with that end carries (Q-TRANSFER-ENDLIB yes 2026-10-08)')
    sid = float(o['S'])
    d3 = pdv.get('D3@BORE')
    if d3 is not None and sid <= d3 + 1e-9:
        put('6410', pn, 'D@BORE', sid, f'OMi WebsitePartsInfo2026.xlsx Items!S{row} Side 1 ID = {o["S"]} (fits inside the -{fd:02d} port drill D3 {d3})')
    else:
        HELD.append((pn, f'Side 1 ID {o["S"]} (Items!S{row}) wider than the -{fd:02d} port drill D3 {d3} -> D@BORE left empty (Q-BOREWIDE, same as the 10 held rows)'))
    put('6410', pn, '$PRP@Threads', f'{thr(o["U"])} (male ORB) x {thr(o["AA"])} (female ORB)',
        f'OMi WebsitePartsInfo2026.xlsx Items!U{row} "{o["U"]}" / AA{row} "{o["AA"]}" (written in the table\'s form)')

# ---- apply ------------------------------------------------------------------------------------------------------------------------
books, log = {}, []


def book(t):
    if t not in books:
        p = os.path.join(DT, f'DesignTable_{t}.xlsx'); books[t] = (p, openpyxl.load_workbook(p))
    return books[t][1].active


def rowof(ws, pn):
    rr = [c.row for c in ws['A'] if c.value == pn]
    assert len(rr) == 1, (pn, rr)
    return rr[0]


def fmt(v):
    return '' if v is None else repr(v) if isinstance(v, float) else str(v)


for t, pn, h, v, src in CH:
    ws = book(t)
    hdr = {c.value: c.column for c in ws[2] if c.value is not None}
    r = rowof(ws, pn)
    cell = ws.cell(r, hdr[h])
    assert cell.value is None, (t, pn, h, cell.value)
    cell.value = v
    log.append((f'DesignTable_{t}.xlsx', pn, r, h, cell.coordinate, '', fmt(v), src))


def note(t, pn, ax_add, cm_add, status=None, why=''):
    ws = book(t)
    hdr = {c.value: c.column for c in ws[2] if c.value is not None}
    r = rowof(ws, pn)
    for h, add, sep in (('$PRP@AxialSource', ax_add, ' -> '), ('$COMMENT', cm_add, ' | ')):
        if not add:
            continue
        cell = ws.cell(r, hdr[h]); old = cell.value
        cell.value = ((old + sep) if old else '') + add
        log.append((f'DesignTable_{t}.xlsx', pn, r, h, cell.coordinate, fmt(old), cell.value, 'note: sources of this row' if h == '$PRP@AxialSource' else 'note (appended)'))
    if status:
        cell = ws.cell(r, hdr['$PRP@DataStatus']); old = cell.value; cell.value = status
        log.append((f'DesignTable_{t}.xlsx', pn, r, '$PRP@DataStatus', cell.coordinate, fmt(old), status, why))


for t, pn in DRAWN:
    pdf = TR[(pn, 'bore')]['pdf']
    note(t, pn,
         f'{TAG}: every size from OMi\'s drawing {pdf} (07_RAW_EXTRACTS/omi_answers_20261018, sha256 {sha(os.path.join(ANS, pdf))}...; a Brennan-format '
         f'sheet OMi sent in answer to our 2026-10-09 size request; {USER}); one source (OMi = Brennan data, evidence 2026-09-26) -> ASSUMED; '
         f'tap drill / thread majors from the end library (Q-TRANSFER-ENDLIB); end and body SHAPE from the Brennan 3D files (shape only, user 2026-10-08)',
         f'{TAG}: was MISSING (no Air-Way row / non-forged Air-Way row only); now filled from OMi\'s own drawing; non-forged Air-Way row NOT used',
         'ASSUMED', 'every driving cell now has a source: OMi drawing (one source) + end library -> ASSUMED')

for pn, (v, row, old, cell) in REAL.items():
    if pn in NEW6410:
        held = dict(HELD).get(pn)
        note('6410', pn,
             f'{TAG}: L33 = OMi real length (Lengths to check!{cell}); hex / Side 1 ID OMi website file Items row {row}; ORB end values from the end library '
             f'(Q-TRANSFER-ENDLIB); {USER}; one source -> ' + ('MISSING (held, below)' if held else 'ASSUMED'),
             f'{TAG}: website length {old} was a filler; OMi gave {v}' + (f'; {held}' if held else ''),
             'MISSING' if held else 'ASSUMED',
             'held: ' + held if held else 'every driving cell now has a source (OMi + end library) -> ASSUMED')
        continue
    t = pn.split('-')[0]
    ws = book(t)
    hdr = {c.value: c.column for c in ws[2] if c.value is not None}
    r = rowof(ws, pn)
    if t == '6420':
        note(t, pn, None, f'{TAG}: OMi real Overall Length = {v} (Lengths to check!{cell}; website file said {old}, = the row above). Still no geometry source '
                          f'(6420 has no Air-Way series and no generator SPEC) -> stays MISSING')
        continue
    lc = {'6405': 'L31@DRIVE', '6410': 'L33@DRIVE'}[t]
    cur = ws.cell(r, hdr[lc]).value
    assert cur is not None and abs(cur - v) < 0.005, (pn, cur, v)           # OMi's real length = the Air-Way value already in the table
    note(t, pn, None, f'{TAG}: OMi\'s real Overall Length {v} (Lengths to check!{cell}; website file said {old}, a filler) = the Air-Way {lc.split("@")[0]} {cur} '
                      f'already in this row -> no change; OMi and Air-Way agree on the length')

for t, (p, wb) in books.items():
    wb.save(p)

with open(os.path.join(WORK, 'CHANGES.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(['table', 'part_number', 'excel_row', 'column', 'cell', 'old', 'new', 'source'])
    for x in sorted(log, key=lambda x: (x[0], x[2], x[4])):
        w.writerow(x)
print(len(log), 'cells written;', len(DRAWN), 'rows from drawings;', len(NEW6410), '6410 rows from lengths;', 'held:', HELD)
