"""build_tables.py (group C, 2026-10-21): the compression (bite-type, SAE J514/2 flareless) review tables DesignTable_C<root>.xlsx.

Rows = OMi's 230 compression part numbers (07_RAW_EXTRACTS/omi_feed/omi_skus_compression.txt) - every one BUILD (OMi decides which
parts exist). Dimensions: Air-Way has no compression series and OMi's website parts file prints only the thread callouts (Items, Side n
Thread); the only dimension source is Brennan's CAD parameter record of the complete fitting (tools/_out/partcommunity_20261021/
pc_final.csv, exact / alias matches; user 2026-10-21 "yes you can use them following my rule") -> $PRP@DimSource BRENNAN, DataStatus
ASSUMED (one source). A part without a record keeps empty cells and DataStatus MISSING (it is in the unsent Triveni package).
"-M1" = OMi's "(Body)" = the fitting sold without nut and sleeve: the body cells of the complete fitting of the same size, ASSUMED (OMi's
description; the question is in the unsent Triveni email 10_DRAWINGS/for_OMi_2026-10-21). Nut C0318 / sleeve C0319 = separate parts.
Record letters were decoded on the ten Brennan -08 STEPs (README.md, measure/measure_c.txt). Writes stage/04_DESIGN_TABLES/ + rows.csv."""
import os, re, sys, csv, openpyxl
from openpyxl.styles import Font
H = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, H)
import records as R
STAGE = R.STAGE; OUTD = os.path.join(STAGE, 'stage', '04_DESIGN_TABLES'); os.makedirs(OUTD, exist_ok=True)
TODAY = '2026-10-21'
REC = R.records(); LIB = R.bite_lib()
OMI = os.path.join(R.ROOT, '07_RAW_EXTRACTS', 'omi_dimensions', 'WebsitePartsInfo2026.xlsx')
_ws = openpyxl.load_workbook(OMI, read_only=True)['Items']; _rows = list(_ws.iter_rows(values_only=True)); _h = _rows[0]
OMIF = {r[0]: dict(zip(_h, r)) for r in _rows[1:] if r and r[0]}

g = lambda d, k: R.num(d.get(k)) if d else None
endl = lambda d: round(g(d, 'BB1') + g(d, 'MB1'), 4) if d and g(d, 'BB1') is not None else None
BITE = [('BITE_SOCKD@DRIVE', 'AB1'), ('BITE_SOCKL@DRIVE', 'KB1'), ('BITE_NECKD@DRIVE', 'EB1'), ('BITE_NECKL@DRIVE', 'CB1')]
# per root: title, columns [(column, record code | callable(rec) | None)], thread text
T = {
 'C2402': ('Compression Male Connector (bite type x male 37 deg JIC)',
           [('L@DRIVE', 'L'), ('HEXB@HEX', 'C'), ('THDMAJ_A@DRIVE', 'B1THD'), ('THDMAJ_B@DRIVE', 'B2THD'), ('D1@BORE', 'B1D'), ('D2@BORE', 'B2D')] + BITE +
           [('BITE_ENDL@DRIVE', endl)]),
 'C2403': ('Compression Union (bite type x bite type)',
           [('L@DRIVE', 'L'), ('HEXB@HEX', 'C'), ('THDMAJ_A@DRIVE', 'B1THD'), ('THDMAJ_B@DRIVE', 'B2THD'), ('D1@BORE', 'B1D'), ('D2@BORE', 'B2D')] + BITE +
           [('BITE_ENDL@DRIVE', endl)]),
 'C2404': ('Compression Male Pipe Adapter (bite type x male NPTF)',
           [('L@DRIVE', 'L'), ('HEXB@HEX', 'C'), ('THDMAJ_A@DRIVE', 'B1THD'), ('PIPEOD@DRIVE', 'A2THD'), ('D1@BORE', 'B1D'), ('D2@BORE', 'A2D')] + BITE +
           [('BITE_ENDL@DRIVE', endl)]),
 'C2405': ('Compression Female Pipe Adapter (bite type x female NPTF)',
           [('L@DRIVE', 'L'), ('HEXB@HEX', 'C'), ('THDMAJ_A@DRIVE', 'B1THD'), ('PIPEOD@DRIVE', 'A2THD'), ('D1@BORE', 'B1D'), ('TAPDRILL@BORE', 'CA2'),
            ('FPT_P1@DRIVE', 'EA2'), ('PORTTHDL@REF', 'BA2'), ('PORTCSKD@REF', 'DA2')] + BITE + [('BITE_ENDL@DRIVE', endl)]),
 'C2500': ('Compression 90 deg Union Elbow, forged (bite type x bite type)',
           [('M1@DRIVE', 'M1'), ('M2@DRIVE', 'M2'), ('Y@REF', 'Y'), ('THDMAJ_B1@DRIVE', 'B1THD'), ('THDMAJ_B2@DRIVE', 'B1THD'), ('D1@BORE', 'B1D'), ('D2@BORE', 'B2D')] + BITE +
           [('BITE_ENDL@DRIVE', endl)]),
 'C2501': ('Compression 90 deg Male Pipe Elbow, forged (bite type x male NPTF)',
           [('M@DRIVE', 'M'), ('N@DRIVE', 'N'), ('Y@REF', 'Y'), ('THDMAJ_A@DRIVE', 'B1THD'), ('PIPEOD@DRIVE', 'A2THD'), ('D@BORE', 'B1D'), ('D1@DRILL', 'A2D')] + BITE +
           [('BITE_ENDL@DRIVE', endl)]),
 'C2502': ('Compression 90 deg Female Pipe Elbow, forged (bite type x female NPTF)',
           [('M@DRIVE', 'M'), ('N@DRIVE', 'N'), ('Y@REF', 'Y'), ('THDMAJ_A@DRIVE', 'B1THD'), ('PIPEOD@DRIVE', 'A2THD'), ('D@BORE', 'B1D'), ('TAPDRILL@BORE', 'CA2'),
            ('PORTDEPTH@BORE', 'BA2'), ('PORTCSKD@REF', 'DA2')] + BITE + [('BITE_ENDL@DRIVE', endl)]),
 'C2601': ('Compression Male Pipe Branch Tee, forged (bite type run x male NPTF branch)',
           [('M1@DRIVE', 'M1'), ('M2@DRIVE', 'M2'), ('N@DRIVE', 'N'), ('Y@REF', 'Y'), ('THDMAJ_B1@DRIVE', 'B1THD'), ('THDMAJ_B2@DRIVE', 'B1THD'),
            ('THDMAJ_A@DRIVE', 'A3THD'), ('D1@BORE', 'B1D'), ('D2@BORE', 'B2D'), ('DDRILL@BORE', 'A3D')] + BITE + [('BITE_ENDL@DRIVE', endl)]),
 'C6801': ('Compression 90 deg Adjustable Male ORB Elbow, forged (bite type x adjustable male ORB)',
           [('LEG_BITE@DRIVE', 'M'), ('LEG_ORB@DRIVE', 'M7'), ('Y@REF', 'Y'), ('HEXB@HEX', 'C'), ('THDMAJ_A@DRIVE', 'B1THD'), ('THDMAJ_B@DRIVE', 'B2THD'),
            ('D1@BORE', 'B1D'), ('D2@BORE', 'B2D')] + BITE + [('BITE_ENDL@DRIVE', endl)]),
 'C0318': ('Compression Tube Nut (bite type)',
           [('L1@DRIVE', 'L'), ('CHEX@HEX', 'C'), ('THDMAJ_B@DRIVE', 'BTHD'), ('THDL@DRIVE', 'F'), ('CONEX@DRIVE', 'E'), ('TUBEHOLE@BORE', 'D'),
            ('CSKDIA@DRIVE', 'B'), ('CHAMDIA@REF', 'A')]),
 'C0319': ('Compression Sleeve / Ferrule (bite type)',
           # 2026-10-21 (user: "I sent the files below for question 4"): every DRIVE cell = the row's own Brennan C0319 STEP (measure/c0319_profile.csv,
           # scripts/measure_c0319.py); P_i = (x_i, r_i) of the 14-line half-section; the record letters are kept in $COMMENT and cross-checked
           [(f'{nm}@DRIVE', (lambda d, k=k: SLV[d['_pn']][k])) for nm, k in (
               ('NOSER', 'r1'), ('BORER', 'r0'), ('GROOVER', 'r3'), ('COLLARR', 'r6'), ('SPIGOTR', 'r8'), ('REARIR', 'r11'), ('APEXR', 'r2'),
               ('STEPR', 'r5'), ('TAILR', 'r10'), ('APEXX', 'x2'), ('GROOVEX', 'x3'), ('CHAMX', 'x6'), ('COLLARX', 'x7'), ('SPIGOTX', 'x8'),
               ('TAILX', 'x9'), ('REARIX', 'x12'), ('BOREX', 'x13'), ('STEPX', 'x4'), ('L', 'x10'))]),
}
SLV = {r['part']: {k: (float(v) if k[0] in 'xr' and k[1:].isdigit() else v) for k, v in r.items()}
       for r in csv.DictReader(open(os.path.join(STAGE, 'measure', 'c0319_profile.csv')))}
SERIES = {k: k for k in T}
# record letters -> meaning, decoded on the ten Brennan -08 STEPs (measure/measure_c.txt); written into $PRP@AxialSource of every filled row
DECODE = {
 'bite': 'bite-type end: AB tube socket dia x KB socket depth from the body tip, CB / EB relief neck length / dia behind the thread, BB + MB = '
         'body tip to hex face (BB = hex face to nut face, MB = nut face to body tip, assembled), B1THD thread major, B1D through bore '
         '(decoded on BRENNAN_C2402/C2403/C2405-08 and the -FG -08 STEPs: 0.514 x 0.305 socket, neck 0.660 x 0.094, 0.562 end; equal on every record of a dash)',
 'L': 'L = body tip to the far tip / port face (Brennan C2402-08-08 1.500, C2403 1.422, C2405 1.470 = the STEPs)',
 'C': 'C = body hex across flats (STEP hex corners)',
 'M': 'M / M1 / M2 = axis crossing to the bite body tip; N = crossing to the pipe tip / port face; Y = forged hex A/F; M7 = crossing to the ORB tip (C6801)',
}


def omi_thread(pn, side):
    r = OMIF.get(pn) or {}
    return r.get(f'Side {side} Thread')


def body_pn(pn):
    return pn[:-3] if pn.endswith('-M1') else pn


def series(pn):
    return re.match(r'C\d{4}', pn).group(0)


rows_out = []
by = {}
for pn in R.skus(): by.setdefault(series(pn), []).append(pn)
assert sum(len(v) for v in by.values()) == 230 and set(by) == set(T), (len(by), set(by) ^ set(T))
for root, pns in by.items():
    title, cols = T[root]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'DesignTable'
    hdr = [None] + [c for c, _ in cols] + ['$PRP@PartNo', '$PRP@Threads', '$PRP@DataStatus', '$PRP@AxialSource', '$COMMENT', '$PRP@ScopeAction',
                                           '$PRP@OMiSKU', '$PRP@DimSource', '$PRP@BrennanRecord'] + (['$PRP@TubeOD'] if root == 'C0319' else [])
    ws.append([f'OMi {root} {title} - review table (group C stage {TODAY}) | dimensions: Brennan CAD parameter records (PARTcommunity, '
               f'user 2026-10-21) - no Air-Way series, OMi file prints threads only | NOT SHIPPED'])
    ws.append(hdr)
    n_f = n_m = 0
    for pn in pns:
        bp = body_pn(pn); d = REC.get(bp)
        thr = ' x '.join(t for t in (omi_thread(pn, 1), omi_thread(pn, 2), omi_thread(pn, 3)) if t)
        vals = []
        for c, k in cols:
            v = k(d) if callable(k) else (g(d, k) if d else None)
            vals.append(v)
        # Brennan record inconsistency (C0318-24: BTHD 1.750 printed with callout 1 7/8-12, B 1.880, D 1.510): the callout wins, flagged
        if d:
            st, src = 'ASSUMED', 'BRENNAN'
            axs = (f'[BRENNAN-CAD {TODAY}: every dimension cell] Brennan CAD parameter record {d["_pn"]} ({d["_family"]}, PARTcommunity reversemap, '
                   f'{d["_match"]} match) - single source (no Air-Way compression series; OMi file prints threads only)')
            if root == 'C0319':
                axs = (f'[BRENNAN-STEP {TODAY}: every DRIVE cell measured on BRENNAN_{d["_pn"]}.stp] Brennan 3D model of the part (user upload 2026-10-21, '
                       f'PARTcommunity) - single source (no Air-Way compression series; OMi file prints threads only); the 14-line half-section read by '
                       f'groupC scripts/measure_c0319.py; Brennan record letters vs the STEP: {SLV[d["_pn"]]["record_check"]} (G / M are not exactly the groove '
                       f'start / end on the STEP - the STEP is used)')
            if root not in ('C0318', 'C0319'): axs += '; ' + DECODE['bite'] + '; ' + (DECODE['L'] + '; ' + DECODE['C'] if root.startswith('C24') else DECODE['M'])
            cm = f'Brennan record {d["_pn"]}: ' + ' '.join(f'{k}={v}' for k, v in d.items() if not k.startswith(('_', 'MAT', 'WBVAR', 'FONT', 'PLATING', 'WEIGHT', 'TT', 'B21', 'A21', 'A31')))
            if pn.endswith('-M1'):
                st = 'ASSUMED'
                axs += (f'; -M1 = OMi "(Body)": the fitting sold WITHOUT nut and sleeve - body cells = those of the complete fitting {bp} (ASSUMPTION: same '
                        f'body; question to OMi in the unsent Triveni email 10_DRAWINGS/for_OMi_2026-10-21)')
            n_f += 1
        else:
            st, src, axs = 'MISSING', None, ('no dimension source: Air-Way has no compression series, OMi file prints threads only, Brennan has no CAD '
                                             'record for ' + bp + ' - in the OMi / Triveni package (10_DRAWINGS/for_OMi_2026-10-21, not sent)')
            cm = 'MISSING: needs OMi dimensions (or a Brennan record / STEP)'
            n_m += 1
        if d and root == 'C0318' and d['_pn'] == 'C0318-24':
            i = [c for c, _ in cols].index('THDMAJ_B@DRIVE'); vals[i] = 1.875
            cm += (' | THDMAJ_B 1.875 = the record\'s own callout B1 1 7/8-12 (OMi file: 1 7/8-12 UN); the record\'s BTHD 1.750 contradicts its callout, its '
                   'countersink B 1.880 and the bite-end library -24 thread 1.875 (C2403-24-24 B1THD) - record typo, flagged for the user')
        if root == 'C0319': tube_ = {4: '1/4', 6: '3/8', 8: '1/2', 10: '5/8', 12: '3/4', 16: '1', 20: '1 1/4', 24: '1 1/2', 32: '2'}[int(pn.split('-')[1])]
        ws.append([pn] + vals + [pn, thr or None, st, axs, cm, 'BUILD', pn, src, d['_pn'] if d else None] + ([tube_] if root == 'C0319' else []))
        rows_out.append(dict(root=root, part=pn, status=st, source=src or '', record=d['_pn'] if d else ''))
    ws['A1'].font = Font(bold=True)
    wb.save(os.path.join(OUTD, f'DesignTable_{root}.xlsx'))
    print(f'{root}: {len(pns)} rows, {n_f} with Brennan cells, {n_m} MISSING')
with open(os.path.join(STAGE, 'rows.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=list(rows_out[0])); w.writeheader(); w.writerows(rows_out)
print('total', len(rows_out))
