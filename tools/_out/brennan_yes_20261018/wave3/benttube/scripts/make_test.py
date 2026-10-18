"""make_test.py (wave 3 bent tube): gen_bent test tables from the stage -Tube rows + cells.csv, in three variants:
  tables_real   - cells applied, DataStatus as in the stage (MISSING: no row is ready)          -> run A
  tables_probe  - same, DataStatus forced ASSUMED (gate-off probe: what the generator says)      -> run B
  tables_whatif - same + TEST placeholders for the inputs no source gives (TUBEOD, RB, HEXL, D), DataStatus TEST (GB_TEST=1) -> runs C / D
                  NOT part data: shows only whether the catalog lengths fit the model type.
Also writes gen_bent copies: gen/gen_bent.py (unchanged copy) and gen/gen_bent_c62cat.py (FLG62_MEAS extended with the Brennan 3300-62
catalog heads -12..-32: O = A, TH = D, GO = H/2, GI = GO - 0.155, GD = 0.110 - the 3300 drawing's groove; WHAT-IF only)."""
import csv, json, math, os, re, shutil
import openpyxl

H = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
W = os.path.join(H, r'tools\_out\brennan_yes_20261018\wave3\benttube')
T = os.path.join(W, 'test')
STAGE = os.path.join(H, r'tools\_out\brennan_yes_20261018\stage\04_DESIGN_TABLES')
GB = os.path.join(H, r'tools\_out\newtypes_20261014\generators\gen_bent.py')
cells = list(csv.DictReader(open(os.path.join(W, 'cells.csv'), encoding='utf-8')))
pairs = list(csv.DictReader(open(os.path.join(W, 'pairs_45_90.csv'), encoding='utf-8')))
CAT = json.load(open(os.path.join(H, r'tools\_out\brennan_catalog_20261001\catalog_verified.json')))

HDR = [None, 'M1@DRIVE', 'L1@DRIVE', 'CHEX@HEX', 'HEXL@HEX', 'O@DRIVE', 'D@BORE', 'TUBEOD@TUBE', 'RB@PATH', '$PRP@FlangeCode',
       '$PRP@TubeOD', '$PRP@PartNo', '$PRP@DataStatus', '$PRP@AxialSource', '$COMMENT', '$PRP@ScopeAction', '$PRP@OMiSKU']
HDR_SW = [None, 'M1@DRIVE', 'L1@DRIVE', 'CHEX@HEX', 'HEXL@HEX', 'NUTHEX@NUTHEXSK', 'NOSETIP@SWSK', 'D@BORE', 'TUBEOD@TUBE', 'RB@PATH',
          '$PRP@TubeOD', '$PRP@PartNo', '$PRP@DataStatus', '$PRP@AxialSource', '$COMMENT', '$PRP@ScopeAction', '$PRP@OMiSKU']
GROUPS = {'1704T': ('1704', '1804'), '1703T': ('1703', '1803'), '6701T': ('6701',)}

C = {}
for c in cells:
    C.setdefault((c['table'], c['config']), {})[c['column']] = c['value']
RB_PAIR = {}
for p in pairs:
    a, b = p['pair'].split(' / ')
    if abs(float(p['diff'])) <= 0.05:
        RB_PAIR[a] = RB_PAIR[b] = float(p['implied_RB'])


def stage_rows(t):
    ws = openpyxl.load_workbook(os.path.join(STAGE, f'DesignTable_{t}.xlsx'), data_only=True).worksheets[0]
    rr = list(ws.iter_rows(values_only=True)); h = rr[1]
    return [dict(zip(h, x)) for x in rr[2:] if x[0] and (t == '6701' or str(x[0]).endswith('-Tube'))]


def wall(od):
    return 0.065 if od <= 0.5 else (0.083 if od <= 1.0 else 0.109)


def build(variant):
    out = os.path.join(T, f'tables_{variant}')
    os.makedirs(out, exist_ok=True)
    for root, tabs in GROUPS.items():
        hdr = HDR_SW if root == '6701T' else HDR
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
        ws.append([f'WAVE-3 BENT-TUBE TEST TABLE {root} ({variant}) - stage -Tube rows + wave3 cells.csv | 2026-10-18 | '
                   + ('TEST PLACEHOLDERS for TUBEOD / RB / HEXL / D - NOT PART DATA' if variant == 'whatif' else 'no placeholders')])
        ws.append(hdr)
        for t in tabs:
            for r in stage_rows(t):
                pn = r['$PRP@PartNo']
                v = {k: C.get((t, pn), {}).get(k) for k in hdr if k}
                v = {k: (float(x) if x not in (None, '') and k != '$PRP@FlangeCode' else x) for k, x in v.items()}
                ds = r['$PRP@DataStatus']
                src = 'wave3 cells.csv (Brennan catalog / Air-Way s115 / end library)'
                if variant == 'probe':
                    ds = 'ASSUMED'
                if variant == 'whatif':
                    ds = 'TEST'
                    a = int(pn.replace('-L-', '-').split('-')[1])
                    tod = a / 16
                    ph = []
                    if v.get('TUBEOD@TUBE') is None:
                        v['TUBEOD@TUBE'] = tod; ph.append(f'TUBEOD {tod} = JIC dash')
                    if v.get('HEXL@HEX') is None:
                        v['HEXL@HEX'] = 0.30; ph.append('HEXL 0.30')
                    if v.get('RB@PATH') is None:
                        k = pn.replace('-Tube', '')
                        rb = RB_PAIR.get(k, round(1.5 * tod, 3))
                        v['RB@PATH'] = rb; ph.append(f'RB {rb} ' + ('(implied by the 45/90 pair)' if k in RB_PAIR else '(1.5 x tube)'))
                    if v.get('D@BORE') is None:
                        d = round(tod - 2 * wall(tod), 3)
                        v['D@BORE'] = d; ph.append(f'D {d} (tube - 2 x TEST wall)')
                    src = 'TEST ROW (wave 3 what-if, NOT part data) | ' + src + ' | TEST placeholders: ' + ', '.join(ph)
                row = [pn] + [v.get(k) if k not in ('$PRP@PartNo', '$PRP@DataStatus', '$PRP@AxialSource', '$COMMENT', '$PRP@ScopeAction', '$PRP@OMiSKU', '$PRP@TubeOD')
                              else {'$PRP@PartNo': pn, '$PRP@DataStatus': ds, '$PRP@AxialSource': src, '$COMMENT': r.get('$COMMENT'),
                                    '$PRP@ScopeAction': r.get('$PRP@ScopeAction'), '$PRP@OMiSKU': r.get('$PRP@OMiSKU'),
                                    '$PRP@TubeOD': (f'{v["TUBEOD@TUBE"]:g}' if isinstance(v.get('TUBEOD@TUBE'), float) else None)}[k]
                              for k in hdr[1:]]
                ws.append(row)
        wb.save(os.path.join(out, f'DesignTable_{root}.xlsx'))
    return out


for v in ('real', 'probe', 'whatif'):
    print(build(v))

# generator copies
G = os.path.join(T, 'gen'); os.makedirs(G, exist_ok=True)
shutil.copy2(GB, os.path.join(G, 'gen_bent.py'))
s = open(GB, encoding='utf-8').read()
h62 = {}
for t in CAT['tables']:
    if t['series'] == '3300-62':
        for r in t['rows']:
            fl = int(r['pn'].split('-')[3]); vv = r['values']
            e = dict(O=float(vv['A']), TH=float(vv['D']), GO=round(float(vv['H']) / 2, 4))
            assert h62.setdefault(fl, e) == e, (fl, e, h62[fl])
add = ', '.join(f"{fl}: dict(O={e['O']}, TH={e['TH']}, GI={round(e['GO'] - 0.155, 4)}, GO={e['GO']}, GD=0.110)" for fl, e in sorted(h62.items()))
old = "FLG62_MEAS = {8: dict(O=1.250, TH=0.305, GI=0.345, GO=0.500, GD=0.110)}"
assert old in s
s = s.replace(old, old[:-1] + ', ' + add + '}  # WHAT-IF (wave 3): -12..-32 from Brennan catalog 3300-62 A / D / H, groove .155 x .110 (3300 drawing)')
open(os.path.join(G, 'gen_bent_c62cat.py'), 'w', encoding='utf-8').write(s)
print('FLG62 what-if heads:', add)
