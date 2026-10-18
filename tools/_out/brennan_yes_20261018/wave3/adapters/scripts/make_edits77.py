"""make_edits77.py (wave 3 adapters, 2026-10-18): convert tools/_out/adapters77_20261016/CHANGES.csv (the 37 wrongly-DROP Air-Way rows,
staged 2026-10-16 against the then-live tables) into edits against TODAY's stage tables (brennan_yes_20261018/stage/04_DESIGN_TABLES,
batch 16 included). Each line: table,config,column,old,new,why with old = the stage's CURRENT value. A line whose 2026-10-16 'old' no
longer matches the stage is reported in edits_77_mismatch.csv (not silently rebased). '$COMMENT (append)' stays an append
(old = current comment, new = the text to append).
    python -I make_edits77.py
"""
import csv, os, collections
import openpyxl

ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
SRC = os.path.join(ROOT, r'tools\_out\adapters77_20261016\CHANGES.csv')
STAGE = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018\stage\04_DESIGN_TABLES')
OUT = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018\wave3\adapters')


def same(a, b):
    if a in (None, '') and b in (None, ''):
        return True
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


def fmt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return repr(v)
    return str(v)


rows = list(csv.DictReader(open(SRC, encoding='utf-8-sig')))
books = {}
out, mism = [], []
for c in rows:
    tab = c['review table']
    t = tab.replace('DesignTable_', '').replace('.xlsx', '')
    if t not in books:
        ws = openpyxl.load_workbook(os.path.join(STAGE, tab)).worksheets[0]
        hdr = {str(x.value): x.column for x in ws[2] if x.value is not None}
        rowof = collections.defaultdict(list)
        for r in range(3, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if v:
                rowof[str(v).strip()].append(r)
        books[t] = (ws, hdr, rowof)
    ws, hdr, rowof = books[t]
    cfg = c['row (config name)']
    col = c['column'].replace(' (append)', '')
    rr = rowof.get(cfg, [])
    if len(rr) != 1 or col not in hdr:
        mism.append(dict(table=t, config=cfg, column=c['column'], old_20261016=c['old'], stage_now='ROW/COLUMN NOT FOUND' if not rr else f'{len(rr)} rows', new=c['new'], why=c['reason']))
        continue
    cur = ws.cell(rr[0], hdr[col]).value
    if c['column'].endswith('(append)'):
        out.append(dict(table=t, config=cfg, column=c['column'], old=fmt(cur), new=c['new'], why=c['reason']))
        continue
    if not same(cur, c['old']):
        mism.append(dict(table=t, config=cfg, column=col, old_20261016=c['old'], stage_now=fmt(cur), new=c['new'], why=c['reason']))
        continue
    out.append(dict(table=t, config=cfg, column=col, old=fmt(cur) if cur not in (None,) else '', new=c['new'], why=c['reason']))

# ---- rebase on a user decision taken after 2026-10-16: A11 Q-BOREWIDE (2026-10-18, 00_DOCS/ANSWERS_2026-10-18.md): "keep / cap the bore at the
# port drill and report to OMi". The two 6405 D3 edits put OMi's 0.844 (= Brennan B DRILL) into a bore that runs into a -06 / -08 female pipe
# port of tap drill 0.57 / 0.70; the batch-16 generator blocks them ("bore 0.844 > pocket 0.57", test/compare.txt first run). Capped at the port
# drill = the stage value, so the edit line is dropped and the row's comment says why (the stage itself did the same on 6404-06-08 D1).
A11 = {('6405', '6405-16-06-O', 'D3@BORE'): ('0.57', '-06'), ('6405', '6405-16-08-O', 'D3@BORE'): ('0.7', '-08')}
keep = []
for e in out:
    k = (e['table'], e['config'], e['column'])
    if k in A11:
        assert same(e['old'], A11[k][0]), (k, e['old'])
        print('  A11 cap: edit dropped', k, e['old'], '->', e['new'])
        continue
    keep.append(e)
for e in keep:
    for (t, c, col), (cap, dash) in A11.items():
        if e['table'] == t and e['config'] == c and e['column'] == '$COMMENT (append)':
            e['new'] += (f' | 2026-10-18 rebase (A11 Q-BOREWIDE yes: cap at the port drill, report to OMi): {col} stays {cap} = the {dash} female pipe tap drill; '
                         f'OMi 0.844 (= Brennan B DRILL) is wider than the port it runs into - the OMi {col} value above is NOT applied')
            e['why'] += ' | A11 cap note added 2026-10-18'
out = keep

with open(os.path.join(OUT, 'edits_77.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=['table', 'config', 'column', 'old', 'new', 'why'])
    w.writeheader(); w.writerows(out)
with open(os.path.join(OUT, 'edits_77_mismatch.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=['table', 'config', 'column', 'old_20261016', 'stage_now', 'new', 'why'])
    w.writeheader(); w.writerows(mism)
print(len(rows), 'source lines;', len(out), 'edits;', len(mism), 'mismatches')
for m in mism:
    print('  MISMATCH', m['table'], m['config'], m['column'], 'was', m['old_20261016'], 'now', m['stage_now'])
