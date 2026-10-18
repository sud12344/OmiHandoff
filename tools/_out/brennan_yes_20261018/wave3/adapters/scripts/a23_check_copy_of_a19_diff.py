# COPY of tools/_out/a19_omi_diffs_20261018/a19_diff.py with OUT / STAGE pointed at the wave3 adapters test (applied tables)
"""A19 / A20 / A21 / A23 (user answers 2026-10-18): READ-ONLY comparison of OMi's dimension file against the CURRENT tables.
Current tables = tools/_out/brennan_yes_20261018/stage/04_DESIGN_TABLES (latest staged state).
Built configurations = names present in the live 04_DESIGN_TABLES/DesignTable_*_SW2.xlsx.
Column mappings = 00_DOCS/OMi_Dimension_Audit_2026-09-25 (tools/_out/omi_audit/mapping_final.json, field 'certainty').
Rounding rule (user 2026-09-25, QA #158): equal after rounding OMi half-up to Air-Way's printed decimals, or OMi cut to them,
or the table holds the exact fraction and OMi rounds / cuts it to its own decimals (OMi with >= 2 decimals).
Writes only into this folder."""
import openpyxl, glob, os, re, json, csv, hashlib, math, collections
from fractions import Fraction
from decimal import Decimal

R = r"C:\Users\sudha\OneDrive\Desktop\OMi_Handoff"
OUT = os.path.join(R, r"tools\_out\brennan_yes_20261018\wave3\adapters\a23_check")
STAGE = os.path.join(R, r"tools\_out\brennan_yes_20261018\wave3\adapters\test\root_applied\04_DESIGN_TABLES")
LIVE = os.path.join(R, "04_DESIGN_TABLES")
OMIF = os.path.join(R, r"07_RAW_EXTRACTS\omi_dimensions\WebsitePartsInfo2026.xlsx")
ORIG1404 = os.path.join(R, r"tools\_out\omi_apply_20260925\originals_before_apply\DesignTable_1404.xlsx")

COL = {"OverallLength": 12, "BodyDiameter": 13, "M1": 14, "M2": 15, "L1": 16, "L2": 17, "HexFlat": 18, "S1ID": 19,
       "NutHex1": 22, "S2ID": 24, "NutHex2": 25, "S3ID": 29, "S4ID": 33}
LET = {"OverallLength": "L", "BodyDiameter": "M", "M1": "N", "M2": "O", "L1": "P", "L2": "Q", "HexFlat": "R", "S1ID": "S",
       "NutHex1": "V", "S2ID": "X", "NutHex2": "Y", "S3ID": "AC", "S4ID": "AG"}
HDR = {"OverallLength": "Overall Length", "BodyDiameter": "Body Diameter", "M1": "M 1 Size", "M2": "M 2 Size", "L1": "L 1 Size",
       "L2": "L 2 Size", "HexFlat": "Hex Flat", "S1ID": "Side 1 ID", "NutHex1": "Nut Hex 1", "S2ID": "Side 2 ID",
       "NutHex2": "Nut Hex 2", "S3ID": "Side 3 ID", "S4ID": "Side 4 ID"}

# leg / length readings the audit called uncertain but that the Brennan TFA catalog's own drawing letters settled in batch 15
# (tools/_out/batch15_20261002/work/*/NOTES.md; the user approved batch 15 on that basis 2026-10-02)
SETTLED = {}
for i in ("M071", "M072", "M076", "M077", "M085", "M086"):
    SETTLED[i] = "Q36-5 reading; settled by Brennan TFA catalog letters in batch 15 (Q-B-ELBOW, TEE-2501 / TEE-OTHER NOTES)"
for i in ("M099", "M100"):
    SETTLED[i] = "Q36-6 reading; settled by Brennan TFA catalog letters in batch 15 (Q-B-5500, TEE-OTHER NOTES)"
for i in ("M166", "M167", "M168"):
    SETTLED[i] = "Q36-7 reading; settled by Brennan TFA catalog letters in batch 15 (Q-B-TEE, TEE-OTHER NOTES)"
for i in ("M113", "M114", "M115", "M134", "M135", "M136"):
    SETTLED[i] = "Q36-8 reading; settled by Brennan TFA catalog letters in batch 15 (Q-B-SWELB, SWIVEL-ELBOWS NOTES)"
for i in ("M127", "M128", "M129"):
    SETTLED[i] = "Q36-9 reading; settled by Brennan TFA catalog letters in batch 15 (Q-B-6501, SWIVEL-ELBOWS NOTES)"
for i in ("M031", "M032", "M038", "M039"):
    SETTLED[i] = "Q36-10 reading; settled by Brennan TFA catalog letters in batch 15 (Q-B-SWSTR, SWIVEL-STRAIGHTS NOTES)"
SETTLED["M026"] = "Q36-14 reading; settled by Brennan TFA catalog letters in batch 15 (Q-B-OTHER-6504, SWIVEL-STRAIGHTS NOTES)"
OPEN = {"M081": "Q36-5", "M082": "Q36-5", "M109": "Q36-5", "M110": "Q36-5", "M141": "Q36-9", "M142": "Q36-9", "M143": "Q36-9",
        "M173": "Q36-7", "M174": "Q36-7", "M175": "Q36-7", "M176": "Q36-7 (hex vs boss; challenger downgraded)",
        "M179": "Q36-7", "M180": "Q36-7", "M181": "Q36-7", "M186": "Q36-7", "M187": "Q36-7", "M188": "Q36-7",
        "M020": "Q36-10", "M021": "Q36-10", "M051": "Q36-13", "M013": "Q36-14", "M016": "Q36-14", "M069": "Q36-14"}
SKIP_MAP = {"M050": "alternative of M049 (user A20: Overall Length = W)", "M070": "alternative of M069 (Q36-14: D1)"}
NOTE_MAP = {"M026": "OMi gives no L14: moving L13 alone moves the swivel nose against the nut face (batch 15 kept 6504-06-08 at "
                    "Air-Way for this reason, SWIVEL-STRAIGHTS NOTES 'Deliberately not changed')"}


def dec(s):
    if s is None: return None
    if isinstance(s, bool): return None
    if isinstance(s, int): return Fraction(s)
    if isinstance(s, float): s = repr(s)
    s = str(s).strip()
    if s == '': return None
    try: return Fraction(Decimal(s))
    except Exception: return None


def sval(v):
    if v is None: return None
    if isinstance(v, float): return repr(v)
    return str(v).strip()


def ndec(s):
    s = sval(s)
    if s is None or 'e' in s.lower(): return None
    return len(s.split('.')[1]) if '.' in s else 0


def rhu(fr, nd):
    q = Fraction(1, 10 ** nd); n = fr / q
    return (math.floor(n + Fraction(1, 2)) if n >= 0 else -math.floor(-n + Fraction(1, 2))) * q


def trunc(fr, nd):
    q = Fraction(1, 10 ** nd); n = fr / q
    return (math.floor(n) if n >= 0 else -math.floor(-n)) * q


def fmt(fr):
    if fr is None: return ''
    d = Decimal(fr.numerator) / Decimal(fr.denominator)
    return format(d.normalize(), 'f')


def compare(o_s, r_v, printed=None):
    """-> 'equal' | 'rounding' | 'real' | None"""
    o = dec(o_s); r = dec(r_v)
    if o is None or r is None: return None
    if o == r: return 'equal'
    pd = ndec(printed) if (printed is not None and dec(printed) == r) else max(2, ndec(r_v) or 0)
    if rhu(o, pd) == r or trunc(o, pd) == r: return 'rounding'
    od, rd = ndec(o_s), ndec(r_v)
    if od is not None and rd is not None and 2 <= od < rd and (rhu(r, od) == o or trunc(r, od) == o): return 'rounding'
    return 'real'


# ---------- OMi file
sha = hashlib.sha256(open(OMIF, 'rb').read()).hexdigest()
ws = openpyxl.load_workbook(OMIF, read_only=True, data_only=True)["Items"]
OMI = collections.defaultdict(list)
for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
    if i < 4 or not r or r[0] is None: continue
    pn = str(r[0]).strip()
    OMI[pn].append((i, {k: sval(r[c - 1]) if c - 1 < len(r) else None for k, c in COL.items()}))
dup = {k: v for k, v in OMI.items() if len(v) > 1 and any(any(x[1].values()) for x in v)}

# ---------- mappings, Air-Way printed values (audit), Task A held cells
MAP = json.load(open(os.path.join(R, r"tools\_out\omi_audit\mapping_final.json"), encoding='utf-8'))
mapidx = collections.defaultdict(list)
for m in MAP: mapidx[m["review_table"]].append(m)
DIMS = json.load(open(os.path.join(R, r"tools\_out\omi_audit\dims.json")))
PRINTED = {(d["table"], d["config"], d.get("review_col")): (d.get("printed"), d.get("printed_src")) for d in DIMS["rows"] if d.get("review_col")}
HELD = {(h["table"], h["config"], h["review_col"]): h["why"] for h in json.load(open(os.path.join(R, r"tools\_out\omi_apply_20260925\held.json"), encoding='utf-8'))}


def applies(m, cfg):
    f = m.get("row_filter")
    if not f: return True
    if "regex" in f: return re.search(f["regex"], cfg or '') is not None
    return cfg in f["configs"]


def read_table(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[1]); letters = {h: openpyxl.utils.get_column_letter(j + 1) for j, h in enumerate(hdr) if h}
    out = []
    for xl, r in enumerate(rows[2:], start=3):
        if not r or r[0] is None: continue
        d = {hdr[j]: r[j] for j in range(min(len(hdr), len(r))) if hdr[j]}
        d["_name"] = str(r[0]).strip(); d["_xl"] = xl
        out.append(d)
    return out, letters


# ---------- built configurations
BUILT = {}
for p in glob.glob(os.path.join(LIVE, "DesignTable_*_SW2.xlsx")):
    ws = openpyxl.load_workbook(p, read_only=True, data_only=True).worksheets[0]
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1 or not r or r[0] is None: continue
        BUILT.setdefault(str(r[0]).strip(), os.path.basename(p))


def built_as(row):
    for k in (row["_name"], row.get("$PRP@OMiSKU"), row.get("$PRP@PartNo")):
        if k and str(k).strip() in BUILT: return BUILT[str(k).strip()]
    return None


def source_of(row, col, val, printed):
    ax = str(row.get("$PRP@AxialSource") or '')
    base = col.split('@')[0]
    if re.search(r'(?<![A-Za-z0-9_])' + re.escape(col) + r': OMi', ax) or re.search(r'(?<![A-Za-z0-9_])' + re.escape(base) + r'(@\w+)?: OMi', ax):
        return 'OMi'
    if printed is not None and dec(printed) is not None and dec(printed) == dec(val):
        return 'Air-Way'
    m = re.search(r'(?<![A-Za-z0-9_])' + re.escape(base) + r'(@\w+)?[^;]{0,80}?(Brennan|MS5\d+|caliper|catalog|rule|OMi)', ax, re.I)
    if m: return 'other: ' + m.group(0)[:100]
    if printed is not None and dec(printed) is not None and compare(sval(val), printed) in ('equal', 'rounding'):
        return 'Air-Way (exact value of the printed ' + printed + ')'
    if ax.lower().startswith('air-way'):
        return 'Air-Way (row AxialSource: ' + ax.split(';')[0][:60] + ')'
    m = re.search(r'[^;]{0,30}Air-Way[^;]{0,50}', ax)
    if m:
        return 'Air-Way (row AxialSource: ' + m.group(0).strip()[:80] + ')'
    return 'unclear (AxialSource: ' + ax[:80] + ')'


tables = sorted({m["review_table"] for m in MAP})
recs = []       # every comparison on BUILD rows
for key in tables:
    path = os.path.join(STAGE, f"DesignTable_{key}.xlsx")
    if not os.path.exists(path): path = os.path.join(LIVE, f"DesignTable_{key}.xlsx")
    if not os.path.exists(path): continue
    rows, letters = read_table(path)
    for row in rows:
        if str(row.get("$PRP@ScopeAction") or '').strip().upper() != 'BUILD': continue
        sku = str(row.get("$PRP@OMiSKU") or '').strip()
        if sku not in OMI: continue
        hits = [x for x in OMI[sku] if any(x[1].values())] or OMI[sku]
        orow, ov = hits[0]
        b = built_as(row)
        for m in mapidx[key]:
            if not applies(m, row["_name"]): continue
            rc = m.get("review_column")
            if not rc or rc not in letters: continue
            o_s = ov.get(m["omi_column"]); r_v = row.get(rc)
            if o_s in (None, ''): continue
            printed, psrc = PRINTED.get((key, row["_name"], rc), (None, None))
            cls = compare(o_s, r_v, printed)
            recs.append(dict(table=key, config=row["_name"], xl=row["_xl"], column=rc, cell=f"{letters[rc]}{row['_xl']}",
                             current=sval(r_v), omi=o_s, omi_cell=f"Items!{LET[m['omi_column']]}{orow}", omi_hdr=HDR[m['omi_column']],
                             cls=cls if cls else ('only OMi' if dec(o_s) is not None else 'OMi not a number'),
                             delta=fmt(dec(o_s) - dec(r_v)) if cls else '', map_id=m["id"], certainty=m["certainty"],
                             built=b, status=row.get("$PRP@DataStatus"), source=source_of(row, rc, r_v, printed),
                             printed=printed, dimsource=row.get("$PRP@DimSource")))

# ---------- classify real differences
A19, UNC, EXCL, A20 = [], [], [], []
for x in recs:
    if x["cls"] != 'real': continue
    k = (x["table"], x["config"], x["column"])
    if x["map_id"] in SKIP_MAP: continue
    if x["table"] == '1404' and x["column"] == 'W@HEX':
        A20.append(x); continue
    if x["table"] == '1404' and x["column"] in ('D@BORE', 'D2@BORE'):
        x["why"] = 'A21: 1404 bores go back to Air-Way (A21_1404_bores.csv)'; EXCL.append(x); continue
    if not x["built"]:
        continue
    if x["table"] == '2404_FINAL':
        x["why"] = "2404 held: 'keep 2404 exactly as it is' (user 2026-10-02) until OMi clarifies the bores"; EXCL.append(x); continue
    if x["table"] == '6565' and x["column"] in ('L31@DRIVE', 'L32@DRIVE'):
        x["why"] = 'A22: 6565 lengths kept as built until OMi or Brennan 6565-08-08 confirms'; EXCL.append(x); continue
    if k in HELD:
        x["why"] = ('Task A held; A11 = keep / cap the bore at the port drill: ' if 'BORE' in x["column"] else 'Task A held: ') + HELD[k]; EXCL.append(x); continue
    if x["omi_hdr"] == 'Overall Length' and x["omi"] in ('1.000', '1', '1.0') and x["table"] in ('6410', '6405'):
        x["why"] = "OMi Overall Length exactly 1.000 = OMi's placeholder pattern (audit 3.1 / question 8; Task A skipped these)"; EXCL.append(x); continue
    if x["source"].startswith('other') or x["source"].startswith('unclear'):
        x["note"] = 'current value is not Air-Way: ' + x["source"]
    if x["certainty"] == 'certain':
        A19.append(x)
    else:
        x["reading"] = SETTLED.get(x["map_id"]) or ('open: ' + OPEN.get(x["map_id"], 'audit uncertain'))
        x["note"] = NOTE_MAP.get(x["map_id"], x.get("note", ''))
        UNC.append(x)


def src_short(x):
    s = x["source"]
    return 'Air-Way' if s.startswith('Air-Way') else s


def write(name, lst, extra=()):
    with open(os.path.join(OUT, name), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(["table", "config", "column", "current_value", "current_source", "omi_value", "omi_cell", "delta", "mapping_id"] + [e[0] for e in extra])
        for x in sorted(lst, key=lambda x: (x["table"], x["config"], x["column"])):
            w.writerow([x["table"], x["config"], x["column"], x["current"], src_short(x), x["omi"], x["omi_cell"], x["delta"], x["map_id"]] + [x.get(e[1], '') for e in extra])


write("A19_changes.csv", A19, [("table_cell", "cell"), ("omi_header", "omi_hdr"), ("data_status", "status"), ("note", "note")])
write("A19_uncertain_readings.csv", UNC, [("table_cell", "cell"), ("omi_header", "omi_hdr"), ("reading_status", "reading"), ("note", "note")])
write("A19_excluded.csv", EXCL, [("table_cell", "cell"), ("built", "built"), ("reason", "why")])
for x in A20: x["b"] = 'built' if x["built"] else 'not built'
write("A20_1404_nutface.csv", A20, [("table_cell", "cell"), ("built", "b")])

# notes: possible leg swaps, known patterns
byrow = collections.defaultdict(list)
for x in UNC: byrow[(x["table"], x["config"])].append(x)
for (t, c), lst in byrow.items():
    legs = [x for x in recs if x["table"] == t and x["config"] == c and x["map_id"] in [y["id"] for y in MAP if y["review_table"] == t and y["omi_column"] in ("M1", "M2", "L1")]]
    for a in lst:
        for b in legs:
            if b is a or b["column"] == a["column"]: continue
            if dec(a["omi"]) is not None and dec(b["current"]) is not None and abs(dec(a["omi"]) - dec(b["current"])) <= Fraction(25, 1000):
                a["note"] = ((a.get("note") or '') + f"; possible leg swap: OMi {a['omi_hdr']} {a['omi']} ~ table {b['column']} {b['current']} - check before applying").lstrip('; ')
for x in UNC:
    if x["table"] == '0318':
        x["note"] = "OMi = MS51531D nominal; Air-Way prints 0.01 shorter on every 0318 row (audit 3.6)"
    if x["config"] in ('2503-04-08-FG', '2503-10-12-FG'):
        x["note"] = ((x.get("note") or '') + "; audit 3.5: OMi M 1 / L 1 look swapped against Air-Way M2 / N2 on this row").lstrip('; ')
    if x["table"] == '1404' and x["omi"] == '1':
        x["note"] = "OMi writes the integer text '1' (read as 1.000)"
for x in A19:
    if x["table"] == '6400':
        x["note"] = "6400 is the hand-made master; A6 = rebuild 6400 with the generator (fold 6400-L in)"
write("A19_uncertain_readings.csv", UNC, [("table_cell", "cell"), ("omi_header", "omi_hdr"), ("reading_status", "reading"), ("note", "note")])
write("A19_changes.csv", A19, [("table_cell", "cell"), ("omi_header", "omi_hdr"), ("data_status", "status"), ("note", "note")])

# ---------- A21: 1404 bores back to Air-Way
cur, letters = read_table(os.path.join(STAGE, "DesignTable_1404.xlsx"))
orig, _ = read_table(ORIG1404)
od = {r["_name"]: r for r in orig}
A21 = []
for r in cur:
    o = od.get(r["_name"])
    if not o: continue
    for c in ('D@BORE', 'D2@BORE'):
        if dec(r.get(c)) != dec(o.get(c)):
            printed, psrc = PRINTED.get(('1404', r["_name"], c), (None, None))
            sku = str(r.get("$PRP@OMiSKU") or '').strip()
            orow = OMI[sku][0][0] if sku in OMI else ''
            ocol = 'X' if c == 'D@BORE' else 'S'
            A21.append(dict(table='1404', config=r["_name"], column=c, current=sval(r.get(c)), current_source='OMi (Task A, batch 14)',
                            airway=sval(o.get(c)), printed=printed or '', printed_src=psrc or '', omi_cell=f"Items!{ocol}{orow}",
                            delta=fmt(dec(o.get(c)) - dec(r.get(c))), built='built' if built_as(r) else 'not built',
                            cell=f"{letters[c]}{r['_xl']}"))
with open(os.path.join(OUT, "A21_1404_bores.csv"), 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(["table", "config", "column", "current_value", "current_source", "airway_value", "omi_cell", "delta", "mapping_id",
                "table_cell", "airway_printed_check", "built"])
    for x in A21:
        w.writerow([x["table"], x["config"], x["column"], x["current"], x["current_source"], x["airway"], x["omi_cell"], x["delta"],
                    'M055' if x["column"] == 'D@BORE' else 'M054', x["cell"],
                    (x["printed"] + ' (' + x["printed_src"] + ')') if x["printed"] else 'no printed value in the audit', x["built"]])

# ---------- A23: rows where every OMi-mapped cell agrees with the table's Air-Way value
rowsets = collections.defaultdict(list)
for x in recs:
    if x["cls"] in (None, 'only OMi', 'OMi not a number'): continue
    if x["map_id"] in SKIP_MAP: continue
    if x["certainty"] != 'certain' and x["map_id"] not in SETTLED and not (x["table"] == '1404' and x["column"] == 'W@HEX'):
        continue
    rowsets[(x["table"], x["config"])].append(x)
A23 = []
for (t, c), lst in sorted(rowsets.items()):
    agree, dis, notaw = [], [], []
    for x in lst:
        tag = f"{x['column']}={x['current']}~{x['omi']}"
        if x["source"] == 'OMi':
            dis.append(tag + ' (table holds OMi value; Air-Way differed)')
        elif not x["source"].startswith('Air-Way'):
            notaw.append(tag + ' (' + x["source"][:60] + ')')
        elif x["cls"] in ('equal', 'rounding'):
            agree.append(tag)
        else:
            dis.append(tag)
    other = [x for x in recs if x["table"] == t and x["config"] == c and x["cls"] == 'real' and x not in lst and x["map_id"] not in SKIP_MAP]
    note = '; '.join(f"{x['column']} {x['current']} vs OMi {x['omi']} differs on an unsettled reading ({x['map_id']})" for x in other)
    if t == '2404_FINAL': note = ('2404 held (keep exactly as it is); ' + note).rstrip('; ')
    if lst[0]["status"] == 'MISSING': note = ('row is MISSING for other cells - agreement alone does not make it VERIFIED; ' + note).rstrip('; ')
    A23.append(dict(note=note, table=t, config=c, n=len(lst), all_agree='yes' if (agree and not dis and not notaw) else 'no',
                    agree='; '.join(agree), dis='; '.join(dis), notaw='; '.join(notaw), status=lst[0]["status"],
                    built='built' if lst[0]["built"] else 'staged (not built)',
                    readings='certain only' if all(x["certainty"] == 'certain' for x in lst) else 'incl. batch-15-settled / A20 readings'))
with open(os.path.join(OUT, "A23_verified_candidates.csv"), 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(["table", "config", "cells_compared", "all_agree", "current_DataStatus", "built", "cells_agree", "cells_differ",
                "cells_not_AirWay", "mappings_used", "note"])
    for x in A23:
        w.writerow([x["table"], x["config"], x["n"], x["all_agree"], x["status"], x["built"], x["agree"], x["dis"], x["notaw"], x["readings"], x["note"]])

json.dump(dict(omi_sha256=sha, omi_dup_pn=list(dup)[:50], n_recs=len(recs),
               cls=collections.Counter((r["certainty"], r["cls"], bool(r["built"])) for r in recs).most_common()),
          open(os.path.join(OUT, "run_info.json"), 'w'), indent=1, default=str)
print('sha', sha[:12], 'recs', len(recs), 'dup', len(dup))
print('A19 certain', len(A19), 'cfgs', len({(x['table'], x['config']) for x in A19}))
print('uncertain', len(UNC), 'cfgs', len({(x['table'], x['config']) for x in UNC}))
print('excluded', len(EXCL), 'A20', len(A20), 'A21', len(A21), 'A23 rows', len(A23), 'yes', sum(1 for x in A23 if x['all_agree'] == 'yes'))
