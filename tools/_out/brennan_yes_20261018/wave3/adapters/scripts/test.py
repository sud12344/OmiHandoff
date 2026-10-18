"""test.py (wave 3 adapters, 2026-10-18): copy today's STAGE tables twice -
  base    = stage copy, nothing applied (control)
  applied = stage copy + edits_77.csv + edits_5406P_alias.csv (old values asserted) + newrows.csv (rows appended at the end of their table,
            missing columns added at the end of header row 2)
then regenerate EVERY root of the batch-16 generators (puckfix_20261009/tools_new, one process per root, python -I) on both and compare
(SW2 cell by cell for every pre-existing configuration, new configurations listed, macro lines). Outputs in ../test.
    python -I -B test.py [apply|regen|compare|checks|all]
"""
import os, sys, csv, json, shutil, subprocess, re, collections
import openpyxl

ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
BY = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018')
OUT = os.path.join(BY, r'wave3\adapters')
TEST = os.path.join(OUT, 'test')
STAGE_DT = os.path.join(BY, r'stage\04_DESIGN_TABLES')
TOOLS = os.path.join(ROOT, r'tools\_out\puckfix_20261009\tools_new')
SCEN = {'base': os.path.join(TEST, 'root_base'), 'applied': os.path.join(TEST, 'root_applied')}


def num_or_str(s):
    s = str(s)
    return float(s) if re.fullmatch(r'-?\d*\.\d+|-?\d+', s.strip()) else s


def same(a, b):
    if a in (None, '') and b in (None, ''):
        return True
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


class Book:
    def __init__(self, dt, t):
        self.p = os.path.join(dt, f'DesignTable_{t}.xlsx')
        self.wb = openpyxl.load_workbook(self.p); self.ws = self.wb.worksheets[0]
        self.hdr = {c.value: c.column for c in self.ws[2] if c.value is not None}

    def row(self, cfg):
        rr = [c.row for c in self.ws['A'] if c.value is not None and str(c.value).strip() == cfg]
        assert len(rr) == 1, (self.p, cfg, rr)
        return rr[0]

    def col(self, h):
        if h not in self.hdr:
            c = max(self.ws.max_column, max(self.hdr.values())) + 1
            self.ws.cell(2, c, h); self.hdr[h] = c
        return self.hdr[h]


def apply(dt):
    log = []
    books = {}
    bk = lambda t: books.setdefault(t, Book(dt, t))
    for f in ('edits_77.csv', 'edits_5406P_alias.csv'):
        for e in csv.DictReader(open(os.path.join(OUT, f), encoding='utf-8')):
            b = bk(e['table']); r = b.row(e['config'])
            colname = e['column'].replace(' (append)', '')
            cell = b.ws.cell(r, b.hdr[colname])
            if e['column'].endswith('(append)'):
                assert same(cell.value, e['old']), (f, e, cell.value)
                cell.value = (str(cell.value) + ' | ' if cell.value else '') + e['new']
            else:
                assert same(cell.value, e['old']), (f, e, cell.value)
                cell.value = num_or_str(e['new']) if not colname.startswith('$') else e['new']
            log.append(f"{f}: {e['table']} {e['config']} {e['column']} {e['old'][:30]!r} -> {e['new'][:50]!r}")
    new = collections.OrderedDict()
    for x in csv.DictReader(open(os.path.join(OUT, 'newrows.csv'), encoding='utf-8')):
        new.setdefault((x['table'], x['config']), []).append(x)
    for (t, cfg), cells in new.items():
        b = bk(t)
        assert not [c for c in b.ws['A'] if c.value is not None and str(c.value).strip() == cfg], (t, cfg, 'exists')
        r = max(c.row for c in b.ws['A'] if c.value is not None) + 1
        b.ws.cell(r, 1, cfg)
        for x in cells:
            if x['value'] == '':
                continue
            c = b.col(x['column'])
            b.ws.cell(r, c, num_or_str(x['value']) if not x['column'].startswith('$') else x['value'])
        log.append(f'newrows.csv: {t} {cfg} appended at row {r} ({sum(1 for x in cells if x["value"] != "")} cells)')
    for b in books.values():
        b.wb.save(b.p)
    open(os.path.join(TEST, 'apply_log.txt'), 'w', encoding='utf-8').write('\n'.join(log) + '\n')
    print(len(log), 'apply log lines')


def regen(scen):
    out = os.path.join(TEST, f'gen_{scen}')
    assert not os.path.exists(out), out
    os.makedirs(os.path.join(out, 'logs'))
    env = dict(os.environ, OMI_ROOT=SCEN[scen], GS_OUT=os.path.join(out, 'straight'), GT_OUT=os.path.join(out, 'tee'), PYTHONDONTWRITEBYTECODE='1')
    code = ("import sys,os,json; t=sys.argv[1]; sys.path[:0]=[os.path.join(t,'generators'),t]; import gen_straight2 as G, gen_tee as T; "
            "print(json.dumps(dict(straight=list(G.SPEC), tee=list(T.SPEC))))")
    k = subprocess.run([sys.executable, '-I', '-B', '-c', code, TOOLS], env=dict(env, GS_OUT=os.path.join(out, '_keys'), GT_OUT=os.path.join(out, '_keys')),
                       capture_output=True, text=True, encoding='utf-8', errors='replace')
    if k.returncode:
        print(k.stderr); sys.exit(1)
    roots = json.loads(k.stdout.strip().splitlines()[-1])
    json.dump(roots, open(os.path.join(out, 'roots.json'), 'w'))
    status = {}
    for sub, gen in (('straight', 'gen_straight2.py'), ('tee', 'gen_tee.py')):
        os.makedirs(os.path.join(out, sub), exist_ok=True)
        sj = os.path.join(out, sub, 'gen_straight2_summary.json' if sub == 'straight' else 'gen_tee_summary.json')
        for r in roots[sub]:
            if os.path.exists(sj): os.remove(sj)
            p = subprocess.run([sys.executable, '-I', '-B', os.path.join(TOOLS, 'generators', gen), r], env=env, capture_output=True, text=True,
                               encoding='utf-8', errors='replace')
            open(os.path.join(out, 'logs', f'{sub}_{r}.txt'), 'w', encoding='utf-8').write(p.stdout + ('\n--- STDERR ---\n' + p.stderr if p.stderr else ''))
            if os.path.exists(sj): os.replace(sj, os.path.join(out, sub, f'summary_{r}.json'))
            last = (p.stderr.strip().splitlines() or [''])[-1][:300]
            status[f'{sub}:{r}'] = dict(rc=p.returncode, err=last if p.returncode else '', out=(p.stdout.strip().splitlines() or [''])[-1][:300])
            print(scen, sub, r, 'rc', p.returncode, last if p.returncode else '', flush=True)
    json.dump(status, open(os.path.join(TEST, f'status_{scen}.json'), 'w'), indent=1)


def sw2(p):
    ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
    v = list(ws.iter_rows(values_only=True)); h = v[0]
    return h, {r[0]: dict(zip(h, r)) for r in v[1:] if r[0]}


def macro(p):
    return [l for l in open(p, encoding='utf-8', errors='replace').read().replace('\r\n', '\n').split('\n')
            if not re.search(r'[Gg]enerated|20\d\d-\d\d-\d\d', l)]


def targets():
    t = {}
    for e in csv.DictReader(open(os.path.join(OUT, 'edits_77.csv'), encoding='utf-8')):
        t[e['config']] = 'a: DROP->BUILD (edits_77)'
    for e in csv.DictReader(open(os.path.join(OUT, 'edits_5406P_alias.csv'), encoding='utf-8')):
        t[e['config']] = 'b: alias row DROP->BUILD'
    rd = {r['config']: r['ready'] for r in csv.DictReader(open(os.path.join(OUT, 'rows.csv'), encoding='utf-8'))}
    for x in csv.DictReader(open(os.path.join(OUT, 'newrows.csv'), encoding='utf-8')):
        t.setdefault(x['config'], f'b: new row (ready={rd.get(x["config"])})')
    return t


def compare():
    st = {s: json.load(open(os.path.join(TEST, f'status_{s}.json'))) for s in ('base', 'applied')}
    T = targets()
    rep, built, oldchanged = ['===== applied vs base (stage)'], {}, 0
    for key in sorted(st['base']):
        sub, r = key.split(':')
        u, a = st['base'][key], st['applied'][key]
        gu, ga = os.path.join(TEST, 'gen_base', sub), os.path.join(TEST, 'gen_applied', sub)
        mu, ma = os.path.join(gu, f'Build_{r}_MASTER.bas'), os.path.join(ga, f'Build_{r}_MASTER.bas')
        tu, ta = os.path.join(gu, f'DesignTable_{r}_SW2.xlsx'), os.path.join(ga, f'DesignTable_{r}_SW2.xlsx')
        line = f'{key:22} rc {u["rc"]}->{a["rc"]}'
        if u['rc'] != a['rc']:
            line += f' | base: {u["err"] or u["out"]} | applied: {a["err"] or a["out"]}'
        if not os.path.exists(tu) and not os.path.exists(ta):
            rep.append(line + ' | no SW2 in either'); continue
        hu, ru = sw2(tu) if os.path.exists(tu) else ((), {})
        ha, ra = sw2(ta) if os.path.exists(ta) else ((), {})
        changed = {}
        for cfg, row in ru.items():
            if cfg not in ra:
                changed[cfg] = 'LOST'; continue
            d = [c for c in set(row) | set(ra[cfg]) if row.get(c) != ra[cfg].get(c)]
            if d: changed[cfg] = sorted(d)
        oldchanged += len(changed)
        new = [c for c in ra if c not in ru]
        same_macro, nd = None, 0
        if os.path.exists(mu) and os.path.exists(ma):
            lu, la = macro(mu), macro(ma)
            same_macro = lu == la
            nd = sum(1 for x, y in zip(lu, la) if x != y) + abs(len(lu) - len(la))
        line += f' | configs {len(ru)}->{len(ra)} new {len(new)} old-changed {len(changed)} | macro ' + ('identical' if same_macro else ('n/a' if same_macro is None else f'differs ({nd} lines)'))
        if hu and ha and set(hu) != set(ha):
            line += f' | SW2 columns +{sorted(set(ha) - set(hu))} -{sorted(set(hu) - set(ha))}'
        rep.append(line)
        for c in new:
            rep.append(f'      NEW {c}  [{T.get(c, "NOT A TARGET ROW!")}]')
            built[c] = key
        for c, d in changed.items():
            rep.append(f'      OLD CONFIG CHANGED {c}: {d if d == "LOST" else d[:12]}')
    # target rows not built: why
    why = {}
    for key in st['applied']:
        sub, r = key.split(':')
        sm = os.path.join(TEST, 'gen_applied', sub, f'summary_{r}.json')
        js = json.load(open(sm, encoding='utf-8')) if os.path.exists(sm) else {}
        for rk, v in js.items():
            if not isinstance(v, dict):
                continue
            for item in v.get('skipped', []) + v.get('blocked', []):
                if item and item[0] in T and item[0] not in built:
                    why.setdefault(item[0], f'{key}: ' + ' | '.join(str(x) for x in item[1:])[:400])
    rep.append('\n===== target rows')
    res = {}
    for c, k in T.items():
        res[c] = dict(kind=k, result=f'CONFIG ({built[c]})' if c in built else ('not built: ' + why.get(c, 'no generator / no summary entry')))
        rep.append(f'{c:24} {k:34} {res[c]["result"][:300]}')
    rep.append(f'\npre-existing configurations changed: {oldchanged}')
    json.dump(res, open(os.path.join(TEST, 'rows_result.json'), 'w'), indent=1)
    open(os.path.join(TEST, 'compare.txt'), 'w', encoding='utf-8').write('\n'.join(rep) + '\n')
    print('\n'.join(l for l in rep if 'NEW' in l or 'CHANGED' in l or 'rc 0->1' in l or 'rc 1->0' in l or l.startswith('=====') or 'changed:' in l))
    print(collections.Counter(v['result'].split(' ')[0] for v in res.values()))


def checks():
    os.makedirs(os.path.join(TEST, 'checks'), exist_ok=True)
    C = os.path.join(TOOLS, 'checks')
    res = json.load(open(os.path.join(TEST, 'rows_result.json')))
    roots = sorted({v['result'].split('(')[1].rstrip(')') for v in res.values() if v['result'].startswith('CONFIG')})
    sroots = [k.split(':')[1] for k in roots if k.startswith('straight:')]
    troots = [k.split(':')[1] for k in roots if k.startswith('tee:')]
    g = os.path.join(TEST, 'gen_applied')
    env = dict(os.environ, OMI_ROOT=SCEN['applied'], GS_OUT=os.path.join(g, 'straight'), GT_OUT=os.path.join(g, 'tee'), PYTHONDONTWRITEBYTECODE='1')
    runs = [('vbacheck', [os.path.join(C, 'vbacheck.py')] + [os.path.join(g, 'straight', f'Build_{r}_MASTER.bas') for r in sroots] + [os.path.join(g, 'tee', f'Build_{r}_MASTER.bas') for r in troots]),
            ('dofcheck_straight', [os.path.join(C, 'dofcheck_straight.py')] + [os.path.join(g, 'straight', f'Build_{r}_MASTER.bas') for r in sroots]),
            ('dofcheck_tee', [os.path.join(C, 'dofcheck_tee.py')] + [os.path.join(g, 'tee', f'Build_{r}_MASTER.bas') for r in troots]),
            ('allcfg_straight', [os.path.join(C, 'allcfg_check.py'), 's'] + sroots),
            ('allcfg_tee', [os.path.join(C, 'allcfg_check.py'), 't'] + troots),
            ('regress_straight', [os.path.join(C, 'regress.py'), os.path.join(TEST, 'gen_base', 'straight'), os.path.join(g, 'straight')] + sroots),
            ('regress_tee', [os.path.join(C, 'regress.py'), os.path.join(TEST, 'gen_base', 'tee'), os.path.join(g, 'tee')] + troots)]
    for name, args in runs:
        p = subprocess.run([sys.executable, '-I', '-B'] + args, env=env, capture_output=True, text=True, encoding='utf-8', errors='replace', cwd=os.path.join(TEST, 'checks'))
        open(os.path.join(TEST, 'checks', f'{name}.txt'), 'w', encoding='utf-8').write(p.stdout + ('\n--- STDERR ---\n' + p.stderr if p.stderr else ''))
        print('=====', name, 'rc', p.returncode); print(p.stdout[-2500:], p.stderr[-1500:])


if __name__ == '__main__':
    what = sys.argv[1] if len(sys.argv) > 1 else 'all'
    if what in ('apply', 'all'):
        for s in SCEN:
            assert not os.path.exists(SCEN[s]), SCEN[s]
            shutil.copytree(STAGE_DT, os.path.join(SCEN[s], '04_DESIGN_TABLES'))
        apply(os.path.join(SCEN['applied'], '04_DESIGN_TABLES'))
    if what in ('regen', 'all'):
        for s in SCEN:
            regen(s)
    if what in ('compare', 'all'):
        compare()
    if what in ('checks', 'all'):
        checks()
