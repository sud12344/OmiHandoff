import openpyxl, glob, os, csv
RT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff\tools\_out\orfs_airway_20261015\review_tables'
want = {r['config'] for r in csv.DictReader(open(r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff\tools\_out\brennan_yes_20261018\brennan_rows.csv')) if r['group']=='C-ORFS'}
SKIP = ('$PRP@AxialSource', '$COMMENT', '$PRP@Description', '$PRP@OnWebsite', '$PRP@OMiSKU', '$PRP@PartNo')
for p in sorted(glob.glob(os.path.join(RT, 'DesignTable_FS*.xlsx'))):
    wb = openpyxl.load_workbook(p, data_only=True); ws = wb.active
    hdr = [c.value for c in ws[2]]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=3)]
    hit = [r for r in rows if r[0] in want]
    if not hit: continue
    print('=====', os.path.basename(p), 'nrows', len(rows), 'want', len(hit))
    print('  HDR', [h for h in hdr if h not in SKIP])
    for r in rows:
        d = {h: v for h, v in zip(hdr, r) if h not in SKIP and h is not None}
        print(('* ' if r[0] in want else '  ') + str(r[0]), '|', ' '.join(f'{k}={v}' for k, v in d.items() if v not in (None, '')))
