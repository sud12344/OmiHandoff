"""shared loaders for wave3 ORFS (read-only on sources)"""
import openpyxl, glob, os, re, json, csv
ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
RT = os.path.join(ROOT, r'tools\_out\orfs_airway_20261015\review_tables')
CAT = os.path.join(ROOT, r'tools\_out\brennan_catalog_20261001\catalog_verified.json')
W3 = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018\wave3\orfs')

def tables():
    out = {}
    for p in sorted(glob.glob(os.path.join(RT, 'DesignTable_FS*.xlsx'))):
        ws = openpyxl.load_workbook(p, data_only=True).active
        hdr = [c.value for c in ws[2]]
        rows = []
        for r in ws.iter_rows(min_row=3, values_only=True):
            rows.append({h: v for h, v in zip(hdr, r)})
        out[os.path.basename(p)[12:-5]] = dict(hdr=hdr, rows=rows, path=p)
    return out

def dashes(pn):
    s = str(pn).upper().split(' ')[0].split(':')[-1]
    s = re.sub(r'^(FF|FS)', '', s)
    s = re.sub(r'^(0?\d{3,4}(-[A-Z]{1,2})?)', '', s)
    return tuple(int(x) for x in re.findall(r'-(\d{1,2})(?=-|$|[A-Z])', s))

def catalog():
    d = json.load(open(CAT, encoding='utf-8'))
    idx = {}
    for t in d['tables']:
        if t['catalog'] not in ('ORFS', 'FLANGE'):
            continue
        for r in t['rows']:
            idx.setdefault(r['pn'], []).append((t, r))
    return d, idx

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None
