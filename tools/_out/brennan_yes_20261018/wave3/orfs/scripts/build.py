"""wave3 ORFS (2026-10-18): Brennan catalog values for the 110 C-ORFS rows of brennan_rows.csv.

Read-only on every source. Writes into wave3/orfs/ only:
  cells.csv            cells to apply (only cells empty in the staged review tables): brennan-dim / endlib / other
  cells_proposed.csv   per-dash transfers that would complete rows (bores, stud-derived washer / hex faces, FS2406 collar) - NOT applied,
                       they need the user's OK (Q-FS-XFER, see the hand-back)
  rows.csv             table,config,ready,missing_columns,gates,conflicts,notes
  evidence/mapping.txt letter mappings + in-table / cross-table confirmation counts
Run: python -I build.py"""
import os, sys, re, csv, json, collections
from fractions import Fraction
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lib
import openpyxl

W3 = lib.W3
ROOT = lib.ROOT
BY = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018')
BRENNAN_TAG = '[Brennan = yes 2026-10-18]'
ENDLIB_TAG = '[Q-TRANSFER-ENDLIB yes 2026-10-08]'

# ------------------------------------------------------------------ numbers (rounding rule = compare_orfs.py cmp_num, copied)
def parse_num(s):
    if s is None:
        return None
    t = str(s).strip().replace('"', '')
    if t in ('', 'NA', 'N/A', '-'):
        return None
    m = re.fullmatch(r'(\d*)\.(\d+)|(\d+)', t)
    if m:
        if m.group(3) is not None:
            return Fraction(int(m.group(3))), 0
        return Fraction(Decimal(t)), len(m.group(2))
    return None


def rnd(x, n, mode):
    return (Decimal(x.numerator) / Decimal(x.denominator)).quantize(Decimal(1).scaleb(-n), rounding=mode)


def cmp_num(a, b):
    pa, pb = parse_num(a), parse_num(b)
    if pa is None or pb is None:
        return None
    (va, da), (vb, db) = pa, pb
    if va == vb:
        return 'equal'
    n = min(da, db)
    for mode in (ROUND_HALF_UP, ROUND_DOWN):
        if rnd(va, n, mode) == rnd(vb, n, mode):
            return 'rounding_only'
    if da == db == n and abs(va - vb) == Fraction(1, 10 ** n):
        if (((va + vb) / 2) * 64).denominator == 1:
            return 'rounding_only'
    return 'different'


def fnum(s):
    v = float(str(s).strip())
    return int(v) if v == int(v) else v

# ------------------------------------------------------------------ threads
FS_THR = {4: 0.5625, 6: 0.6875, 8: 0.8125, 10: 1.0, 12: 1.1875, 16: 1.4375, 20: 1.6875, 24: 2.0, 32: 2.5}   # PLAN.md s2 FS_THR
FS_CALL = {4: '9/16-18', 6: '11/16-16', 8: '13/16-16', 10: '1-14', 12: '1 3/16-12', 16: '1 7/16-12', 20: '1 11/16-12', 24: '2-12', 32: '2 1/2-12'}
UN_THR = {2: 0.3125, 3: 0.375, 4: 0.4375, 5: 0.5, 6: 0.5625, 8: 0.75, 10: 0.875, 12: 1.0625, 14: 1.1875, 16: 1.3125, 20: 1.625,
          24: 1.875, 32: 2.5}                                                                                          # SAE J1926 / J514 by dash


def thr_parse(s):
    """'1-3/16-12' / '1 7/16-12 UN' / '13/16-16 UN' / '1-14' / '1/14' -> (major, tpi, table text) or None"""
    t = re.sub(r'\s*(UNF|UN)\s*$', '', str(s).strip())
    t = re.sub(r'\s+', ' ', t)
    m = re.fullmatch(r'(\d+)[- ](\d+)/(\d+)-(\d+)', t)
    if m:
        w, a, b, p = map(int, m.groups())
        return w + a / b, p, f'{w} {a}/{b}-{p}'
    m = re.fullmatch(r'(\d+)/(\d+)-(\d+)', t)
    if m:
        a, b, p = map(int, m.groups())
        return a / b, p, f'{a}/{b}-{p}'
    m = re.fullmatch(r'(\d+)-(\d+)', t)
    if m:
        w, p = map(int, m.groups())
        return float(w), p, f'{w}-{p}'
    return None


def pipe_fmt(s):
    t = str(s).strip()
    m = re.fullmatch(r'(\d+-\d+/\d+|\d+/\d+|\d+)-(\d+)(-1/2)?', t)
    if not m:
        return None
    size = m.group(1).replace('-', ' ')
    return f'{size}-{m.group(2)}' + (' 1/2' if m.group(3) else '')


def tube_fmt(s):
    return re.sub(r'(\d)[- ](\d+/\d+)', r'\1 \2', str(s).strip())

# ------------------------------------------------------------------ per-table spec
# dims: (column, Brennan column, meaning). thr: (THDMAJ column, Brennan thread column, 'FS' | 'UN', dash index).
# prp: ($PRP column, Brennan column, 'thr' | 'pipe' | 'tube' | 'bspp'). need: columns the PLAN (s3 / s1) says the master needs.
# xfer: proposals: ('bore', col, kind, dash index) | ('stud', col, total col, ORB dash index) | ('d8', col)
S = {}
S['FS0304-C'] = dict(page='ORFS p19 FS0304-C "CAP ASSEMBLY INSERT"', new=True,
    dims=[('L@DRIVE', 'L', 'nut length (open face -> closed end, drawing L)'), ('NUTHEX@HEX', 'HEX', 'nut hex across flats'),
          ('H@DRIVE', 'H O.D.', 'insert collar OD (drawing H)'), ('K@DRIVE', 'K LGTH', 'insert length (drawing K, insert inner face -> outer face)')],
    thr=[('THDMAJ_D@DRIVE', 'B1', 'FS', 0)], prp=[('$PRP@Thread_D', 'B1', 'thr'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['L@DRIVE', 'NUTHEX@HEX', 'H@DRIVE', 'K@DRIVE', 'THDMAJ_D@DRIVE'], xfer=[])
S['FS0306'] = dict(page='ORFS p20 FS0306 "BULKHEAD LOCK NUT"',
    dims=[('H2@DRIVE', 'TT WIDTH', 'nut thickness'), ('S4HEX@HEX', 'B HEX', 'nut hex')],
    thr=[('THDMAJ_D@DRIVE', 'A THREAD', 'FS', 0)], prp=[('$PRP@Thread_D', 'A THREAD', 'thr'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['H2@DRIVE', 'S4HEX@HEX', 'THDMAJ_D@DRIVE'], xfer=[])
S['FS0319'] = dict(page='ORFS p20 FS0319 "BRAZE SLEEVE"',
    dims=[('L@DRIVE', 'F LGTH', 'sleeve length'), ('D8@DRIVE', 'A O.D.', 'sleeve collar OD')],
    thr=[], prp=[('$PRP@TubeOD', 'TUBE O.D.', 'tube')], need=['L@DRIVE', 'D8@DRIVE', '$PRP@TubeOD'], xfer=[])
S['FS1700'] = dict(page='Flange p4 FS1700 "MFS-FLANGE STRAIGHT - CODE 61"', new=True, flange=61,
    dims=[('L28@DRIVE', 'L (in.)', 'face-seal face -> flange face (overall)'), ('CHEX@HEX', 'A/F (in.)', 'body wrench flats')],
    thr=[('THDMAJ_D@DRIVE', '1 THREAD', 'FS', 0)], prp=[('$PRP@Thread_D', '1 THREAD', 'thr'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['L28@DRIVE', 'CHEX@HEX', 'O@DRIVE', 'D@BORE', 'THDMAJ_D@DRIVE'], xfer=[])
S['FS1800'] = dict(S['FS1700'], page='Flange p4 FS1800 "MFS-FLANGE STRT CODE 62"', flange=62)
S['FS1704'] = dict(page='Flange p6 FS1704-FG "MFS-FLANGE 90 DEG CODE 61"', new=True, flange=61,
    dims=[('M@DRIVE', 'L1 (in.)', 'face-seal face -> crossing (FS leg)'), ('N2@DRIVE', 'L2 (in.)', 'flange face -> crossing (flange leg)'),
          ('Y@HEX', 'A/F (in.)', 'forged body across flats')],
    thr=[('THDMAJ_D@DRIVE', '1 THREAD', 'FS', 0)], prp=[('$PRP@Thread_D', '1 THREAD', 'thr'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['M@DRIVE', 'N2@DRIVE', 'Y@HEX', 'O@DRIVE', 'D@BORE', 'DFLG@BORE', 'THDMAJ_D@DRIVE'], xfer=[])
S['FS1804'] = dict(S['FS1704'], page='Flange p7 FS1804-FG "MFS-FLANGE 90 DEG CODE 62"', flange=62)
S['FS2403'] = dict(page='ORFS p3 FS2403 "MFS-MFS STRAIGHT"',
    dims=[('L6@DRIVE', 'FF LGTH', 'overall face -> face'), ('S1HEX@HEX', 'H HEX', 'body hex')],
    thr=[('THDMAJ_D_1@DRIVE', 'A1 THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'A2 THREAD', 'FS', 1)],
    prp=[('$PRP@Thread_D_1', 'A1 THREAD', 'thr'), ('$PRP@Thread_D_2', 'A2 THREAD', 'thr')],
    need=['L6@DRIVE', 'S1HEX@HEX', 'D1_1@BORE', 'D1_2@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE'],
    xfer=[('bore', 'D1_1@BORE', 'FSM', 0), ('bore', 'D1_2@BORE', 'FSM', 1)])
S['FS2406'] = dict(page='ORFS p4 FS2406 "FFS-MFS TUBE END REDUCER"',
    dims=[('L36@DRIVE', 'L LGTH', 'collar face -> male face'), ('NUTHEX@HEX', 'C HEX', 'nut hex (FS0318 nut of the first dash, Q-FS2406-NUT)')],
    thr=[('THDMAJ_D@DRIVE', 'A2 THREAD', 'FS', 1)], prp=[('$PRP@Thread_D', 'A2 THREAD', 'thr')],
    need=['L36@DRIVE', 'D8@DRIVE', 'D1@BORE', 'NUTHEX@HEX', 'THDMAJ_D@DRIVE'],
    xfer=[('d8', 'D8@DRIVE', None, 0), ('bore', 'D1@BORE', 'FSM', 1)])
S['FS2500'] = dict(page='ORFS p7 FS2500-FG "MFS-MFS 90 DEG ELBOW"',
    dims=[('L7_1@DRIVE', 'M1 LGTH', 'leg 1 (A1) face -> crossing'), ('L7_2@DRIVE', 'M2 LGTH', 'leg 2 (A2) face -> crossing'),
          ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D_1@DRIVE', 'A1 THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'A2 THREAD', 'FS', 1)],
    prp=[('$PRP@Thread_D_1', 'A1 THREAD', 'thr'), ('$PRP@Thread_D_2', 'A2 THREAD', 'thr')],
    need=['L7_1@DRIVE', 'L7_2@DRIVE', 'S2HEX@HEX', 'D1_1@BORE', 'D1_2@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE'],
    xfer=[('bore', 'D1_1@BORE', 'FSM', 0), ('bore', 'D1_2@BORE', 'FSM', 1)])
S['FS2501'] = dict(page='ORFS p7-8 FS2501-FG "MFS-MP 90 DEG ELBOW"',
    dims=[('L7@DRIVE', 'M LGTH', 'FS leg face -> crossing'), ('N@DRIVE', 'N LGTH', 'pipe leg tip -> crossing'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D@DRIVE', 'A THREAD', 'FS', 0)],
    prp=[('$PRP@Thread_D', 'A THREAD', 'thr'), ('$PRP@Thread_A', 'B THREAD', 'pipe'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['L7@DRIVE', 'N@DRIVE', 'S2HEX@HEX', 'D1@BORE', 'D2@BORE', 'THDMAJ_D@DRIVE', '$PRP@Thread_A'],
    xfer=[('bore', 'D1@BORE', 'FSM', 0), ('bore', 'D2@BORE', 'NPTF', 1)])
S['FS2503'] = dict(S['FS2501'], page='ORFS p8 FS2503-FG "MFS-MP 45 DEG ELBOW"',
    dims=[('L12@DRIVE', 'M LGTH', 'FS leg face -> crossing'), ('N2@DRIVE', 'N LGTH', 'pipe leg tip -> crossing'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    need=['L12@DRIVE', 'N2@DRIVE', 'S2HEX@HEX', 'D1@BORE', 'D2@BORE', 'THDMAJ_D@DRIVE', '$PRP@Thread_A'])
S['FS2601'] = dict(page='ORFS p14 FS2601-FG "MFS-MFS-MP TEE"', new=True,
    dims=[('L7_1@DRIVE', 'M1 LGTH', 'run leg A1 face -> crossing'), ('L7_2@DRIVE', 'M2 LGTH', 'run leg A2 face -> crossing'),
          ('N@DRIVE', 'N LGTH', 'pipe branch tip -> crossing'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D_1@DRIVE', 'A1 THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'A2 THREAD', 'FS', 1)],
    prp=[('$PRP@Thread_D_1', 'A1 THREAD', 'thr'), ('$PRP@Thread_D_2', 'A2 THREAD', 'thr'), ('$PRP@Thread_A', 'B THREAD', 'pipe'),
         ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['L7_1@DRIVE', 'L7_2@DRIVE', 'N@DRIVE', 'S2HEX@HEX', 'D1_1@BORE', 'D1_2@BORE', 'D2@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE',
          '$PRP@Thread_A'],
    xfer=[('bore', 'D1_1@BORE', 'FSM', 0), ('bore', 'D1_2@BORE', 'FSM', 1), ('bore', 'D2@BORE', 'NPTF', 2)])
S['FS2701'] = dict(page='ORFS p8 FS2701-FG "MFS-MFS BULKHEAD 90 DEG ELBOW"',
    dims=[('L16@DRIVE', 'R LGTH', 'FS leg face -> crossing'), ('L17@DRIVE', 'S LGTH', 'bulkhead leg tip -> crossing'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D_1@DRIVE', 'A1 THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'A2 THREAD', 'FS', 1)],
    prp=[('$PRP@Thread_D_1', 'A1 THREAD', 'thr'), ('$PRP@Thread_D_2', 'A2 THREAD', 'thr')],
    need=['L16@DRIVE', 'L17@DRIVE', 'I2@DRIVE', 'I4@DRIVE', 'D17@BORE', 'S2HEX@HEX', 'D1_1@BORE', 'D1_2@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE'],
    xfer=[('bore', 'D1_1@BORE', 'FSM', 0), ('bore', 'D1_2@BORE', 'FSM', 1)])
S['FS2701-LN'] = dict(S['FS2701'])
S['FS2702'] = dict(S['FS2701'], page='ORFS p8 FS2702-FG "MFS-MFS BULKHEAD 45 DEG ELBOW"',
    dims=[('L12@DRIVE', 'P LGTH', 'FS leg face -> crossing'), ('L18@DRIVE', 'ZZ LGTH', 'bulkhead leg tip -> crossing'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    need=['L12@DRIVE', 'L18@DRIVE', 'I2@DRIVE', 'I4@DRIVE', 'D17@BORE', 'S2HEX@HEX', 'D1_1@BORE', 'D1_2@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE'])
S['FS2702-LN'] = dict(S['FS2702'])
S['FS2703'] = dict(page='ORFS p16 FS2703-FG "MFS-MFS-MFS BULKHEAD TEE"',
    dims=[('L16_1@DRIVE', 'R1 LGTH', 'run leg A1 face -> crossing'), ('L16_2@DRIVE', 'R2 LGTH', 'run leg A2 face -> crossing'),
          ('L7@DRIVE', 'S LGTH', 'bulkhead branch tip -> crossing'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D_1@DRIVE', 'A1 THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'A2 THREAD', 'FS', 1), ('THDMAJ_D_3@DRIVE', 'A3 THREAD', 'FS', 2)],
    prp=[('$PRP@Thread_D_1', 'A1 THREAD', 'thr'), ('$PRP@Thread_D_2', 'A2 THREAD', 'thr'), ('$PRP@Thread_D_3', 'A3 THREAD', 'thr')],
    need=['L16_1@DRIVE', 'L16_2@DRIVE', 'L7@DRIVE', 'I2@DRIVE', 'I4@DRIVE', 'D17@BORE', 'S2HEX@HEX', 'D1_1@BORE', 'D1_2@BORE', 'D1_3@BORE',
          'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE', 'THDMAJ_D_3@DRIVE'],
    xfer=[('bore', 'D1_1@BORE', 'FSM', 0), ('bore', 'D1_2@BORE', 'FSM', 1), ('bore', 'D1_3@BORE', 'FSM', 2)])
S['FS2703-LN'] = dict(S['FS2703'])
S['FS6400'] = dict(page='ORFS p5 FS6400-O "MFS-MORB STRAIGHT"',
    dims=[('L8@DRIVE', 'L LGTH', 'FS face -> ORB tip (overall)'), ('S5HEX@HEX', 'H HEX', 'body hex')],
    thr=[('THDMAJ_D@DRIVE', 'A THREAD', 'FS', 0), ('THDMAJ_N@DRIVE', 'B THREAD', 'UN', 1)],
    prp=[('$PRP@Thread_D', 'A THREAD', 'thr'), ('$PRP@Thread_N', 'B THREAD', 'thr')],
    need=['L8@DRIVE', 'L9@DRIVE', 'S5HEX@HEX', 'D1@BORE', 'D20@BORE', 'THDMAJ_D@DRIVE', 'THDMAJ_N@DRIVE'],
    xfer=[('stud', 'L9@DRIVE', 'L8@DRIVE', 1), ('bore', 'D1@BORE', 'FSM', 0), ('bore', 'D20@BORE', 'ORB', 1)])
S['FS6400-L'] = dict(S['FS6400'], page='ORFS p5 FS6400-L-O "MFS-MORB STRAIGHT LONG"',
    dims=[('L32@DRIVE', 'L LGTH', 'FS face -> ORB tip (overall)'), ('S5HEX@HEX', 'H HEX', 'body hex')],
    need=['L32@DRIVE', 'L31@DRIVE', 'S5HEX@HEX', 'D1@BORE', 'D20@BORE', 'THDMAJ_D@DRIVE', 'THDMAJ_N@DRIVE', '$PRP@Style'],
    xfer=[('stud', 'L31@DRIVE', 'L32@DRIVE', 1), ('bore', 'D1@BORE', 'FSM', 0), ('bore', 'D20@BORE', 'ORB', 1)])
S['FS6402'] = dict(page='ORFS p5 FS6402 "MFS-FJS STRAIGHT"', new=True,
    dims=[('L@DRIVE', 'L LGTH', 'FS face -> JIC swivel nut face (overall)'), ('S1HEX@HEX', 'H HEX', 'body hex'),
          ('SHEX@HEX', 'C HEX', 'JIC swivel nut hex')],
    thr=[('THDMAJ_D@DRIVE', 'A THREAD', 'FS', 0), ('THDMAJ_B@DRIVE', 'B THREAD', 'UN', 1)],
    prp=[('$PRP@Thread_D', 'A THREAD', 'thr'), ('$PRP@Thread_B', 'B THREAD', 'thr'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['L@DRIVE', 'S1HEX@HEX', 'SHEX@HEX', 'D1@BORE', 'D2@BORE', 'THDMAJ_D@DRIVE', 'THDMAJ_B@DRIVE'],
    xfer=[('bore', 'D1@BORE', 'FSM', 0)])
S['FS6500'] = dict(page='ORFS p9 FS6500-FG "MFS-FFSS 90 DEG ELBOW"',
    dims=[('L7@DRIVE', 'M LGTH', 'FS leg face -> crossing'), ('L21@DRIVE', 'AAA LGTH', 'swivel leg crossing -> nut face'),
          ('SHEX@HEX', 'C HEX', 'swivel nut hex'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D_1@DRIVE', 'A THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'B THREAD', 'FS', 1)],
    prp=[('$PRP@Thread_D_1', 'A THREAD', 'thr'), ('$PRP@Thread_D_2', 'B THREAD', 'thr')],
    need=['L7@DRIVE', 'L21@DRIVE', 'SHEX@HEX', 'S2HEX@HEX', 'D1@BORE', 'D19@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE'],
    xfer=[('bore', 'D1@BORE', 'FSM', 0), ('bore', 'D19@BORE', 'FSSW', 1)])
S['FS6502'] = dict(S['FS6500'], page='ORFS p9 FS6502-FG "MFS-FFSS 45 DEG ELBOW"',
    dims=[('L7@DRIVE', 'L1 LGTH', 'FS leg face -> crossing'), ('L2@DRIVE', 'L2 LGTH', 'swivel leg crossing -> nut face along the 45 deg axis (NEW column: '
          'mapping to Air-Way L52 not confirmed)'), ('SHEX@HEX', 'C HEX', 'swivel nut hex'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    need=['L7@DRIVE', 'L2@DRIVE', 'SHEX@HEX', 'S2HEX@HEX', 'D1@BORE', 'D19@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE'])
S['FS6505'] = dict(page='ORFS p6 FS6505 "MP-FFSS STRAIGHT"',
    dims=[('L16@DRIVE', 'L LGTH', 'pipe tip -> swivel nut face'), ('SHEX@HEX', 'C HEX', 'swivel nut hex'), ('S1HEX@HEX', 'H HEX', 'body hex'),
          ('D2@BORE', 'A HOLE', 'pipe-end bore'), ('D19@BORE', 'B HOLE', 'swivel-end bore')],
    thr=[('THDMAJ_D@DRIVE', 'B THREAD', 'FS', 1)],
    prp=[('$PRP@Thread_A', 'A THREAD', 'pipe'), ('$PRP@Thread_D', 'B THREAD', 'thr'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['L16@DRIVE', 'SHEX@HEX', 'S1HEX@HEX', 'D2@BORE', 'D19@BORE', 'THDMAJ_D@DRIVE', '$PRP@Thread_A'], xfer=[])
S['FS6565'] = dict(page='ORFS p6-7 FS6565 "FFSS-FFSS STRAIGHT"',
    dims=[('L45@DRIVE', 'L LGTH', 'nut face -> nut face'), ('S1HEX@HEX', 'H HEX', 'body hex'), ('S_1HEX@HEX', 'C HEX', 'swivel nut 1 hex'),
          ('S_2HEX@HEX', 'C HEX', 'swivel nut 2 hex (Brennan prints one C HEX for both nuts)')],
    thr=[('THDMAJ_D_1@DRIVE', 'A THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'B THREAD', 'FS', 1)],
    prp=[('$PRP@Thread_D_1', 'A THREAD', 'thr'), ('$PRP@Thread_D_2', 'B THREAD', 'thr')],
    need=['L45@DRIVE', 'S1HEX@HEX', 'S_1HEX@HEX', 'S_2HEX@HEX', 'D19_1@BORE', 'D19_2@BORE', 'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE'],
    xfer=[('bore', 'D19_1@BORE', 'FSSW', 0), ('bore', 'D19_2@BORE', 'FSSW', 1)])
S['FS6602'] = dict(page='ORFS p17 FS6602-FG "MFS-FFSS-MFS RUN TEE"',
    dims=[('L7_1@DRIVE', 'M1 LGTH', 'run leg A1 face -> crossing'), ('L21@DRIVE', 'AAA LGTH', 'run swivel leg crossing -> nut face'),
          ('L7@DRIVE', 'M2 LGTH', 'branch A2 face -> crossing'), ('SHEX@HEX', 'C HEX', 'swivel nut hex'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D_1@DRIVE', 'A1 THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'B THREAD', 'FS', 1), ('THDMAJ_D_3@DRIVE', 'A2 THREAD', 'FS', 2)],
    prp=[('$PRP@Thread_D_1', 'A1 THREAD', 'thr'), ('$PRP@Thread_D_2', 'B THREAD', 'thr'), ('$PRP@Thread_D_3', 'A2 THREAD', 'thr')],
    need=['L7_1@DRIVE', 'L21@DRIVE', 'L7@DRIVE', 'SHEX@HEX', 'S2HEX@HEX', 'D1_1@BORE', 'D19@BORE', 'D1_3@BORE', 'THDMAJ_D_1@DRIVE',
          'THDMAJ_D_2@DRIVE', 'THDMAJ_D_3@DRIVE'],
    xfer=[('bore', 'D1_1@BORE', 'FSM', 0), ('bore', 'D19@BORE', 'FSSW', 1), ('bore', 'D1_3@BORE', 'FSM', 2)])
S['FS6801'] = dict(page='ORFS p10 FS6801-NWO-FG "MFS-MAORB 90 DEG ELBOW"',
    dims=[('L7@DRIVE', 'M LGTH', 'FS leg face -> crossing'), ('L10@DRIVE', 'N LGTH', 'adjustable stud tip -> crossing'),
          ('S3HEX@HEX', 'C HEX', 'lock nut hex'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D@DRIVE', 'A THREAD', 'FS', 0), ('THDMAJ_N@DRIVE', 'B THREAD', 'UN', 1)],
    prp=[('$PRP@Thread_D', 'A THREAD', 'thr'), ('$PRP@Thread_N', 'B THREAD', 'thr')],
    need=['L7@DRIVE', 'L10@DRIVE', 'L11@DRIVE', 'S3HEX@HEX', 'S2HEX@HEX', 'D1@BORE', 'D20@BORE', 'THDMAJ_D@DRIVE', 'THDMAJ_N@DRIVE'],
    xfer=[('stud', 'L11@DRIVE', 'L10@DRIVE', 1), ('bore', 'D1@BORE', 'FSM', 0), ('bore', 'D20@BORE', 'ORB', 1)])
S['FS6802'] = dict(S['FS6801'], page='ORFS p10 FS6802-NWO-FG "MFS-MAORB 45 DEG ELBOW"',
    dims=[('L12@DRIVE', 'M LGTH', 'FS leg face -> crossing'), ('L13@DRIVE', 'N LGTH', 'adjustable stud tip -> crossing'),
          ('S3HEX@HEX', 'C HEX', 'lock nut hex'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    need=['L12@DRIVE', 'L13@DRIVE', 'L14@DRIVE', 'S3HEX@HEX', 'S2HEX@HEX', 'D1@BORE', 'D20@BORE', 'THDMAJ_D@DRIVE', 'THDMAJ_N@DRIVE'],
    xfer=[('stud', 'L14@DRIVE', 'L13@DRIVE', 1), ('bore', 'D1@BORE', 'FSM', 0), ('bore', 'D20@BORE', 'ORB', 1)])
S['FS6803'] = dict(page='ORFS p17 FS6803-NWO-FG "MFS-MFS-MAORB TEE"',
    dims=[('L7_1@DRIVE', 'M1 LGTH', 'run leg A1 face -> crossing'), ('L7_2@DRIVE', 'M2 LGTH', 'run leg A2 face -> crossing'),
          ('L10@DRIVE', 'N LGTH', 'adjustable stud tip -> crossing'), ('S3HEX@HEX', 'C HEX', 'lock nut hex'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats')],
    thr=[('THDMAJ_D_1@DRIVE', 'A1 THREAD', 'FS', 0), ('THDMAJ_D_2@DRIVE', 'A2 THREAD', 'FS', 1), ('THDMAJ_N@DRIVE', 'B THREAD', 'UN', 2)],
    prp=[('$PRP@Thread_D_1', 'A1 THREAD', 'thr'), ('$PRP@Thread_D_2', 'A2 THREAD', 'thr'), ('$PRP@Thread_N', 'B THREAD', 'thr')],
    need=['L7_1@DRIVE', 'L7_2@DRIVE', 'L10@DRIVE', 'L11@DRIVE', 'S3HEX@HEX', 'S2HEX@HEX', 'D1_1@BORE', 'D1_2@BORE', 'D20@BORE',
          'THDMAJ_D_1@DRIVE', 'THDMAJ_D_2@DRIVE', 'THDMAJ_N@DRIVE'],
    xfer=[('stud', 'L11@DRIVE', 'L10@DRIVE', 2), ('bore', 'D1_1@BORE', 'FSM', 0), ('bore', 'D1_2@BORE', 'FSM', 1), ('bore', 'D20@BORE', 'ORB', 2)])
S['FS6809'] = dict(page='ORFS p11 FS6809-NWO-FG "FFSS-MAORB 90 DEG ELBOW"',
    dims=[('L21@DRIVE', 'AAA LGTH', 'swivel leg crossing -> nut face'), ('L10@DRIVE', 'N LGTH', 'adjustable stud tip -> crossing'),
          ('SHEX@HEX', 'C1 HEX', 'swivel nut hex'), ('S3HEX@HEX', 'C2 HEX', 'lock nut hex'), ('S2HEX@HEX', 'JJ FLATS', 'forged flats'),
          ('D19@BORE', 'A HOLE', 'swivel-end bore'), ('D20@BORE', 'B HOLE', 'stud-end bore')],
    thr=[('THDMAJ_D@DRIVE', 'A THREAD', 'FS', 0), ('THDMAJ_N@DRIVE', 'B THREAD', 'UN', 1)],
    prp=[('$PRP@Thread_D', 'A THREAD', 'thr'), ('$PRP@Thread_N', 'B THREAD', 'thr')],
    need=['L21@DRIVE', 'L10@DRIVE', 'L11@DRIVE', 'SHEX@HEX', 'S2HEX@HEX', 'S3HEX@HEX', 'D19@BORE', 'D20@BORE', 'THDMAJ_D@DRIVE', 'THDMAJ_N@DRIVE'],
    xfer=[('stud', 'L11@DRIVE', 'L10@DRIVE', 1)])
S['FS7202'] = dict(page='ORFS p11 FS7202-NWO-FG "MFS-MBSPP 90 DEG ELBOW" (lengths and hexes printed in MILLIMETRES)', new=True, mm=True,
    dims=[('L7@DRIVE', 'M LGTH', 'FS leg face -> crossing'), ('L10@DRIVE', 'N LGTH', 'BSPP stud tip -> crossing'),
          ('S3HEX@HEX', 'C4 HEX', 'lock nut hex'), ('S2HEX@HEX', 'JJ HEX', 'forged body across flats')],
    thr=[('THDMAJ_D@DRIVE', 'A THREAD', 'FS', 0)],
    prp=[('$PRP@Thread_D', 'A THREAD', 'thr'), ('$PRP@Thread_N', 'B THREAD', 'bspp'), ('$PRP@TubeOD', 'TUBE O.D.', 'tube')],
    need=['L7@DRIVE', 'L10@DRIVE', 'L11@DRIVE', 'S3HEX@HEX', 'S2HEX@HEX', 'D1@BORE', 'D20@BORE', 'THDMAJ_D@DRIVE', 'THDMAJ_N@DRIVE'],
    xfer=[('bore', 'D1@BORE', 'FSM', 0)])

NUTDASH = {'FS2406': 0, 'FS6505': 1, 'FS6565': 0, 'FS6500': 1, 'FS6502': 1, 'FS6602': 1, 'FS6809': 0, 'FS0304-C': 0}   # dash of the nut's end
FILE = {'FS0304': 'FS0304-C'}      # brennan_rows.csv table -> review-table stem when they differ

# bore donors: table -> [(column, end kind, dash index)] (meanings: PLAN.md s3, positions checked on the staged rows)
DONORS = {
    'FS2403': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)], 'FS2404': [('D1@BORE', 'FSM', 0), ('D2@BORE', 'NPTF', 1)],
    'FS2406': [('D1@BORE', 'FSM', 1)], 'FS2500': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)],
    'FS2501': [('D1@BORE', 'FSM', 0), ('D2@BORE', 'NPTF', 1)], 'FS2503': [('D1@BORE', 'FSM', 0), ('D2@BORE', 'NPTF', 1)],
    'FS2603': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1), ('D1_3@BORE', 'FSM', 2)],
    'FS2700': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)], 'FS2700-LN': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)],
    'FS2701': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)], 'FS2701-LN': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)],
    'FS2702': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)], 'FS2702-LN': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1)],
    'FS2703': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1), ('D1_3@BORE', 'FSM', 2)],
    'FS2703-LN': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1), ('D1_3@BORE', 'FSM', 2)],
    'FS2704': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1), ('D1_3@BORE', 'FSM', 2)],
    'FS6400': [('D1@BORE', 'FSM', 0), ('D20@BORE', 'ORB', 1)], 'FS6400-L': [('D1@BORE', 'FSM', 0), ('D20@BORE', 'ORB', 1)],
    'FS6403': [('D1@BORE', 'FSM', 0)], 'FS6500': [('D1@BORE', 'FSM', 0), ('D19@BORE', 'FSSW', 1)],
    'FS6504': [('D19@BORE', 'FSSW', 0)], 'FS6505': [('D2@BORE', 'NPTF', 0), ('D19@BORE', 'FSSW', 1)],
    'FS6540': [('D19@BORE', 'FSSW', 0), ('D20@BORE', 'ORB', 1)], 'FS6565': [('D19_1@BORE', 'FSSW', 0), ('D19_2@BORE', 'FSSW', 1)],
    'FS6600': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1), ('D19@BORE', 'FSSW', 2)],
    'FS6602': [('D1_1@BORE', 'FSM', 0), ('D19@BORE', 'FSSW', 1), ('D1_3@BORE', 'FSM', 2)],
    'FS6801': [('D1@BORE', 'FSM', 0), ('D20@BORE', 'ORB', 1)], 'FS6801-L': [('D1@BORE', 'FSM', 0), ('D20@BORE', 'ORB', 1)],
    'FS6802': [('D1@BORE', 'FSM', 0), ('D20@BORE', 'ORB', 1)],
    'FS6803': [('D1_1@BORE', 'FSM', 0), ('D1_2@BORE', 'FSM', 1), ('D20@BORE', 'ORB', 2)],
    'FS6804': [('D1_1@BORE', 'FSM', 0), ('D20@BORE', 'ORB', 1), ('D1_3@BORE', 'FSM', 2)],
    'FS6809': [('D19@BORE', 'FSSW', 0), ('D20@BORE', 'ORB', 1)],
}
FORGED_OLD = {t for t in DONORS if t.startswith(('FS25', 'FS26', 'FS27', 'FS65', 'FS66', 'FS68')) and t not in ('FS2700', 'FS2700-LN', 'FS6504', 'FS6505', 'FS6540', 'FS6565')}


def main():
    T = lib.tables()
    cat, idx = lib.catalog()
    comp = {r['omi_pn']: r for r in json.load(open(os.path.join(ROOT, r'tools\_out\orfs_airway_20261015\compare_orfs.json'), encoding='utf-8'))}
    want = [r for r in csv.DictReader(open(os.path.join(BY, 'brennan_rows.csv'))) if r['group'] == 'C-ORFS']
    assert len(want) == 110, len(want)
    ff318 = {}
    for ln in open(os.path.join(ROOT, r'07_RAW_EXTRACTS\airway_passes\FF318_p1.txt'), encoding='utf-8').read().strip().split('\n')[1:]:
        v = ln.split('|')
        ff318[int(v[0].split('-')[1])] = v[4]          # dash -> S (nut hex)
    o61 = collections.defaultdict(set)                 # code 61 head OD by flange dash, from the 1700 master (base table)
    ws = openpyxl.load_workbook(os.path.join(BY, r'base_root\04_DESIGN_TABLES\DesignTable_1700.xlsx'), data_only=True).active
    hdr = [c.value for c in ws[2]]
    o61rows = collections.defaultdict(list)
    for r in ws.iter_rows(min_row=3, values_only=True):
        d = dict(zip(hdr, r))
        if d.get('O@DRIVE') is not None and '-61-' in str(d[None]):
            fd = int(str(d[None]).split('-')[3])
            o61[fd].add(d['O@DRIVE']); o61rows[fd].append(f'{d[None]} {d["O@DRIVE"]} ({d["$PRP@ScopeAction"]})')

    global FORGED
    FORGED = {t for t in T if any(str(r[None]).endswith('-FG') for r in T[t]['rows'])}
    cells, prop, rows_out, maplog = [], [], [], []

    # ---------------- confirmation of every mapping on rows of the same table with both values
    confirm = {}
    for stem, sp in S.items():
        if stem not in T:
            continue
        for col, bcol, meaning in sp['dims']:
            res = collections.Counter(); ex = []
            for r in T[stem]['rows']:
                pn = str(r[None])
                if r.get(col) in (None, '') or ':' in pn:
                    continue
                bpn = pn.replace('-ASY', '').replace('-LN-', '-')
                hit = idx.get(bpn)
                if not hit:
                    continue
                bv = hit[0][1]['values'].get(bcol)
                c = cmp_num(r[col], bv) if not sp.get('mm') else None
                res[c] += 1
                if c == 'different':
                    ex.append(f'{pn} {r[col]} vs {bv}')
            confirm[(stem, col)] = (res, ex)

    def conf_txt(stem, col):
        res, ex = confirm.get((stem, col), (collections.Counter(), []))
        n = sum(res.values())
        if not n:
            return 'no row of this table has both an Air-Way and a Brennan value (drawing mapping only; cross-check in evidence/crosscheck.txt)'
        ok = res['equal'] + res['rounding_only']
        return f'same-table check: {ok}/{n} rows with an Air-Way value agree to rounding' + (f' ({len(ex)} differ, e.g. {ex[0]})' if ex else '')

    for stem, sp in S.items():
        if stem not in T:
            continue
        maplog.append(f'== {stem}: {sp["page"]}' + (' - NEW columns (table had none)' if sp.get('new') else ''))
        for col, bcol, meaning in sp['dims']:
            maplog.append(f'   {col} <- "{bcol}" = {meaning}; {conf_txt(stem, col)}')
            for e in confirm.get((stem, col), (None, []))[1]:
                maplog.append(f'        differs: {e}')

    # ---------------- bore donor pool
    pool = collections.defaultdict(list)      # (kind, dash, forged?) -> [(value, table, pn)]
    for t, cols in DONORS.items():
        if t not in T:
            continue
        for r in T[t]['rows']:
            pn = str(r[None])
            ds = lib.dashes(pn)
            for col, kind, i in cols:
                v = r.get(col)
                if v in (None, '') or i >= len(ds):
                    continue
                pool[(kind, ds[i])].append((v, t, pn))

    def bore_prop(stem, kind, dash):
        same = [x for x in pool[(kind, dash)] if x[1] == stem]
        if same:
            vals = sorted({x[0] for x in same})
            return (vals[0] if len(vals) == 1 else None), f'same table {stem}: {kind} -{dash:02d} bore on {len(same)} Air-Way row(s): ' + \
                ', '.join(sorted({str(x[0]) for x in same})) + (' (unanimous)' if len(vals) == 1 else ' (NOT unanimous - no proposal)')
        fam = stem in FORGED
        grp = [x for x in pool[(kind, dash)] if (x[1] in FORGED) == fam]
        if not grp:
            return None, f'no Air-Way {kind} -{dash:02d} bore in any {"forged" if fam else "straight"} ORFS table'
        vals = sorted({x[0] for x in grp})
        c = collections.Counter(x[0] for x in grp)
        return (vals[0] if len(vals) == 1 else None), f'{"forged" if fam else "straight"} ORFS tables: {kind} -{dash:02d} bore on {len(grp)} Air-Way row(s) in ' + \
            f'{len({x[1] for x in grp})} tables: ' + ', '.join(f'{k} x{n}' for k, n in sorted(c.items())) + (' (unanimous)' if len(vals) == 1 else ' (NOT unanimous - no proposal)')

    # ---------------- stud pool (Air-Way: total leg - face/washer position, per ORB dash, same table)
    def stud_prop(stem, col, tot, di):
        vals = collections.defaultdict(set)
        for r in T[stem]['rows']:
            pn = str(r[None])
            a, b = lib.num(r.get(tot)), lib.num(r.get(col))
            ds = lib.dashes(pn)
            if a is None or b is None or di >= len(ds):
                continue
            vals[ds[di]].add(round(a - b, 3))
        return vals

    # ---------------- per row
    for w in want:
        stem = FILE.get(w['table'], w['table'])
        pn = w['config']
        if '-LN-' in pn and stem + '-LN' in T:
            stem = stem + '-LN'
        if pn.startswith('FS6400-L-'):
            stem = 'FS6400-L'
        sp = S[stem]
        tab = T[stem]
        row = [r for r in tab['rows'] if r[None] == pn]
        assert len(row) == 1, (stem, pn)
        row = row[0]
        bpn = pn.replace('-ASY', '').replace('-LN-', '-')
        hit = idx.get(bpn)
        gates = ['ORFS model type (no ORFS generator yet; waits for the 9 pilot STEPs, STEPS_NEEDED S1)']
        conflicts, notes, filled, missing = [], [], {}, []
        if not hit:
            rows_out.append(dict(table=stem, config=pn, ready='no', missing_columns='every dimension column', gates=' ; '.join(gates), conflicts='',
                                 notes='NO SOURCE: Brennan prints no row'))
            continue
        bt, br = hit[0]
        bv = br['values']
        where = f'Brennan {bt["catalog"]} catalog p{bt["page"]} table {bt["table_id"]} "{bt["series"]} {bt["title"]}" row {br["pn"]}'
        pdf = 'ORFS Catalog.pdf' if bt['catalog'] == 'ORFS' else 'Flange Catalog.pdf'
        if pn != bpn:
            notes.append(f'Brennan prints {bpn} (' + ('the cap assembly insert & nut; OMi -ASY = the staked assembly' if '-ASY' in pn else
                         'the bare bulkhead fitting; the -LN kit adds the FS0306 lock nut, which is not modelled (PLAN)') + ')')
        held = pn == 'FS0319-24'
        if held:
            gates.append('HELD plan D5 / compare_orfs HOLD: Air-Way contradicts itself (FF319-24 d8 1.71 vs FF2407 collar 2.00); Brennan A O.D. 1.910 '
                         'is a third value - the user picks (never pick between conflicting sources)')

        def add(col, val, src, kind, explain=''):
            if row.get(col) not in (None, ''):
                return
            filled[col] = val
            if not held:
                cells.append(dict(table=stem, config=pn, column=col, value=val, source=src, kind=kind))

        for col, bcol, meaning in sp['dims']:
            raw = bv.get(bcol)
            if raw in (None, ''):
                continue
            if pn == 'FS6402-12-12' and col == 'L@DRIVE':
                conflicts.append(f'L: Brennan prints L LGTH "0.609" for FS6402-12-12 (the -10 row 2.060, -16 2.460; a misprint) - not used, L@DRIVE stays empty')
                continue
            if sp.get('mm'):
                val = round(float(raw) / 25.4, 4)
                src = (f'{where} col "{bcol}" = "{raw}" MILLIMETRES -> {val} in (/25.4) ({pdf}; catalog_verified.json line_y {br["line_y"]}, '
                       f'page verified {bt.get("verified")}) {BRENNAN_TAG} - {meaning}')
            else:
                val = fnum(raw)
                src = (f'{where} col "{bcol}" = "{raw}" (07_RAW_EXTRACTS/brennan_catalogs/{pdf}; tools/_out/brennan_catalog_20261001/catalog_verified.json, '
                       f'line_y {br["line_y"]}, page verified {bt.get("verified")}) {BRENNAN_TAG} - {meaning}; mapping: {conf_txt(stem, col)}')
            # Q-FS-NUT: swivel / reducer nut hex vs Air-Way FF318 S of the same dash
            if col in ('SHEX@HEX', 'NUTHEX@HEX', 'S_1HEX@HEX', 'S_2HEX@HEX') and stem not in ('FS6402',):
                ds = lib.dashes(pn)
                nd = ds[NUTDASH.get(stem, 0)]
                s = ff318.get(nd)
                c = cmp_num(raw, s)
                if c in ('equal', 'rounding_only'):
                    src += f'; Q-FS-NUT: = Air-Way FF318-{nd} S {s} to rounding'
                else:
                    conflicts.append(f'{col}: Brennan {raw} vs Air-Way FF318-{nd} nut S {s} - Q-FS-NUT says use FF318 S where the row disagrees; '
                                     f'the Brennan value is listed (the user decides)')
            res, ex = confirm.get((stem, col), (collections.Counter(), []))
            nn = sum(res.values())
            if nn and (res['equal'] + res['rounding_only']) * 2 < nn:
                conflicts.append(f'{col}: Brennan "{bcol}" agrees with Air-Way {col} on only {res["equal"] + res["rounding_only"]}/{nn} rows of this table '
                                 f'({"; ".join(ex)}) - same meaning on both drawings, different sizes: this row would be Brennan-sized in an Air-Way-sized table')
            add(col, val, src, 'brennan-dim')

        for col, bcol, typ, di in sp['thr']:
            raw = bv.get(bcol)
            p = thr_parse(raw) if raw else None
            ds = lib.dashes(pn)
            ref = (FS_THR if typ == 'FS' else UN_THR).get(ds[di]) if di < len(ds) else None
            if ref is None:
                conflicts.append(f'{col}: no {typ} thread for this dash'); continue
            if p is None:
                conflicts.append(f'{col}: Brennan prints the thread "{raw}" (a printing slip) - built with the -{ds[di]:02d} {typ} thread {ref} (PLAN s8 rule: threads by dash)')
                p = (ref, None, None)
            if abs(p[0] - ref) > 1e-6:
                conflicts.append(f'{col}: Brennan prints "{raw}" ({p[0]}) but the -{ds[di]:02d} {typ} thread is {ref} - built with the dash thread (PLAN s8 rule)')
            add(col, ref, f'end library: {"face-seal (SAE J1453) thread FS_THR" if typ == "FS" else "straight (UN/UNF) thread"} of dash -{ds[di]:02d} '
                          f'= {ref} (PLAN.md s2 FS_THR / the -{ds[di]:02d} thread of every staged ORFS row); = Brennan "{bcol}" "{raw}" of this row {ENDLIB_TAG}',
                'endlib')
        for col, bcol, typ in sp['prp']:
            raw = bv.get(bcol)
            if raw in (None, ''):
                continue
            if typ == 'thr':
                p = thr_parse(raw); v = p[2] if p else None
                if v is None:
                    dd = lib.dashes(pn)
                    k = {'A THREAD': 0, 'A1 THREAD': 0, 'B THREAD': 1, 'A2 THREAD': 1, 'A3 THREAD': 2, 'B1': 0, '1 THREAD': 0}.get(bcol, 0)
                    if stem == 'FS6602' and bcol == 'B THREAD':
                        k = 1
                    if stem == 'FS6602' and bcol == 'A2 THREAD':
                        k = 2
                    v = FS_CALL.get(dd[k]) if k < len(dd) else None
                    if v:
                        add(col, v, f'face-seal callout of dash -{dd[k]:02d} (Brennan prints "{raw}" for this row: printing slip; PLAN s8 rule)', 'other')
                        continue
            elif typ == 'pipe':
                v = pipe_fmt(raw)
            elif typ == 'bspp':
                v = f'{raw} BSPP'
            else:
                v = tube_fmt(raw)
            if v is None:
                conflicts.append(f'{col}: "{raw}" not parsed'); continue
            add(col, v, f'{where} col "{bcol}" "{raw}" in the table\'s format', 'other')

        # flange head OD (code 61 from the 1700 master; code 62 beyond -08 needs a STEP)
        if sp.get('flange'):
            fd = lib.dashes(pn)[1]
            if sp['flange'] == 61 and o61.get(fd):
                vals = sorted(o61[fd])
                v = 2.38 if fd == 24 else (vals[0] if len(vals) == 1 else None)
                if v is not None:
                    if fd == 24:
                        gates.append('Q-FLG61-O24 (the -24 code 61 head is printed 2.37 / 2.38 in the 1700 master; 2.38 used as wave 2 did - recommended, not answered)')
                    add('O@DRIVE', v, f'end library: code 61 head OD of flange -{fd:02d} = {v} as built in the 1700 master (base 1700 rows: '
                        + '; '.join(o61rows[fd]) + f') {ENDLIB_TAG} - FLG61 head by flange size', 'endlib')
            else:
                notes.append(f'O@DRIVE: code 62 -{fd:02d} head not in any table (FLANGE62 measured -08 only) - needs BRENNAN_{pn}.stp or a code 62 head STEP')
            notes.append('D@BORE' + (' / DFLG@BORE' if 'DFLG@BORE' in sp['need'] else '') + f': the catalog prints no bore - needs BRENNAN_{pn}.stp (as the JIC 1700 / 1704 rows)')
            if stem in ('FS1700', 'FS1800'):
                notes.append('A16: Brennan\'s straight flange body is a round collar + two flats (JIC 1700 rows: COLLARD / FLATL / NECKD from the own STEP); '
                             'the same needs this row\'s STEP if the FS master follows A16')

        # comment
        cells.append(dict(table=stem, config=pn, column='$COMMENT', kind='other', source='note appended to the row comment',
                          value=f'2026-10-18 Brennan = yes (wave 3 ORFS, staged, not shipped): sizes from Brennan {bt["catalog"]} p{bt["page"]} {br["pn"]} '
                                f'($PRP@DimSource BRENNAN; to undo clear the BRENNAN cells and regenerate); threads by dash (end library)')) if not held else None

        # ------------- proposals (not applied)
        for x in sp['xfer']:
            if x[0] == 'bore':
                _, col, kind, di = x
                if row.get(col) not in (None, '') or col in filled:
                    continue
                ds = lib.dashes(pn)
                v, why = bore_prop(stem, kind, ds[di])
                if v is not None:
                    prop.append(dict(table=stem, config=pn, column=col, value=v, kind='endlib',
                                     source=f'PROPOSED per-dash bore transfer (Q-FS-XFER, not approved): {why} [Air-Way, not Brennan]'))
                else:
                    notes.append(f'{col}: no proposal - {why}')
            elif x[0] == 'stud':
                _, col, tot, di = x
                if tot not in filled and row.get(tot) in (None, ''):
                    continue
                ds = lib.dashes(pn)
                vals = stud_prop(stem, col, tot, di).get(ds[di], set())
                if len(vals) == 1:
                    st = next(iter(vals))
                    v = round(float(filled.get(tot, row.get(tot))) - st, 3)
                    prop.append(dict(table=stem, config=pn, column=col, value=v, kind='endlib',
                                     source=f'PROPOSED derived (Q-FS-XFER, not approved): {tot} {filled.get(tot, row.get(tot))} (Brennan) - Air-Way stud length {st} '
                                            f'of ORB -{ds[di]:02d} ({tot} - {col} on every Air-Way row of {stem} with this ORB dash; Q-FS-ORBSTUD (a) end ORB_FS / NWO_FS)'))
                else:
                    notes.append(f'{col}: no proposal - Air-Way stud length of ORB -{ds[di]:02d} in {stem}: ' + (', '.join(map(str, sorted(vals))) + ' (not unanimous)' if vals else 'no Air-Way row'))
            elif x[0] == 'd8':
                col = x[1]
                d1 = lib.dashes(pn)[0]
                vals = sorted({r.get(col) for r in tab['rows'] if r.get(col) is not None and lib.dashes(r[None])[:1] == (d1,) and ':' not in str(r[None])})
                if len(vals) == 1:
                    prop.append(dict(table=stem, config=pn, column=col, value=vals[0], kind='endlib',
                                     source=f'PROPOSED per-dash collar transfer (Q-FS-XFER): every Air-Way FS2406 row with female dash -{d1:02d} prints D8 {vals[0]}'))
                else:
                    notes.append(f'{col}: no proposal - Air-Way FS2406 rows with female dash -{d1:02d} print ' + (', '.join(map(str, vals)) or 'nothing') +
                                 f'; Brennan FS0319-{d1:02d} sleeve A O.D. = ' + str((idx.get(f'FS0319-{d1:02d}') or [(None, {"values": {}})])[0][1]['values'].get('A O.D.')))
        if stem in ('FS2701', 'FS2701-LN', 'FS2702', 'FS2702-LN', 'FS2703', 'FS2703-LN'):
            notes.append('I2 / I4 / D17 (bulkhead thread length, collar): Brennan prints none; the only Air-Way -16 FSBH values are FS2704-16-16-16-FG / -LN '
                         'I2 1.65, I4 0.06, D17 1.62 and FS2700-16-16 I2 1.66 (FF2704-20-20-16 D17 1.83) - not unanimous, no proposal')
        if stem == 'FS6400-L':
            notes.append('$PRP@Style: Brennan prints no style (Air-Way A / B; only A builds) - needs the user / the pilot')
        if stem == 'FS7202':
            gates.append('BSPP stud (international thread): out of scope by the skill rule unless the user keeps it (PLAN s1); THDMAJ_N / ORB-type stud data for BSPP not in any library')
        if stem == 'FS6502':
            conflicts.append('L2@DRIVE: Brennan L2 LGTH vs Air-Way FF6502 L52 differ on all 6 sizes (0.94/0.99, 1.06/1.12, 1.32/1.49, 1.44/1.53, 1.48/1.73, '
                             '1.52/1.87) - L52 meaning open (PLAN s1 "the L52 meaning, Air-Way drawing s30"); Brennan L2 put in a NEW column L2@DRIVE, L52@DRIVE left empty')
        if stem == 'FS6809' and 'D20@BORE' in filled:
            conflicts.append('D20@BORE: Brennan B HOLE runs below Air-Way d20 on all 3 shared sizes (-06 0.264/0.30, -16 0.811/0.85, -20 1.024/1.08); A HOLE = d19 3/3')
        if pn == 'FS2601-10-10-08-FG':
            conflicts.append('N@DRIVE: Brennan prints N 1.59 here, but its own FS2501-10-08-FG pipe leg is 1.470 (1.59 = FS2501-10-12-FG); the other FS2601 rows equal the FS2501 leg of the same pair (7/8, evidence/crosscheck.txt) - printed value used, worth a check')
        if stem == 'FS0304-C':
            notes.append('the insert shape (inner face, staking) needs a STEP (PLAN s1: wave-1 0304 cap-nut pattern + an insert); nut L agrees with Brennan FS0318 Q / Air-Way FF318 h within 0.01 (evidence/crosscheck.txt)')
        cr = comp.get(pn, {})
        if 'differs' in cr.get('reason', ''):
            conflicts.append('Air-Way NON-forged row (not used, forged rule 2026-10-15) differs from Brennan forged: ' + cr['reason'].split('differs from Brennan forged on ')[-1].split(' - stays')[0])

        miss = [c for c in sp['need'] if row.get(c) in (None, '') and c not in filled]
        propd = {p['column'] for p in prop if p['config'] == pn}
        left = [c for c in miss if c not in propd]
        if held:
            status = 'HELD (conflict, values not applied: L 0.610 / D8 1.910 Brennan)'
        elif not miss:
            status = 'DATA COMPLETE'
        elif not left:
            status = 'PARTIAL - complete with cells_proposed.csv (' + ', '.join(miss) + ')'
        else:
            status = 'PARTIAL - missing ' + ', '.join(left) + ((' (proposed: ' + ', '.join(sorted(propd & set(miss))) + ')') if propd & set(miss) else '')
        rows_out.append(dict(table=stem, config=pn, ready='no', missing_columns=' ; '.join(miss), gates=' ; '.join(gates),
                             conflicts=' ; '.join(conflicts), notes=' ; '.join([status] + notes)))

    os.makedirs(os.path.join(W3, 'evidence'), exist_ok=True)
    for name, data, cols in (('cells.csv', cells, ['table', 'config', 'column', 'value', 'source', 'kind']),
                             ('cells_proposed.csv', prop, ['table', 'config', 'column', 'value', 'source', 'kind']),
                             ('rows.csv', rows_out, ['table', 'config', 'ready', 'missing_columns', 'gates', 'conflicts', 'notes'])):
        with open(os.path.join(W3, name), 'w', newline='', encoding='utf-8') as f:
            wr = csv.DictWriter(f, fieldnames=cols); wr.writeheader(); wr.writerows(data)
    open(os.path.join(W3, 'evidence', 'mapping.txt'), 'w', encoding='utf-8').write('\n'.join(maplog) + '\n')
    st = collections.Counter(r['notes'].split(' ;')[0].split(' - ')[0].split(' (')[0] for r in rows_out)
    print(len(cells), 'cells;', len(prop), 'proposed;', dict(st))


if __name__ == '__main__':
    main()
