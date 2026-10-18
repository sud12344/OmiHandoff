import openpyxl, glob, os, csv, sys
RT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff\tools\_out\orfs_airway_20261015\review_tables'
want = {r['config'] for r in csv.DictReader(open(r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff\tools\_out\brennan_yes_20261018\brennan_rows.csv')) if r['group']=='C-ORFS'}
for p in sorted(glob.glob(os.path.join(RT, 'DesignTable_FS*.xlsx'))):
    wb = openpyxl.load_workbook(p, data_only=True); ws = wb.active
    hdr = [c.value for c in ws[2]]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=3)]
    hit = [r for r in rows if r[0] in want]
    if not hit and '-a' not in sys.argv: continue
    print('=====', os.path.basename(p), 'row1:', [c.value for c in ws[1]][:3], 'nrows', len(rows), 'want', len(hit))
    print(hdr)
    for r in rows:
        if r[0] in want or '-v' in sys.argv:
            print(('*' if r[0] in want else ' '), {h: v for h, v in zip(hdr, r) if v not in (None, '') and h not in ('$PRP@AxialSource',)})
