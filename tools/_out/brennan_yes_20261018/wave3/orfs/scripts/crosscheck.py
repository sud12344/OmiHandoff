"""cross-table confirmation of the letter mappings of the ORFS tables that have no Air-Way value of their own (read-only)
writes wave3/orfs/evidence/crosscheck.txt"""
import os, sys, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lib, build
import openpyxl

T = lib.tables()
cat, idx = lib.catalog()
AW = os.path.join(lib.ROOT, r'07_RAW_EXTRACTS\airway_passes')


def aw(code):
    L = open(os.path.join(AW, f'{code}_p1.txt'), encoding='utf-8').read().strip().split('\n')
    h = L[0].split('|')
    return {l.split('|')[0]: dict(zip(h, l.split('|'))) for l in L[1:]}


out = []


def rep(title, pairs):
    c = collections.Counter(build.cmp_num(a, b) for _, a, b in pairs)
    ok = c['equal'] + c['rounding_only']
    out.append(f'{title}: {ok}/{len(pairs)} agree to rounding' +
               ''.join(f'\n      differs: {k} {a} vs {b}' for k, a, b in pairs if build.cmp_num(a, b) not in ('equal', 'rounding_only')))


def br(pn, col):
    h = idx.get(pn)
    return h[0][1]['values'].get(col) if h else None


ff318, ff304, ff6502 = aw('FF318'), aw('FF304-C'), aw('FF6502')
d = [4, 6, 8, 10, 12, 16, 20, 24]
rep('FS0304-C "L" (nut length) vs Air-Way FF318 h (tube nut, same dash)',
    [(f'-{x:02d}', br(f'FS0304-C-{x:02d}', 'L'), ff318[f'FF318-{x}']['h']) for x in d if f'FF318-{x}' in ff318])
rep('FS0304-C "HEX" vs Air-Way FF318 S (tube nut hex)',
    [(f'-{x:02d}', br(f'FS0304-C-{x:02d}', 'HEX'), ff318[f'FF318-{x}']['S']) for x in d if f'FF318-{x}' in ff318])
rep('FS0304-C "HEX" vs Air-Way FF304-C S (cap nut hex)',
    [(f'-{x:02d}', br(f'FS0304-C-{x:02d}', 'HEX'), ff304[f'FF304-C-{x}']['S']) for x in d if f'FF304-C-{x}' in ff304])
rep('FS0304-C "H O.D." (insert collar) vs Brennan FS0319 "A O.D." (braze sleeve collar, same dash)',
    [(f'-{x:02d}', br(f'FS0304-C-{x:02d}', 'H O.D.'), br(f'FS0319-{x:02d}', 'A O.D.')) for x in d])
rep('FS0304-C "L" vs Brennan FS0318 "Q WIDTH" (nut length)', [(f'-{x:02d}', br(f'FS0304-C-{x:02d}', 'L'), br(f'FS0318-{x:02d}', 'Q WIDTH')) for x in d])

legs = {}
for t in ('FS2500', 'FS2603', 'FS2650'):
    for r in T[t]['rows']:
        ds = lib.dashes(r[None])
        if r.get('L7_1@DRIVE') is not None and len(set(ds)) == 1 and ':' not in str(r[None]):
            legs.setdefault(ds[0], set()).add(r['L7_1@DRIVE'])
pairs = []
for pn in sorted(p for p in idx if p.startswith('FS2601-')):
    ds = lib.dashes(pn)
    for k in ('M1 LGTH', 'M2 LGTH'):
        for v in sorted(legs.get(ds[0], [])):
            pairs.append((pn + ' ' + k, br(pn, k), v))
rep('FS2601 "M1 / M2" (run legs) vs Air-Way L7 of the same FS dash on FS2500 / FS2603 / FS2650 rows', pairs)
n2501 = {lib.dashes(r[None]): r.get('N@DRIVE') for r in T['FS2501']['rows'] if r.get('N@DRIVE') is not None and ':' not in str(r[None])}
rep('FS2601 "N" (pipe branch) vs Air-Way FS2501 N of the same (FS, pipe) dash pair',
    [(pn, br(pn, 'N LGTH'), n2501[(lib.dashes(pn)[0], lib.dashes(pn)[2])]) for pn in sorted(idx) if pn.startswith('FS2601-') and (lib.dashes(pn)[0], lib.dashes(pn)[2]) in n2501])
rep('FS2601 "N" vs Brennan FS2501-FG "N LGTH" of the same (FS, pipe) pair',
    [(pn, br(pn, 'N LGTH'), br(f'FS2501-{lib.dashes(pn)[0]:02d}-{lib.dashes(pn)[2]:02d}-FG', 'N LGTH')) for pn in sorted(idx)
     if pn.startswith('FS2601-') and br(f'FS2501-{lib.dashes(pn)[0]:02d}-{lib.dashes(pn)[2]:02d}-FG', 'N LGTH')])

ws = openpyxl.load_workbook(os.path.join(lib.ROOT, r'tools\_out\brennan_yes_20261018\base_root\04_DESIGN_TABLES\DesignTable_6402.xlsx'), data_only=True).active
h = [c.value for c in ws[2]]
nut = collections.defaultdict(set)
for r in ws.iter_rows(min_row=3, values_only=True):
    x = dict(zip(h, r))
    if x.get('C3HEX@HEX') is not None:
        nut[lib.dashes(x[None])[1]].add(x['C3HEX@HEX'])
rep('FS6402 "C HEX" vs the JIC swivel nut hex C3HEX of the same JIC dash in the 6402 master table',
    [(pn, br(pn, 'C HEX'), v) for pn in sorted(idx) if pn.startswith('FS6402-') for v in sorted(nut.get(lib.dashes(pn)[1], []))])
rep('FS6402 "H HEX" vs Brennan FS6403 "H HEX" of the same pair (body hex of the MFS x MJ straight)',
    [(pn, br(pn, 'H HEX'), br(pn.replace('FS6402', 'FS6403'), 'H HEX')) for pn in sorted(idx) if pn.startswith('FS6402-') and br(pn.replace('FS6402', 'FS6403'), 'H HEX')])

f = []
for pn in sorted(p for p in idx if p.startswith('FS6502-')):
    ds = lib.dashes(pn)
    f.append((pn, 'FF6502-' + '-'.join(str(x) for x in ds)))
rep('FS6502-FG "L1 LGTH" vs Air-Way FF6502 L7 (plain)', [(a, br(a, 'L1 LGTH'), ff6502[b]['L7']) for a, b in f if b in ff6502])
rep('FS6502-FG "L2 LGTH" vs Air-Way FF6502 L52 (plain)', [(a, br(a, 'L2 LGTH'), ff6502[b]['L52']) for a, b in f if b in ff6502])
rep('FS6502-FG "C HEX" vs Air-Way FF6502 S (swivel nut)', [(a, br(a, 'C HEX'), ff6502[b]['S']) for a, b in f if b in ff6502])
rep('FS6502-FG "JJ FLATS" vs Air-Way FF6502 S2', [(a, br(a, 'JJ FLATS'), ff6502[b]['S2']) for a, b in f if b in ff6502])
rep('FS6502-FG "L2 LGTH" vs Brennan FS6500-FG "AAA LGTH" (90 deg swivel leg, same size) - information only',
    [(a, br(a, 'L2 LGTH'), br(a.replace('FS6502', 'FS6500'), 'AAA LGTH')) for a, b in f if br(a.replace('FS6502', 'FS6500'), 'AAA LGTH')])

pairs = []
for pn in sorted(p for p in idx if p.startswith('FS7202-')):
    k = lib.dashes(pn)[0]
    ref = br(f'FS6801-{k:02d}-{k:02d}-NWO-FG', 'M LGTH')
    if ref:
        pairs.append((pn, f'{float(br(pn, "M LGTH")) / 25.4:.2f}', ref))
rep('FS7202 "M LGTH" mm/25.4 vs Brennan FS6801-NWO-FG "M LGTH" of the same FS dash (inches, 0.01)', pairs)
out.append('FS1700 / FS1800 / FS1704 / FS1804: no Air-Way ORFS flange series; the letters follow the Flange catalog drawings (p4 / p6 / p7: L = '
           'face-seal face -> flange face, L1 = FS leg to the crossing, L2 = flange face to the crossing, A/F = flats), the drawing convention wave 2 '
           'checked on the JIC 1700 / 1703 / 1704 / 1804 rows against their own Brennan STEPs (catalog L 9/9; L1 / L2 / A/F reproduced). '
           'No FS flange STEP is on disk.')
os.makedirs(os.path.join(lib.W3, 'evidence'), exist_ok=True)
open(os.path.join(lib.W3, 'evidence', 'crosscheck.txt'), 'w', encoding='utf-8').write('\n'.join(out) + '\n')
print('\n'.join(out))
