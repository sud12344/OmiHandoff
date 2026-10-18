"""apply_brennan.py <stage_dir> <cells.csv> [<cells.csv> ...] [--rows <rows.csv> ...]

Applies Brennan cell lists (from the wave agents) to a COPY of the base tables, with the reversible marking the user asked for on
2026-10-18 ("yes use brennan but still differentiate the two in tables in case i need to change it later"):
  - column $PRP@DimSource is added to a table when missing (tested: generators ignore it - marker_test);
  - a row that takes any 'brennan-dim' cell gets $PRP@DimSource = BRENNAN and its $PRP@AxialSource is prefixed
    "[BRENNAN-ONLY 2026-10-18] " and appended with the per-cell sources;
  - rows listed ready=yes in a rows.csv go to DataStatus ASSUMED (one source); others keep their status;
  - rows filled today from OMi's own answers (omi_answers_20261018/CHANGES.csv) get DimSource OMI-DRAWING (the 14 drawn parts) or OMI.
Only EMPTY cells are filled (asserted). <stage_dir>/04_DESIGN_TABLES is created from base_root on first use.
Writes <stage_dir>/CHANGES.csv (every cell, old -> new, source)."""
import os, sys, csv, shutil, collections
import openpyxl

ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
BY = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018')
BASE = os.path.join(BY, r'base_root\04_DESIGN_TABLES')
OMI_CH = os.path.join(ROOT, r'tools\_out\omi_answers_20261018\CHANGES.csv')
TAG = '[BRENNAN-ONLY 2026-10-18] '
COL = '$PRP@DimSource'

args = sys.argv[1:]
stage = args.pop(0)
cells_f, rows_f = [], []
cur = cells_f
for a in args:
    if a == '--rows':
        cur = rows_f; continue
    cur.append(a)
DT = os.path.join(stage, '04_DESIGN_TABLES')
if not os.path.exists(DT):
    shutil.copytree(BASE, DT)

cells = [r for f in cells_f for r in csv.DictReader(open(f, encoding='utf-8-sig'))]
ready = {(r['table'], r['config']) for f in rows_f for r in csv.DictReader(open(f, encoding='utf-8-sig')) if r['ready'].strip().lower() == 'yes'}
books, log = {}, []


def sheet(t):
    if t not in books:
        p = os.path.join(DT, f'DesignTable_{t}.xlsx')
        wb = openpyxl.load_workbook(p); ws = wb.active
        hdr = {c.value: c.column for c in ws[2] if c.value is not None}
        if COL not in hdr:
            c = ws.max_column + 1; ws.cell(2, c, COL); hdr[COL] = c
            log.append((f'DesignTable_{t}.xlsx', '(header)', 2, COL, ws.cell(2, c).coordinate, '', COL, 'new column: dimension source of the row (2026-10-18 marking)'))
        books[t] = (p, wb, ws, hdr)
    return books[t]


def rowof(ws, pn):
    rr = [c.row for c in ws['A'] if c.value == pn]
    assert len(rr) == 1, (ws.title, pn, rr)
    return rr[0]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


per_row = collections.defaultdict(list)
MARKING = ('$PRP@DataStatus', '$PRP@DimSource', '$PRP@AxialSource')    # set below by this script, never taken from a cell list
for c in cells:
    t, pn, h = c['table'].strip(), c['config'].strip(), c['column'].strip()
    if h in MARKING:
        continue
    _, _, ws, hdr = sheet(t)
    if h not in hdr:                         # a new dimension column some masters need (e.g. SOCKAF@SOCKET): added at the end of row 2
        col = ws.max_column + 1; ws.cell(2, col, h); hdr[h] = col
        log.append((f'DesignTable_{t}.xlsx', '(header)', 2, h, ws.cell(2, col).coordinate, '', h, 'new column needed by this row'))
    r = rowof(ws, pn)
    cell = ws.cell(r, hdr[h])
    if h == '$COMMENT':                      # notes are appended, never replace
        old = cell.value or ''
        v = c['value'].split(' | before: ')[0]           # agents may echo the old comment after "| before:" - keep it once
        add = v if not old or v not in old else ''
        if add:
            cell.value = (old + ' | ' if old else '') + add
            log.append((f'DesignTable_{t}.xlsx', pn, r, h, cell.coordinate, old, cell.value, f"[note] {c['source']}"))
        continue
    assert cell.value in (None, ''), (t, pn, h, cell.value, 'cell not empty')
    cell.value = num(c['value']) if not h.startswith('$') else c['value']
    log.append((f'DesignTable_{t}.xlsx', pn, r, h, cell.coordinate, '', c['value'], f"[{c['kind']}] {c['source']}"))
    per_row[(t, pn)].append(c)

for (t, pn), cs in per_row.items():
    _, _, ws, hdr = sheet(t)
    r = rowof(ws, pn)
    brennan = any(c['kind'] == 'brennan-dim' for c in cs)
    if brennan:
        cell = ws.cell(r, hdr[COL]); old = cell.value
        cell.value = 'BRENNAN'
        log.append((f'DesignTable_{t}.xlsx', pn, r, COL, cell.coordinate, old or '', 'BRENNAN', 'row takes sizes from Brennan alone'))
        cell = ws.cell(r, hdr['$PRP@AxialSource']); old = cell.value or ''
        srcs = '; '.join(f"{c['column']}: {c['source']}" for c in cs)
        cell.value = TAG + old + (' -> ' if old else '') + f'2026-10-18 Brennan allowed (user: "yes use brennan but still differentiate"): {srcs}'
        log.append((f'DesignTable_{t}.xlsx', pn, r, '$PRP@AxialSource', cell.coordinate, old, cell.value, 'marking + per-cell sources'))
    if (t, pn) in ready:
        cell = ws.cell(r, hdr['$PRP@DataStatus']); old = cell.value
        if old != 'ASSUMED':
            cell.value = 'ASSUMED'
            log.append((f'DesignTable_{t}.xlsx', pn, r, '$PRP@DataStatus', cell.coordinate, old or '', 'ASSUMED', 'every driving cell has a source; Brennan alone = one source'))

# today's OMi rows: OMI-DRAWING (the 14 drawn parts) / OMI (6410 rows filled from OMi's length answer + website file)
if os.path.exists(OMI_CH):
    seen = set()
    for c in csv.DictReader(open(OMI_CH, encoding='utf-8')):
        t = c['table'].replace('DesignTable_', '').replace('.xlsx', '')
        key = (t, c['part_number'])
        if key in seen or c['column'] != '$PRP@DataStatus':
            continue
        if not os.path.exists(os.path.join(DT, f'DesignTable_{t}.xlsx')):
            continue                         # stage without that table (e.g. the ORFS stage)
        seen.add(key)
        _, _, ws, hdr = sheet(t)
        r = rowof(ws, c['part_number'])
        v = 'OMI-DRAWING' if t in ('1405', '1502', '1504', '1602') else 'OMI'
        cell = ws.cell(r, hdr[COL])
        if cell.value in (None, ''):
            cell.value = v
            log.append((f'DesignTable_{t}.xlsx', c['part_number'], r, COL, cell.coordinate, '', v, 'sizes from OMi\'s 2026-10-18 answer (omi_answers_20261018)'))

for t, (p, wb, ws, hdr) in books.items():
    wb.save(p)
out = os.path.join(stage, 'CHANGES.csv')
new = not os.path.exists(out)
with open(out, 'a', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    if new:
        w.writerow(['table', 'part_number', 'excel_row', 'column', 'cell', 'old', 'new', 'source'])
    w.writerows(log)
print(len(log), 'log lines;', len(per_row), 'rows touched;', sum(1 for k in per_row if any(c['kind'] == 'brennan-dim' for c in per_row[k])), 'marked BRENNAN;',
      len(ready & set(per_row)), 'set ASSUMED')
