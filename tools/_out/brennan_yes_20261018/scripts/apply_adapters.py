"""apply_adapters.py <stage_dir>: wave 3 adapters into a stage (2026-10-18).
  1. wave3/adapters/edits_77.csv          - the 37 wrongly-DROP rows (2026-10-16 change set rebased on today's stage; old values asserted)
  2. wave3/adapters/edits_5406P_alias.csv - 5406-08P / -12P = OMi 5406-P-08 / -12 (A4 alias yes)
  3. wave3/adapters/newrows.csv           - 31 new rows (Brennan rows marked $PRP@DimSource BRENNAN + AxialSource prefix; -LN copies of plain rows)
  4. wave3/adapters/pending_omi_A19.csv   - user 2026-10-18: "Yes, apply the 7 pending OMi values ... like A19, and set 6500-08-06-FG to
                                            ASSUMED" (9 cells on 5 rows + the status)
Same edit logic as wave3/adapters/scripts/test.py apply() (which the agent tested), plus the marking and step 4. Logs to <stage>/CHANGES.csv."""
import os, sys, csv, re, collections
import openpyxl

A = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff\tools\_out\brennan_yes_20261018\wave3\adapters'
DT = os.path.join(sys.argv[1], '04_DESIGN_TABLES')
TAG = '[BRENNAN-ONLY 2026-10-18] '


def num_or_str(s):
    s = str(s)
    return float(s) if re.fullmatch(r'-?\d*\.\d+|-?\d+', s.strip()) else s


def same(a, b):
    if a in (None, '') and b in (None, ''):
        return True
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


books, log = {}, []


def bk(t):
    if t not in books:
        p = os.path.join(DT, f'DesignTable_{t}.xlsx')
        wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
        books[t] = [p, wb, ws, {c.value: c.column for c in ws[2] if c.value is not None}]
    return books[t]


def col(t, h):
    p, wb, ws, hdr = bk(t)
    if h not in hdr:
        c = max(ws.max_column, max(hdr.values())) + 1
        ws.cell(2, c, h); hdr[h] = c
        log.append((f'DesignTable_{t}.xlsx', '(header)', 2, h, ws.cell(2, c).coordinate, '', h, 'new column'))
    return hdr[h]


def row(t, cfg):
    ws = bk(t)[2]
    rr = [c.row for c in ws['A'] if c.value is not None and str(c.value).strip() == cfg]
    assert len(rr) == 1, (t, cfg, rr)
    return rr[0]


for f in ('edits_77.csv', 'edits_5406P_alias.csv'):
    for e in csv.DictReader(open(os.path.join(A, f), encoding='utf-8')):
        t = e['table']; r = row(t, e['config'])
        h = e['column'].replace(' (append)', '')
        cell = bk(t)[2].cell(r, col(t, h))
        assert same(cell.value, e['old']), (f, e, cell.value)
        old = cell.value
        if e['column'].endswith('(append)'):
            cell.value = (str(cell.value) + ' | ' if cell.value else '') + e['new']
        else:
            cell.value = num_or_str(e['new']) if not h.startswith('$') else e['new']
        log.append((f'DesignTable_{t}.xlsx', e['config'], r, h, cell.coordinate, old, cell.value, f"{f}: {e['why']}"))

new = collections.OrderedDict()
for x in csv.DictReader(open(os.path.join(A, 'newrows.csv'), encoding='utf-8')):
    new.setdefault((x['table'], x['config']), []).append(x)
for (t, cfg), cells in new.items():
    ws = bk(t)[2]
    assert not [c for c in ws['A'] if c.value is not None and str(c.value).strip() == cfg], (t, cfg, 'exists')
    r = max(c.row for c in ws['A'] if c.value is not None) + 1
    ws.cell(r, 1, cfg)
    brennan = any(x['kind'] == 'brennan-dim' and x['value'] != '' for x in cells)
    for x in cells:
        if x['value'] == '':
            continue
        v = num_or_str(x['value']) if not x['column'].startswith('$') else x['value']
        if x['column'] == '$PRP@AxialSource' and brennan and not str(v).startswith(TAG):
            v = TAG + str(v)
        ws.cell(r, col(t, x['column']), v)
    if brennan:
        ws.cell(r, col(t, '$PRP@DimSource'), 'BRENNAN')
    log.append((f'DesignTable_{t}.xlsx', cfg, r, '(new row)', f'A{r}', '', f"{sum(1 for x in cells if x['value'] != '')} cells" + (' BRENNAN' if brennan else ''),
                'newrows.csv (wave 3 adapters)'))

for x in csv.DictReader(open(os.path.join(A, 'pending_omi_A19.csv'), encoding='utf-8')):
    t = x['table']; r = row(t, x['config']); h = x['column']
    cell = bk(t)[2].cell(r, col(t, h))
    assert same(cell.value, x['table_now'].split(' ')[0]), (x, cell.value)
    old = cell.value
    cell.value = num_or_str(x['proposed']) if not h.startswith('$') else x['proposed']
    cm = bk(t)[2].cell(r, col(t, '$COMMENT'))
    cm.value = (cm.value + ' | ' if cm.value else '') + f"A19-style (user 2026-10-18 'yes, apply the 7 pending OMi values'): {h} {old} -> {cell.value}"
    log.append((f'DesignTable_{t}.xlsx', x['config'], r, h, cell.coordinate, old, cell.value, f"user 2026-10-18 pending OMi values: {x['omi_cell']} {x['why'][:80]}"))

for p, wb, ws, hdr in books.values():
    wb.save(p)
with open(os.path.join(sys.argv[1], 'CHANGES.csv'), 'a', newline='', encoding='utf-8') as f:
    csv.writer(f).writerows(log)
print(len(log), 'log lines;', len(new), 'new rows')
