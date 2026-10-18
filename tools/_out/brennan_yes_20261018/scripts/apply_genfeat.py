"""apply_genfeat.py <stage_dir>: data for the rows the 5 generator features build (genfeat_20261018, merged 2026-10-18).
- built rows = rows_result.csv result 'built' of hoseb / perlegsw / angled / flange / straights;
- their cells from: wave2/{straight,tee_elbow,flange_hose}/cells.csv, wave1/w1brennan_cells.csv, genfeat hoseb/cells_extra.csv,
  genfeat straights/test/cells_feature.csv (first source wins per cell; only EMPTY stage cells are filled; a non-empty different value
  is reported as a conflict and left);
- angled's edits test/root/CHANGES_angled.csv (A15 'look at Brennan': 5503 drills, PORTDEPTH, 1504-04-04-FG thread major, 6902 rename),
  old values asserted, DataStatus lines ignored;
- every built row that is MISSING -> ASSUMED; rows that took a brennan-dim cell -> $PRP@DimSource BRENNAN + AxialSource prefix.
Logs to <stage>/CHANGES.csv and prints conflicts."""
import os, sys, csv, re, collections
import openpyxl

ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
BY = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018')
GF = os.path.join(ROOT, r'tools\_out\genfeat_20261018')
DT = os.path.join(sys.argv[1], '04_DESIGN_TABLES')
TAG = '[BRENNAN-ONLY 2026-10-18] '


def num_or_str(s):
    s = str(s)
    return float(s) if re.fullmatch(r'-?\d*\.\d+|-?\d+', s.strip()) else s


def same(a, b):
    if a in (None, '') and b in (None, ''):
        return True
    try:
        return abs(float(a) - float(b)) < 1e-6
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


books, log, conflicts = {}, [], []


def bk(t):
    if t not in books:
        p = os.path.join(DT, f'DesignTable_{t}.xlsx')
        if not os.path.exists(p):
            return None
        wb = openpyxl.load_workbook(p); ws = wb.worksheets[0]
        books[t] = [p, wb, ws, {c.value: c.column for c in ws[2] if c.value is not None}]
    return books[t]


def locate(t, cfg):
    """(table, row) of cfg: the named table, else the table named by the leading digits (1700RC -> 1700, 2501LL -> 2501LL / 2501)"""
    for cand in (t, re.sub(r'[A-Z]+$', '', t), t.rstrip('T')):
        b = bk(cand)
        if b:
            norm = lambda s: str(s).strip().replace(' ', '-')     # Air-Way names "4601-06-06 NWO-FG" = OMi "4601-06-06-NWO-FG"
            rr = [c.row for c in b[2]['A'] if c.value is not None and str(c.value).strip() == cfg]
            if not rr:
                rr = [c.row for c in b[2]['A'] if c.value is not None and norm(c.value) == norm(cfg)]
            if len(rr) == 1:
                return cand, rr[0]
    return None, None


def col(t, h):
    p, wb, ws, hdr = bk(t)
    if h not in hdr:
        c = max(ws.max_column, max(hdr.values())) + 1
        ws.cell(2, c, h); hdr[h] = c
        log.append((f'DesignTable_{t}.xlsx', '(header)', 2, h, ws.cell(2, c).coordinate, '', h, 'new column'))
    return hdr[h]


built = set()
for f in ('hoseb', 'perlegsw', 'angled', 'flange', 'straights'):
    for x in csv.DictReader(open(os.path.join(GF, f, 'rows_result.csv'), encoding='utf-8-sig')):
        if x['result'].strip() == 'built':
            built.add((x['table'], x['config']))

# angled edits first (renames change column A)
for e in csv.DictReader(open(os.path.join(GF, r'angled\test\root\CHANGES_angled.csv'), encoding='utf-8-sig')):
    if e['column'] == '$PRP@DataStatus' or e['config'] == '(header)':
        continue                                    # status: set below; new columns: col() adds them
    t, r = locate(e['table'], e['config'])
    if t is None:                                   # a rename: the row still has its old name
        t, r = locate(e['table'], e['old'])
    assert t, e
    h = e['column']
    if h.startswith('(config name'):               # rename: column A + $PRP@PartNo
        p_, wb_, ws_, hdr_ = bk(t)
        for cc in (ws_.cell(r, 1), ws_.cell(r, hdr_['$PRP@PartNo'])):
            if cc.value == e['old']:
                cc.value = e['new']; log.append((f'DesignTable_{t}.xlsx', e['new'], r, 'rename', cc.coordinate, e['old'], e['new'], 'A15 (angled): OMi part-number form'))
        continue
    c = bk(t)[2].cell(r, col(t, h))
    if e['old'] == '(appended)':                   # note appended to a text cell
        if e['new'] not in (c.value or ''):
            old = c.value; c.value = (c.value + ' | ' if c.value else '') + e['new']
            log.append((f'DesignTable_{t}.xlsx', e['config'], r, h, c.coordinate, old, c.value, 'A15 (angled): note appended'))
        continue
    if not same(c.value, e['old']):
        conflicts.append(('angled edit', t, e['config'], h, c.value, e['old'], e['new'])); continue
    old = c.value; c.value = num_or_str(e['new']) if not h.startswith('$') else e['new']
    log.append((f'DesignTable_{t}.xlsx', e['config'], r, h, c.coordinate, old, c.value, f"A15 (angled feature): {e['why'][:200]}"))
    if h == '$PRP@PartNo':                          # keep column A in step with a rename
        a = bk(t)[2].cell(r, 1)
        if a.value == old:
            a.value = e['new']; log.append((f'DesignTable_{t}.xlsx', e['config'], r, 'A', a.coordinate, old, e['new'], 'rename (A15 angled)'))

SRC = [os.path.join(BY, r'wave2\straight\cells.csv'), os.path.join(BY, r'wave2\tee_elbow\cells.csv'), os.path.join(BY, r'wave2\flange_hose\cells.csv'),
       os.path.join(BY, r'wave1\w1brennan_cells.csv'), os.path.join(GF, r'hoseb\cells_extra.csv'), os.path.join(GF, r'straights\test\cells_feature.csv')]
done, brennan_rows, touched = set(), set(), set()
for f in SRC:
    for c in csv.DictReader(open(f, encoding='utf-8-sig')):
        key = (c['table'], c['config'])
        h = c['column']
        if key not in built or h in ('$PRP@DataStatus', '$PRP@DimSource', '$PRP@AxialSource') or (key, h) in done:
            continue
        t, r = locate(*key)
        if t is None:
            conflicts.append(('row not found', c['table'], c['config'], h, '', '', '')); continue
        cell = bk(t)[2].cell(r, col(t, h))
        done.add((key, h))
        v = c['value'].split(' | before: ')[0]
        if h == '$COMMENT':
            if v and v not in (cell.value or ''):
                old = cell.value; cell.value = (cell.value + ' | ' if cell.value else '') + v
                log.append((f'DesignTable_{t}.xlsx', key[1], r, h, cell.coordinate, old, cell.value, f'note ({os.path.basename(os.path.dirname(f))})'))
            continue
        if cell.value not in (None, ''):
            if not same(cell.value, v):
                conflicts.append(('non-empty', t, key[1], h, cell.value, '', v))
            continue
        cell.value = num_or_str(v) if not h.startswith('$') else v
        log.append((f'DesignTable_{t}.xlsx', key[1], r, h, cell.coordinate, '', cell.value, f"[{c['kind']}] {c['source'][:300]}"))
        touched.add((t, r, key[1]))
        if c['kind'] == 'brennan-dim':
            brennan_rows.add((t, r, key[1]))

n_ass = 0
for key in sorted(built):
    t, r = locate(*key)
    if t is None:
        conflicts.append(('built row not found', key[0], key[1], '', '', '', '')); continue
    p, wb, ws, hdr = bk(t)
    s = ws.cell(r, hdr['$PRP@DataStatus'])
    if s.value in (None, '', 'MISSING'):
        old = s.value; s.value = 'ASSUMED'; n_ass += 1
        log.append((f'DesignTable_{t}.xlsx', key[1], r, '$PRP@DataStatus', s.coordinate, old or '', 'ASSUMED', 'built by the merged generator features 2026-10-18; every driving cell sourced'))
for t, r, cfg in brennan_rows:
    p, wb, ws, hdr = bk(t)
    d = ws.cell(r, col(t, '$PRP@DimSource'))
    if d.value != 'BRENNAN':
        d.value = 'BRENNAN'
        ax = ws.cell(r, hdr['$PRP@AxialSource'])
        if not str(ax.value or '').startswith(TAG):
            ax.value = TAG + (ax.value or '') + ' -> 2026-10-18 Brennan allowed (user: "yes use brennan but still differentiate"); sources per cell in CHANGES.csv'
        log.append((f'DesignTable_{t}.xlsx', cfg, r, '$PRP@DimSource', d.coordinate, '', 'BRENNAN', 'row takes sizes from Brennan'))

for p, wb, ws, hdr in books.values():
    wb.save(p)
with open(os.path.join(sys.argv[1], 'CHANGES.csv'), 'a', newline='', encoding='utf-8') as f:
    csv.writer(f).writerows(log)
print('built rows', len(built), '| cells filled', sum(1 for l in log if l[3] not in ('$PRP@DataStatus', '$PRP@DimSource', '$COMMENT')),
      '| set ASSUMED', n_ass, '| BRENNAN rows', len(brennan_rows), '| conflicts', len(conflicts))
for c in conflicts:
    print('  CONFLICT', c)
