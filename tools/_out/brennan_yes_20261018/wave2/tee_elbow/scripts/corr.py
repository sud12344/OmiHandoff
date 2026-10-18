import openpyxl, csv, sys, json, re, collections
R = r'C:/Users/sudha/OneDrive/Desktop/OMi_Handoff/'
BASE = R + 'tools/_out/brennan_yes_20261018/base_root/04_DESIGN_TABLES/'
d = json.load(open(R + 'tools/_out/brennan_catalog_20261001/catalog_verified.json'))
cat = {}
for t in d['tables']:
    for rr in t['rows']:
        cat.setdefault(rr['pn'], (t['table_id'], rr['values']))
def num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip()
    if re.fullmatch(r'-?\d*\.?\d+', s): return float(s)
    return None
rows = list(csv.DictReader(open(R + 'tools/_out/brennan_yes_20261018/wave2/rows_tee_elbow.csv')))
tabs = sys.argv[1:] or sorted({r['table'] for r in rows}, key=str)
TWIN = __import__('os').environ.get('TWIN') == '1'
TOL = float(__import__('os').environ.get('TOL', '0.011'))
for t in tabs:
    ws = openpyxl.load_workbook(BASE + f'DesignTable_{t}.xlsx').active
    hdr = [c.value for c in ws[2]]
    tgt = {r['config'] for r in rows if r['table'] == t}
    stats = collections.defaultdict(lambda: [0, 0, []])
    npair = 0; tids = set()
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None or row[0] in tgt: continue
        pn = str(row[0]).strip().replace(' ', '-')
        if TWIN and not pn.endswith('FG'): pn = pn + '-FG' if not pn.endswith('-O') else pn
        if TWIN and pn.endswith('FG') and str(row[0]).strip().replace(' ', '-').endswith('FG'): continue
        if pn not in cat: continue
        dd = dict(zip(hdr, row)); tid, cv = cat[pn]; tids.add(tid)
        if not any(num(dd[h]) is not None for h in hdr[1:] if h and not h.startswith('$')): continue
        npair += 1
        for h in hdr[1:]:
            if not h or h.startswith('$'): continue
            a = num(dd[h])
            if a is None: continue
            for k, v in cv.items():
                b = num(v)
                if b is None: continue
                s = stats[(h, k)]; s[1] += 1
                if abs(a - b) <= TOL: s[0] += 1
                else: s[2].append((row[0], a, b))
    print(f'===== {t}: paired rows with values {npair}, catalog tables {sorted(tids)}')
    for h in hdr[1:]:
        if not h or h.startswith('$'): continue
        c = sorted([(s[0], s[1], k) for (hh, k), s in stats.items() if hh == h], reverse=True)[:3]
        print(f'  {h:18}', ' | '.join(f'{k}: {m}/{n}' for m, n, k in c))
