import openpyxl, csv, sys, json
R = r'C:/Users/sudha/OneDrive/Desktop/OMi_Handoff/'
BASE = R + 'tools/_out/brennan_yes_20261018/base_root/04_DESIGN_TABLES/'
rows = list(csv.DictReader(open(R + 'tools/_out/brennan_yes_20261018/wave2/rows_tee_elbow.csv')))
tabs = sys.argv[1:] or sorted({r['table'] for r in rows})
for t in tabs:
    ws = openpyxl.load_workbook(BASE + f'DesignTable_{t}.xlsx').active
    hdr = [c.value for c in ws[2]]
    tgt = {r['config'] for r in rows if r['table'] == t}
    n = sum(1 for r in ws.iter_rows(min_row=3, values_only=True) if r[0])
    built = sum(1 for r in ws.iter_rows(min_row=3, values_only=True) if r[0] and dict(zip(hdr, r)).get('$PRP@ScopeAction') == 'BUILD' and dict(zip(hdr, r)).get('$PRP@DataStatus') in ('VERIFIED', 'ASSUMED'))
    print(f'===== {t} ({n} rows, {built} BUILD+data) HDR', hdr[1:])
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] in tgt:
            d = dict(zip(hdr, r))
            print('  ', r[0], {k: (str(v)[:70]) for k, v in d.items() if v not in (None, '') and k not in ('$PRP@PartNo', '$PRP@OMiSKU', hdr[0])})
