import openpyxl, csv, sys, json, re
R = r'C:/Users/sudha/OneDrive/Desktop/OMi_Handoff/'
BASE = R + 'tools/_out/brennan_yes_20261018/base_root/04_DESIGN_TABLES/'
d = json.load(open(R + 'tools/_out/brennan_catalog_20261001/catalog_verified.json'))
cat = {}
for t in d['tables']:
    for rr in t['rows']:
        cat.setdefault(rr['pn'], (t['table_id'], rr['values']))
t = sys.argv[1]; cols = sys.argv[2].split(',') if len(sys.argv) > 2 else None
ws = openpyxl.load_workbook(BASE + f'DesignTable_{t}.xlsx').active
hdr = [c.value for c in ws[2]]
for row in ws.iter_rows(min_row=3, values_only=True):
    if row[0] is None or row[0] not in cat: continue
    dd = dict(zip(hdr, row))
    tv = {h: dd[h] for h in hdr[1:] if h and not h.startswith('$') and (cols is None or h in cols)}
    print(row[0], dd.get('$PRP@DataStatus'), dd.get('$PRP@ScopeAction'), '| T:', json.dumps(tv), '| B:', json.dumps({k: v for k, v in cat[row[0]][1].items() if re.match(r'^[\d.]+$', v)}))
