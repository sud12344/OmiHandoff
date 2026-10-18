import openpyxl, csv, sys, json
R = r'C:/Users/sudha/OneDrive/Desktop/OMi_Handoff/'
BASE = R + 'tools/_out/brennan_yes_20261018/base_root/04_DESIGN_TABLES/'
rows = list(csv.DictReader(open(R + 'tools/_out/brennan_yes_20261018/wave2/rows_tee_elbow.csv')))
tabs = sys.argv[1:] or sorted({r['table'] for r in rows})
SKIP = {'$PRP@AxialSource', '$COMMENT', '$PRP@Description'}
for t in tabs:
    ws = openpyxl.load_workbook(BASE + f'DesignTable_{t}.xlsx').active
    hdr = [c.value for c in ws[2]]
    print('=====', t, 'rows', ws.max_row - 2, 'row1:', [c.value for c in ws[1]][:3])
    print('HDR', hdr)
    tgt = {r['config'] for r in rows if r['table'] == t}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None: continue
        d = dict(zip(hdr, row))
        st = d.get('$PRP@DataStatus'); sa = d.get('$PRP@ScopeAction') or d.get('ScopeAction')
        mark = '**' if row[0] in tgt else '  '
        vals = {k: v for k, v in d.items() if v not in (None, '') and k not in SKIP and k != hdr[0]}
        print(mark, row[0], json.dumps(vals, default=str)[:900])
