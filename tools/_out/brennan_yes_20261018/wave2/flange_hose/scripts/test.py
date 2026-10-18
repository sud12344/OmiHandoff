"""test.py - flange_hose TEST (2026-10-18).
  test/root_untouched/04_DESIGN_TABLES  = copy of today's base (control)
  test/root/04_DESIGN_TABLES            = base + cells.csv applied with the main session's apply_brennan.py (ready rows -> ASSUMED; none is ready)
  test/root_B/04_DESIGN_TABLES          = root + DataStatus ASSUMED on every DATA COMPLETE row (informative only: what the rows would give
                                          once their capability is merged; NOT a proposal)
Regenerates EVERY root of the batch-16 generators (puckfix_20261009/tools_new; python -I, one process per root) on untouched and root;
then, informative (scenario B): batch-16 gen_straight2 1700 on root_B, and the unmerged wave2_impl_20261009 HOSEB gen_tee on root_B for
4501 / 4503 / 4601 / 4603 (4501 / 4503 injected from SPEC_WAITING with the hose leg reading DH@BORE, the decided Q-HOSEB-BORE route).
    python -I test.py [apply|regen|b|compare|all]"""
import os, sys, csv, json, shutil, subprocess, re
import openpyxl

ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
BY = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018')
W = os.path.join(BY, r'wave2\flange_hose')
T = os.path.join(W, 'test')
BASE_DT = os.path.join(BY, r'base_root\04_DESIGN_TABLES')
TOOLS = os.path.join(ROOT, r'tools\_out\puckfix_20261009\tools_new')
TOOLS_W2 = os.path.join(ROOT, r'tools\_out\wave2_impl_20261009\tools_new')
APPLY = os.path.join(BY, r'scripts\apply_brennan.py')
SCEN = {'untouched': os.path.join(T, 'root_untouched'), 'A': os.path.join(T, 'root'), 'B': os.path.join(T, 'root_B')}
MY_TABLES = ['1700', '1703', '1704', '1803', '1804', '4501', '4503', '4601', '4603']


def run(cmd, env=None, log=None):
    p = subprocess.run(cmd, env=env, capture_output=True, text=True, encoding='utf-8', errors='replace')
    if log:
        open(log, 'w', encoding='utf-8').write(p.stdout + ('\n--- STDERR ---\n' + p.stderr if p.stderr else ''))
    return p


def apply():
    for s in SCEN.values():
        assert not os.path.exists(s), s
    shutil.copytree(BASE_DT, os.path.join(SCEN['untouched'], '04_DESIGN_TABLES'))
    p = run([sys.executable, '-I', APPLY, SCEN['A'], os.path.join(W, 'cells.csv'), '--rows', os.path.join(W, 'rows.csv')], log=os.path.join(T, 'apply_A_stdout.txt'))
    print(p.stdout, p.stderr)
    assert p.returncode == 0
    shutil.copytree(SCEN['A'], SCEN['B'])
    complete = [r for r in csv.DictReader(open(os.path.join(W, 'rows.csv'), encoding='utf-8')) if r['notes'].startswith('DATA COMPLETE')]
    log = []
    for t in MY_TABLES:
        p = os.path.join(SCEN['B'], '04_DESIGN_TABLES', f'DesignTable_{t}.xlsx')
        wb = openpyxl.load_workbook(p); ws = wb.active
        hdr = {ws.cell(2, j).value: j for j in range(1, ws.max_column + 1) if ws.cell(2, j).value}
        n = 0
        for r in complete:
            if r['table'] != t:
                continue
            i = [c.row for c in ws['A'] if c.value == r['config']][0]
            ws.cell(i, hdr['$PRP@DataStatus'], 'ASSUMED'); n += 1
            log.append(f'{t} {r["config"]} DataStatus ASSUMED (scenario B only)')
        if n:
            wb.save(p)
    open(os.path.join(SCEN['B'], 'B_log.txt'), 'w', encoding='utf-8').write('\n'.join(log) + '\n')


def regen(scen):
    out = os.path.join(T, f'gen_{scen}')
    assert not os.path.exists(out), out
    os.makedirs(os.path.join(out, 'logs'))
    env = dict(os.environ, OMI_ROOT=SCEN[scen], GS_OUT=os.path.join(out, 'straight'), GT_OUT=os.path.join(out, 'tee'), PYTHONDONTWRITEBYTECODE='1')
    code = ("import sys,os,json; t=sys.argv[1]; sys.path[:0]=[os.path.join(t,'generators'),t]; import gen_straight2 as G, gen_tee as T; "
            "print(json.dumps(dict(straight=list(G.SPEC), tee=list(T.SPEC))))")
    k = run([sys.executable, '-I', '-B', '-c', code, TOOLS], env=dict(env, GS_OUT=os.path.join(out, '_keys'), GT_OUT=os.path.join(out, '_keys')))
    assert k.returncode == 0, k.stderr
    roots = json.loads(k.stdout.strip().splitlines()[-1])
    json.dump(roots, open(os.path.join(out, 'roots.json'), 'w'))
    status = {}
    for sub, gen in (('straight', 'gen_straight2.py'), ('tee', 'gen_tee.py')):
        os.makedirs(os.path.join(out, sub), exist_ok=True)
        sj = os.path.join(out, sub, 'gen_straight2_summary.json' if sub == 'straight' else 'gen_tee_summary.json')
        for r in roots[sub]:
            if os.path.exists(sj): os.remove(sj)
            p = run([sys.executable, '-I', '-B', os.path.join(TOOLS, 'generators', gen), r], env=env, log=os.path.join(out, 'logs', f'{sub}_{r}.txt'))
            if os.path.exists(sj): os.replace(sj, os.path.join(out, sub, f'summary_{r}.json'))
            last = (p.stderr.strip().splitlines() or [''])[-1][:300]
            status[f'{sub}:{r}'] = dict(rc=p.returncode, err=last if p.returncode else '', out=(p.stdout.strip().splitlines() or [''])[-1][:300])
            print(scen, sub, r, 'rc', p.returncode, last if p.returncode else '', flush=True)
    json.dump(status, open(os.path.join(T, f'status_{scen}.json'), 'w'), indent=1)


RUNNER = os.path.join(W, 'scripts', 'w2_runner.py')


def scen_b():
    out = os.path.join(T, 'gen_B')
    assert not os.path.exists(out), out
    os.makedirs(os.path.join(out, 'logs'))
    env = dict(os.environ, OMI_ROOT=SCEN['B'], GS_OUT=os.path.join(out, 'straight'), GT_OUT=os.path.join(out, 'tee'), PYTHONDONTWRITEBYTECODE='1')
    p = run([sys.executable, '-I', '-B', os.path.join(TOOLS, 'generators', 'gen_straight2.py'), '1700'], env=env, log=os.path.join(out, 'logs', 'straight_1700_batch16.txt'))
    print('B 1700 (batch-16, full hex - informative)', p.returncode)
    env2 = dict(env, GT_OUT=os.path.join(out, 'tee_w2'))
    inj = {'4601': ['4601', {}], '4603': ['4603', {}],
           '4501': ['4501', {'hose_bore': 'own', 'legs': [['HOSEB', 'L4@DRIVE', 'DH@BORE', None], ['NPTFM', 'N@DRIVE', 'D1@BORE', 'THDMAJ_A@DRIVE']]}],
           '4503': ['4503', {'hose_bore': 'own', 'legs': [['NPTFM', 'N2@DRIVE', 'D1@BORE', 'THDMAJ_A@DRIVE'], ['HOSEB', 'L5@DRIVE', 'DH@BORE', None]]}]}
    for root, v in inj.items():
        if os.path.exists(os.path.join(out, 'tee_w2', 'gen_tee_summary.json')):
            os.remove(os.path.join(out, 'tee_w2', 'gen_tee_summary.json'))
        p = run([sys.executable, '-I', '-B', RUNNER, TOOLS_W2, json.dumps({root: v})], env=env2, log=os.path.join(out, 'logs', f'tee_w2_{root}.txt'))
        sj = os.path.join(out, 'tee_w2', 'gen_tee_summary.json')
        if os.path.exists(sj): os.replace(sj, os.path.join(out, 'tee_w2', f'summary_{root}.json'))
        print('B', root, '(wave2_impl HOSEB, informative)', p.returncode, (p.stderr.strip().splitlines() or [''])[-1][:200])


def sw2(p):
    ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
    v = list(ws.iter_rows(values_only=True)); h = v[0]
    return h, {r[0]: dict(zip(h, r)) for r in v[1:] if r[0]}


def macro(p):
    return [l for l in open(p, encoding='utf-8', errors='replace').read().replace('\r\n', '\n').split('\n')
            if not re.search(r'[Gg]enerated|20\d\d-\d\d-\d\d', l)]


def compare():
    rep = []
    st = {s: json.load(open(os.path.join(T, f'status_{s}.json'))) for s in ('untouched', 'A')}
    target = {r['config'] for r in csv.DictReader(open(os.path.join(W, 'rows.csv'), encoding='utf-8'))}
    rep.append('===== scenario A (cells.csv applied) vs untouched base - every root of the batch-16 generators')
    nsame = 0
    for key in sorted(st['untouched']):
        sub, r = key.split(':')
        u, a = st['untouched'][key], st['A'][key]
        gu, ga = os.path.join(T, 'gen_untouched', sub), os.path.join(T, 'gen_A', sub)
        mu, ma = os.path.join(gu, f'Build_{r}_MASTER.bas'), os.path.join(ga, f'Build_{r}_MASTER.bas')
        tu, ta = os.path.join(gu, f'DesignTable_{r}_SW2.xlsx'), os.path.join(ga, f'DesignTable_{r}_SW2.xlsx')
        line = f'{key:22} rc {u["rc"]}->{a["rc"]}'
        if u['rc'] != a['rc']:
            line += f' | untouched: {u["err"] or u["out"]} | A: {a["err"] or a["out"]}'
        if not os.path.exists(tu) and not os.path.exists(ta):
            rep.append(line + ' | no SW2 in either'); continue
        hu, ru = sw2(tu) if os.path.exists(tu) else ((), {})
        ha, ra = sw2(ta) if os.path.exists(ta) else ((), {})
        changed = {}
        for cfg, row in ru.items():
            if cfg not in ra:
                changed[cfg] = 'LOST'; continue
            d = [c for c in set(row) | set(ra[cfg]) if row.get(c) != ra[cfg].get(c)]
            if d: changed[cfg] = sorted(d)
        new = [c for c in ra if c not in ru]
        same = os.path.exists(mu) and os.path.exists(ma) and macro(mu) == macro(ma)
        if same and not changed and not new and u['rc'] == a['rc'] and set(hu) == set(ha):
            nsame += 1
        line += f' | configs {len(ru)}->{len(ra)} new {len(new)} old-changed {len(changed)} | macro ' + ('identical' if same else 'DIFFERS')
        rep.append(line)
        for c in new:
            rep.append(f'      NEW {c}' + (' (target row)' if c in target else ' (NOT a target row!)'))
        for c, d in changed.items():
            rep.append(f'      OLD CONFIG CHANGED {c}: {d}')
    rep.insert(1, f'roots identical (macro apart from the date line, every SW2 cell, return code): {nsame} of {len(st["untouched"])}')
    # 1700 in A: why the 3 rows are not configurations
    for s, lg in (('A', os.path.join(T, 'gen_A', 'logs', 'straight_1700.txt')), ('B', os.path.join(T, 'gen_B', 'logs', 'straight_1700_batch16.txt'))):
        if os.path.exists(lg):
            txt = open(lg, encoding='utf-8').read()
            rep.append(f'\n===== 1700 log, scenario {s} (lines naming the 3 Brennan rows)')
            rep += ['  ' + l for l in txt.splitlines() if re.search(r'1700-(20-16|24-20|32-24)', l)][:20]
            rep.append('  last: ' + (txt.strip().splitlines() or [''])[-1][:300])
    # scenario B HOSEB (wave2_impl)
    rep.append('\n===== scenario B (informative): unmerged wave2_impl HOSEB gen_tee on root_B (DATA COMPLETE rows set ASSUMED)')
    for root in ('4501', '4503', '4601', '4603'):
        lg = os.path.join(T, 'gen_B', 'logs', f'tee_w2_{root}.txt')
        sm = os.path.join(T, 'gen_B', 'tee_w2', f'summary_{root}.json')
        t2 = os.path.join(T, 'gen_B', 'tee_w2', f'DesignTable_{root}_SW2.xlsx')
        built = sorted(sw2(t2)[1]) if os.path.exists(t2) else []
        rep.append(f'  {root}: configurations {len(built)}: ' + ', '.join(c + (' *' if c in target else '') for c in built))
        if os.path.exists(sm):
            js = json.load(open(sm, encoding='utf-8'))
            js = js.get(root, js)
            for b in js.get('blocked', []) if isinstance(js, dict) else []:
                rep.append(f'      BLOCKED {b}'[:400])
        if os.path.exists(lg):
            txt = open(lg, encoding='utf-8').read()
            rep += ['      ' + l[:300] for l in txt.splitlines() if 'BLOCK' in l.upper() or 'Traceback' in l or 'Error' in l][:30]
    open(os.path.join(T, 'compare.txt'), 'w', encoding='utf-8').write('\n'.join(rep) + '\n')
    print('\n'.join(rep[:3])); print('\n'.join(l for l in rep if 'DIFFERS' in l or 'CHANGED' in l or 'NEW' in l or 'rc 0->1' in l or 'rc 1->0' in l))
    print('\n'.join(rep[rep.index(next(l for l in rep if l.startswith('\n===== 1700'))):]))


if __name__ == '__main__':
    what = sys.argv[1] if len(sys.argv) > 1 else 'all'
    if what in ('apply', 'all'):
        apply()
    if what in ('regen', 'all'):
        regen('untouched'); regen('A')
    if what in ('b', 'all'):
        scen_b()
    if what in ('compare', 'all'):
        compare()
