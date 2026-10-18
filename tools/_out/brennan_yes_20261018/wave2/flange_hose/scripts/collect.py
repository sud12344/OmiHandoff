"""collect.py - wave 2, group flange_hose (2026-10-18): builds ../cells.csv and ../rows.csv for the 55 rows of
wave2/rows_flange_hose.csv (tables 1700 1703 1704 1803 1804 4501 4503 4601 4603). Read-only on every input.

Sources (in order; Brennan only fills cells that are EMPTY in today's base and have no OMi / Air-Way value):
  - Brennan catalog parse tools/_out/brennan_catalog_20261001/catalog_verified.json (page-verified tables only)
  - own Brennan STEPs: evidence/pilot_sections.json (scripts/measure_pilots.py), FLANGE62 flange62_straight.jsonl (validated flange62.py)
    + evidence/1700_bores_prof.txt (tools/measure/prof.py)
  - end libraries (Q-TRANSFER-ENDLIB yes 2026-10-08): gen_straight2 JIC / ORB / ms51500.NPTF of the batch-16 tools, the code 61 head OD
    by flange size as built in the 1700 master, the hose bead dia as built in the 4404 / 4604 masters, the code 62 -08 head of the
    FLANGE62 staging (FLG62_MEAS[8])
Run: python -I collect.py  (imports the batch-16 gen_straight2 on a scratch OMI_ROOT = base_root)."""
import os, sys, re, csv, json, collections
from fractions import Fraction
import openpyxl

ROOT = r'C:\Users\sudha\OneDrive\Desktop\OMi_Handoff'
BY = os.path.join(ROOT, r'tools\_out\brennan_yes_20261018')
W = os.path.join(BY, r'wave2\flange_hose')
EV = os.path.join(W, 'evidence')
BASE = os.path.join(BY, r'base_root\04_DESIGN_TABLES')
CATF = os.path.join(ROOT, r'tools\_out\brennan_catalog_20261001\catalog_verified.json')
F62 = os.path.join(ROOT, r'tools\_out\newmasters_20261002\work\FLANGE62\evidence\measure\flange62_straight.jsonl')
TOOLS = os.path.join(ROOT, r'tools\_out\puckfix_20261009\tools_new')
STEPDIR = os.path.join(ROOT, '09_STEP_FILES')
DEC = '[Brennan = yes 2026-10-18]'
ENDLIB = '[Q-TRANSFER-ENDLIB yes 2026-10-08]'

os.environ.setdefault('OMI_ROOT', os.path.join(BY, 'base_root'))
os.environ.setdefault('GS_OUT', os.path.join(W, r'test\_import_scratch\straight'))
os.environ.setdefault('GT_OUT', os.path.join(W, r'test\_import_scratch\tee'))
sys.path[:0] = [os.path.join(TOOLS, 'generators'), TOOLS]
import gen_straight2 as G   # noqa: E402
import ms51500 as MS        # noqa: E402

# ------------------------------------------------------------------------------------------------------------------ inputs
ROWS = list(csv.DictReader(open(os.path.join(BY, r'wave2\rows_flange_hose.csv'), encoding='utf-8-sig')))
CAT = json.load(open(CATF, encoding='utf-8'))
CATPN = collections.defaultdict(list)
for t in CAT['tables']:
    for r in t['rows']:
        CATPN[r['pn']].append((t, r))
PILOT = json.load(open(os.path.join(EV, 'pilot_sections.json'), encoding='utf-8'))
F62J = {json.loads(l)['step']: json.loads(l) for l in open(F62, encoding='utf-8')}


def load_table(t):
    ws = openpyxl.load_workbook(os.path.join(BASE, f'DesignTable_{t}.xlsx'), data_only=True).worksheets[0]
    hdr = [c.value for c in ws[2]]
    rows = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0]:
            rows[str(r[0]).strip()] = dict(zip(hdr, r))
    return hdr, rows


TABLES = {t: load_table(t) for t in sorted({r['table'] for r in ROWS} | {'4404', '4604'})}


def num(s):
    """printed catalog string -> float (decimals or fractions like '1 7/8', '1.0625')."""
    s = str(s).strip().replace("-", " ") if re.fullmatch(r"\d+-\d+/\d+", str(s).strip()) else str(s).strip()     # "1-1/4" -> "1 1/4"
    if re.fullmatch(r'\d*\.\d+|\d+', s):
        return float(s)
    if re.fullmatch(r'(\d+ )?\d+/\d+', s):
        return float(sum(Fraction(p) for p in s.split()))
    raise ValueError(s)


def fmt(v):
    return ('%.4f' % v).rstrip('0').rstrip('.')


def catrow(pn):
    """the page-verified forged catalog row of exactly this part number (never a non-FG or swept row)."""
    hits = [(t, r) for t, r in CATPN.get(pn, []) if t.get('verified')]
    assert len(hits) <= 1, (pn, [t['table_id'] for t, _ in hits])
    return hits[0] if hits else (None, None)


def catsrc(t, r, col, catalog_name):
    pdf = {'TFA': 'TFA Catalog.pdf', 'FLANGE': 'Flange Catalog.pdf'}[t['catalog']]
    return (f'Brennan {catalog_name} catalog p.{t["page"]} table {t["table_id"]} "{t["series"]} {t["title"]}" row {r["pn"]} col "{col}" = '
            f'"{r["values"][col]}" (07_RAW_EXTRACTS/brennan_catalogs/{pdf}; tools/_out/brennan_catalog_20261001/catalog_verified.json, '
            f'line_y {r["line_y"]}, page verified {t.get("verified")}) {DEC}')


def dashes(pn):
    return [int(x) for x in re.findall(r'-(\d\d)(?=-|$)', pn)]


# ------------------------------------------------------------------------------------------------------------------ end libraries
def lib_flg61_o():
    """code 61 head OD by flange dash as built in the 1700 master (Air-Way O of every 1700 row with that flange dash)."""
    hdr, rows = TABLES['1700']
    d = collections.defaultdict(set)
    for pn, rec in rows.items():
        m = re.match(r'1700-61-(\d\d)-(\d\d)$', pn)
        if m and rec.get('O@DRIVE') is not None:
            d[int(m.group(2))].add((rec['O@DRIVE'], pn, rec['$PRP@ScopeAction']))
    return d


FLG61_O = lib_flg61_o()


def flg61_o(fd):
    vals = sorted({v for v, _, _ in FLG61_O[fd]})
    rows = ', '.join(f'{pn} {v:g} ({sa})' for v, pn, sa in sorted(FLG61_O[fd], key=lambda x: x[1]))
    if len(vals) == 1:
        return vals[0], f'end library: code 61 head OD of flange -{fd:02d} = {vals[0]:g} as built in the 1700 master (Air-Way O on every 1700 row with this flange: {rows}) {ENDLIB} - the generator\'s FLG61 head (FLG61_MEAS: thickness / groove) by flange size', None
    # -24: 2.38 / 2.37 = 2.375 rounded both ways (Q-FLG61-O24, rounding-only, recommended 2.38 - FLG61LEG DESIGN 4.3)
    assert vals == [2.37, 2.38], vals
    return 2.38, (f'end library: code 61 head OD of flange -{fd:02d}: the 1700 master prints {rows} = 2.375 rounded both ways (rounding-only; '
                  f'Q-FLG61-O24 recommends 2.38, wave2_shapes_20261008/FLG61LEG/DESIGN.md 4.3; the Brennan 1700-24-24 / 1703-24-24-FG / 1704-24-24-FG '
                  f'STEP heads are 2.375) {ENDLIB}'), 'Q-FLG61-O24 (rounding-only 2.37 / 2.38; recommended 2.38)'


def lib_bead():
    d = collections.defaultdict(set)
    for t, col in (('4404', 'AA@DRIVE'), ('4604', 'AA@DRIVE')):
        for pn, rec in TABLES[t][1].items():
            if rec.get('$PRP@ScopeAction') == 'BUILD' and rec.get(col) is not None and rec.get('$PRP@DataStatus') in ('ASSUMED', 'VERIFIED'):
                d[dashes(pn)[0]].add((rec[col], f'{t}:{pn}'))
    return d


BEAD = lib_bead()


def bead(hd):
    vals = sorted({v for v, _ in BEAD.get(hd, ())})
    if len(vals) != 1:
        return None, f'no single bead dia for hose -{hd:02d} in the built 4404 / 4604 rows: {sorted(BEAD.get(hd, ()))}'
    rows = sorted(p for _, p in BEAD[hd])
    return vals[0], (f'end library: hose bead dia AA of hose -{hd:02d} = {vals[0]:g} on every built (BUILD) row of the 4404 / 4604 masters ({", ".join(rows)}; '
                     f'Air-Way; gen_straight2 end_HOSEB, the barb end every built hose fitting uses; the Brennan 4404-08-06 / -16-16 STEPs show the '
                     f'same bead form) {ENDLIB}')


def jic_tm(d):
    return G.JIC[d]['TM'], f'end library: JIC -{d:02d} thread major {G.JIC[d]["TM"]:g} ({G.JIC[d]["thr"]} UN/UNF; gen_straight2 JIC = templates/DesignTable_2403_SW2.xlsx, MS33656J chain) {ENDLIB}'


def orb_tm(d):
    return G.ORB[d]['TM'], f'end library: ORB -{d:02d} thread major {G.ORB[d]["TM"]:g} ({G.ORB_THR[d]} UN/UNF; gen_straight2 ORB dictionary from the 6401 / 6405 / 6410 tables) {ENDLIB}'


def pipe_od(d):
    return MS.NPTF[d][2], f'end library: pipe OD of the -{d:02d} NPTF thread = {MS.NPTF[d][2]:g} (ASME B1.20.1 Table 2, tools generators/ms51500.py NPTF[{d}], as on every built pipe end) {ENDLIB}'


def tube_frac(d):
    return {2: '1/8', 3: '3/16', 4: '1/4', 5: '5/16', 6: '3/8', 8: '1/2', 10: '5/8', 12: '3/4', 14: '7/8', 16: '1', 20: '1 1/4', 24: '1 1/2', 32: '2'}[d]


def thread_text(t, dash, which):
    """the $PRP@Threads text an existing row of the same table prints for this threaded-end dash (format of the table)."""
    hdr, rows = TABLES[t]
    for pn, rec in rows.items():
        txt = rec.get('$PRP@Threads')
        if not txt:
            continue
        ds = dashes(pn.replace(' ', '-'))
        if len(ds) >= 2 and ds[which] == dash:
            return txt, pn
    return None, None


# ------------------------------------------------------------------------------------------------------------------ per-family rules
FLANGE_COLS = {   # table -> (JIC leg col, flange leg col, hex col); catalog L1 = JIC leg, L2 = flange face to the crossing, A/F = forged body
    '1703': ('M2@DRIVE', 'N3@DRIVE', 'CHEX@HEX'), '1704': ('M@DRIVE', 'N2@DRIVE', 'Y@HEX'),
    '1803': ('M2@DRIVE', 'N3@DRIVE', 'CHEX@HEX'), '1804': ('M@DRIVE', 'N2@DRIVE', 'Y@HEX'),
}
HOSE_COLS = {     # table -> (catalog col -> table col); D-col = threaded-end drill; DH@BORE = hose drill (new column)
    '4501': {'L LGTH': 'L4@DRIVE', 'N LGTH': 'N@DRIVE', 'Y FLATS': 'Y@HEX', 'A2 DRILL': 'D1@BORE', 'B1 DRILL': 'DH@BORE'},
    '4503': {'L LGTH': 'L5@DRIVE', 'N LGTH': 'N2@DRIVE', 'Y FLATS': 'CHEX@HEX', 'A2 DRILL': 'D1@BORE', 'B1 DRILL': 'DH@BORE'},
    # 4601 / 4603: Q-HOSEB-BORE answered (a) 2026-10-18 = the row's own single bore through both legs -> no DH@BORE column (B1 noted in $COMMENT)
    '4601': {'L LGTH': 'L4@DRIVE', 'M LGTH': 'M7@DRIVE', 'C HEX': 'C4HEX@HEX', 'Y FLATS': 'Y@HEX', 'B2 DRILL': 'D@BORE'},
    '4603': {'L LGTH': 'L5@DRIVE', 'M LGTH': 'M8@DRIVE', 'C HEX': 'C4HEX@HEX', 'Y FLATS': 'CHEX@HEX', 'B2 DRILL': 'D@BORE'},
}
MEANING = {'L LGTH': 'axis crossing to the hose-barb tip', 'N LGTH': 'axis crossing to the male NPTF tip', 'M LGTH': 'axis crossing to the adjustable ORB tip',
           'Y FLATS': 'forged body across flats', 'C HEX': 'ORB locknut hex', 'A2 DRILL': 'male NPTF end drill', 'B2 DRILL': 'adjustable ORB end drill',
           'B1 DRILL': 'hose-barb drill (new column DH@BORE: the one-bore tables cannot carry it; see REPORT)',
           'L1 (in.)': 'JIC tip to the crossing (JIC leg)', 'L2 (in.)': 'flange face to the crossing (flange leg)', 'A/F (in.)': 'forged body across flats',
           'L (in.)': 'JIC tip to the flange face', }
NEWCOL = ' [column not in today\'s base table: add it]'


def main():
    cells, rowsout = [], []

    def add(t, pn, col, val, src, kind):
        hdr, rows = TABLES[t]
        if col in hdr and col != '$COMMENT':          # $COMMENT is appended by the apply script, never replaced
            cur = rows[pn].get(col)
            assert cur in (None, ''), (t, pn, col, cur, 'base cell not empty')
        elif col not in hdr:
            src += NEWCOL
        cells.append(dict(table=t, config=pn, column=col, value=val if isinstance(val, str) else fmt(val), source=src, kind=kind))

    for R in ROWS:
        t, pn = R['table'], R['config']
        hdr, rows = TABLES[t]
        assert pn in rows, (t, pn)
        assert all(rows[pn].get(h) in (None, '') for h in hdr if h and ('@' in h) and not h.startswith('$')), (pn, 'base row has values')
        gates, missing, conflicts, notes = [], [], [], []
        comment = []
        ds = dashes(pn)
        tC, rC = catrow(pn)
        stepf = os.path.join(STEPDIR, f'BRENNAN_{pn}.stp')
        has_step = os.path.exists(stepf)
        # ---------------------------------------------------------------- 1700 code 61 straights (round collar bodies)
        if t == '1700':
            jd, fd = ds
            assert tC is not None
            add(t, pn, 'L28@DRIVE', num(rC['values']['L (in.)']), catsrc(tC, rC, 'L (in.)', 'Flange') + ' - JIC tip to the flange face (mapping: drawing Flange p3; the 9 Brennan 1700 STEPs reproduce catalog L 9/9, FLANGE62 validation_1700.txt; this row\'s own STEP L = ' + fmt(F62J[f'BRENNAN_{pn}.stp']['L']) + ')', 'brennan-dim')
            af = num(rC['values']['A/F (in.)']); stp_af = F62J[f'BRENNAN_{pn}.stp']['flats_AF']
            add(t, pn, 'CHEX@HEX', af, catsrc(tC, rC, 'A/F (in.)', 'Flange') + f' - wrench flats on Brennan\'s ROUND collar (not a full hex: Q-NEW-1700BODY); own STEP flats {stp_af:g} (A16 2026-10-18 "do the same as Brennan": the body becomes Brennan\'s round collar + two flats, CHEX = the flats A/F)', 'brennan-dim')
            if abs(stp_af - af) > 0.006:
                notes.append(f'CHEX@HEX: Brennan catalog A/F {af:g} vs own STEP flats {stp_af:g} (Brennan internal, FLANGE62 validation_1700.txt): catalog value as printed - user rule A13 2026-10-18 "if its printed in the spec sheets you don\'t need to check in the steps again"')
            if pn == '1700-32-24':
                notes.append('catalog A/F prints 2.622 (1700-32-32 prints 2.625; 2.622 is not a 64th) - the own STEP reads 2.622 too, so it is listed as printed')
            o, osrc, oq = flg61_o(fd)
            j = F62J[f'BRENNAN_{pn}.stp']
            add(t, pn, 'O@DRIVE', o, osrc + f'; this row\'s own STEP head OD {j["head"]["O"]:g} (flange62.py)', 'endlib')
            if oq: notes.append(oq)
            # bore: Air-Way's 1700 D is the SMALLER of the two end bores on every row (1700-61-16-12 0.61, 1700-61-12-16 0.61, 1700-61-20-24 1.08);
            # the master has one through bore (SPEC bore 'mid', D@BORE both ways) -> D@BORE = the smaller Brennan bore; the other noted
            jb = j['bore']['D']
            prof = {'1700-20-16': (1.078, None), '1700-24-20': (1.078, 1.312), '1700-32-24': (1.500, 1.781)}[pn]
            assert abs((prof[1] or prof[0]) - jb) < 0.0015, (pn, prof, jb)     # prof.py JIC-side bore = the validated flange62.py reading
            src = (f'own Brennan STEP 09_STEP_FILES/BRENNAN_{pn}.stp, tools/measure/prof.py section (evidence/1700_bores_prof.txt) + FLANGE62 flange62.py '
                   f'(validated on the 9 Brennan 1700 STEPs, FLANGE62/evidence/measure/flange62_straight.jsonl: JIC-side bore {jb:g}) {DEC}: ')
            if prof[1]:
                src += (f'Brennan bore is STEPPED: JIC side {prof[1]:g}, flange side {prof[0]:g}; the 1700 master has ONE through bore and Air-Way\'s D is the smaller '
                        f'end bore on every 1700 row -> D@BORE = {prof[0]:g} (modelling assumption, logged; the JIC-side {prof[1]:g} is not modelled)')
                comment.append(f'bore: Brennan stepped {prof[1]:g} (JIC side) / {prof[0]:g} (flange side), modelled as one {prof[0]:g} through bore')
            else:
                src += f'one through bore {prof[0]:g}'
            add(t, pn, 'D@BORE', prof[0], src, 'brennan-dim')
            v, s = jic_tm(jd); add(t, pn, 'THDMAJ_B@DRIVE', v, s, 'endlib')
            add(t, pn, '$PRP@TubeOD', tube_frac(jd), f'tube size of JIC dash -{jd:02d}; = Brennan Flange p3 "TUBE O.D." {rC["values"]["TUBE O.D."]}', 'other')
            add(t, pn, '$PRP@Threads', f'{G.JIC[jd]["thr"]} (B UN/UNF-2A)', f'JIC thread of dash -{jd:02d} (gen_straight2 JIC); = Brennan Flange p3 "1 THREAD" {rC["values"]["1 THREAD"]}' + ('; = OMi Items Side 1 Thread "1 5/8-12"' if pn == '1700-20-16' else ''), 'other')
            # A16 answered 2026-10-18 "do the same as Brennan": round collar + two wrench flats (+ neck) -> collar / flats-zone / neck from the own STEP
            fz = j['collar_face_to_neck_start']
            s62 = (f'own Brennan STEP 09_STEP_FILES/BRENNAN_{pn}.stp, FLANGE62 tools/measure/flange62.py (validated on the 9 Brennan 1700 STEPs: heads = FLG61_MEAS 9/9, '
                   f'L = catalog 9/9; newmasters_20261002/work/FLANGE62/evidence/measure/flange62_straight.jsonl) {DEC}')
            add(t, pn, 'COLLARD@DRIVE', j['collar_dia'], s62 + f': round collar dia {j["collar_dia"]:g} (A16 round-collar body; column of the FLANGE62 1800 SPEC)', 'brennan-dim')
            add(t, pn, 'FLATL@DRIVE', fz, s62 + f': collar face to the neck start {fz:g} (collar {j["collar_round_len"]:g} + 30 deg chamfer {j["collar_chamfer_len"]:g}) = the 1800 FLATL zone; on the 1700 STEPs the flats run out in the chamfer ({j["flats_zone_len"]:g} long) - the 1800 construction ends them at the neck start (logged)', 'brennan-dim')
            add(t, pn, 'NECKD@DRIVE', j['neck_dia'], s62 + f': round neck dia {j["neck_dia"]:g} ({j["neck_len"]:g} long)', 'brennan-dim')
            comment.append('A16 round collar + two flats as Brennan (COLLARD / FLATL / NECKD from the own STEP)')
            gates.append('TOOLCHAIN: A16 (2026-10-18) round-collar body for 1700 - the batch-16 gen_straight2 SPEC 1700 models a full hex; the collar + flats construction exists only in the unmerged FLANGE62 staging (SPEC 1800 option flats=); SPEC 1700 needs it (and must read COLLARD / FLATL / NECKD)')
            notes.append('Air-Way 1700 lengths differ from Brennan on every shared size (L28 vs catalog L: 0 of 13 equal, Brennan bodies 0.09-0.70 longer; evidence/pairing.txt) - these 3 rows would be the only Brennan-sized rows of an Air-Way-sized master')
        # ---------------------------------------------------------------- forged flange elbows 1703 / 1704 (code 61) and 1803 / 1804 (code 62)
        elif t in FLANGE_COLS:
            jd, fd = ds
            cj, cf, ch = FLANGE_COLS[t]
            code = 61 if t in ('1703', '1704') else 62
            if tC is None:
                if pn == '1803-08-08-FG':
                    missing += [cj, cf, ch, 'D@BORE', 'DFLG@BORE']
                    notes.append('no source: Brennan prints 1803-08-08 only as the swept (tube) elbow (Flange p5) and has no 3D file of a forged 1803-08-08-FG (3D_FILES_NEEDED.md "Not to ask for"); OMi Items row 314 empty -> needs OMi')
                else:
                    missing += [cj, cf, ch, 'D@BORE', 'DFLG@BORE']
                    notes.append(f'not in Brennan\'s Flange catalog (forged table {"p4 1703-FG" if t == "1703" else "p5 1803-FG"} has no {pn}) -> every body size only from its own STEP: needs 09_STEP_FILES/BRENNAN_{pn}.stp')
            else:
                v1, v2, vaf = (num(rC['values'][k]) for k in ('L1 (in.)', 'L2 (in.)', 'A/F (in.)'))
                ev = (' (mapping: drawing Flange p4-p6 forged elbows; own STEPs reproduce L1 / L2 / A/F exactly: 1703-24-24-FG 1.780 / 2.380 / 1.875, '
                      '1704-24-24-FG 2.330 / 3.150 / 1.875, 1804-08-08-FG 1.520 / 2.050 / 0.875, evidence/pilot_sections.txt; Air-Way 1703-61-24-24 M2 1.78 / '
                      '1704-61-24-24 M 2.33 / 1704-61-32-32 M 3.06 = catalog L1)')
                add(t, pn, cj, v1, catsrc(tC, rC, 'L1 (in.)', 'Flange') + ' - ' + MEANING['L1 (in.)'] + ev, 'brennan-dim')
                add(t, pn, cf, v2, catsrc(tC, rC, 'L2 (in.)', 'Flange') + ' - ' + MEANING['L2 (in.)'] + ev, 'brennan-dim')
                add(t, pn, ch, vaf, catsrc(tC, rC, 'A/F (in.)', 'Flange') + ' - ' + MEANING['A/F (in.)'] + ev, 'brennan-dim')
                if pn in PILOT:
                    P = PILOT[pn]
                    pj, pf = P['JIC'], P['flange']
                    assert abs(pj['tip'] - v1) < 0.002 and abs(pf['tip'] - v2) < 0.002, (pn, pj['tip'], pf['tip'], v1, v2)
                    bj = min(r for r, a1, a2 in pj['z0_parallel'] if a2 > pj['tip'] - 0.01)
                    bf = min(r for r, a1, a2 in pf['z0_parallel'] if a2 > pf['tip'] - 0.01)
                    srcb = (f'own Brennan STEP 09_STEP_FILES/BRENNAN_{pn}.stp, scripts/measure_pilots.py (tools/measure section.py, OCP; the same sections '
                            f'reproduce catalog L1 {v1:g} / L2 {v2:g} / A/F {vaf:g} first) evidence/pilot_sections.txt {DEC}')
                    add(t, pn, 'D@BORE', round(2 * bj, 4), srcb + f': JIC-leg bore {2 * bj:.4f} from the JIC tip to the crossing', 'brennan-dim')
                    add(t, pn, 'DFLG@BORE', round(2 * bf, 4), srcb + f': flange-leg bore {2 * bf:.4f} from the flange face to the crossing (new column: Brennan\'s flange leg has its own bore; the FLG61LEG / FLG62LEG leg tuple reads it as leg-1 bore column)', 'brennan-dim')
                    if abs(bj - bf) > 0.001:
                        comment.append(f'bores: JIC leg {2 * bj:.3f}, flange leg {2 * bf:.3f} (own STEP)')
                else:
                    missing += ['D@BORE', 'DFLG@BORE']
                    notes.append(f'bores: the catalog prints no drill; only the own STEP gives them -> needs 09_STEP_FILES/BRENNAN_{pn}.stp')
            # head OD
            if code == 61:
                o, osrc, oq = flg61_o(fd)
                add(t, pn, 'O@DRIVE', o, osrc, 'endlib')
                if oq: notes.append(oq)
            else:
                if fd == 8:
                    if pn in PILOT:
                        add(t, pn, 'O@DRIVE', 1.25, f'own Brennan STEP 09_STEP_FILES/BRENNAN_{pn}.stp: code 62 -08 head dia 1.250 (section r 0.6250, evidence/pilot_sections.txt) = FLANGE62 FLG62_MEAS[8] O 1.250 (1800-08-08 and 1804-08-08-FG exact) {DEC}', 'brennan-dim')
                    else:
                        add(t, pn, 'O@DRIVE', 1.25, 'end library (staged, not merged): code 62 -08 head OD 1.250 = FLANGE62 FLG62_MEAS[8] (newmasters_20261002/work/FLANGE62, measured on the Brennan 1800-08-08 and 1804-08-08-FG STEPs) ' + ENDLIB, 'endlib')
                else:
                    missing.append('O@DRIVE')
                    notes.append(f'needs code 62 head STEP (-{fd:02d}): only the -08 code 62 head is measured (FLG62_MEAS[8]); e.g. 1824-{fd:02d} or 1800-{fd:02d}-{fd:02d} (3D_FILES_NEEDED.md part 4)')
            v, s = jic_tm(jd); add(t, pn, 'THDMAJ_B@DRIVE', v, s, 'endlib')
            add(t, pn, '$PRP@TubeOD', tube_frac(jd), f'tube size of JIC dash -{jd:02d}' + (f'; = Brennan Flange "TUBE O.D." {rC["values"]["TUBE O.D."]}' if rC else ''), 'other')
            add(t, pn, '$PRP@Threads', f'{G.JIC[jd]["thr"]} (B UN/UNF-2A)', f'JIC thread of dash -{jd:02d} (gen_straight2 JIC), the format of the 1703 / 1704 Air-Way rows' + (f'; = Brennan Flange "1 THREAD" {rC["values"]["1 THREAD"]}' if rC else ''), 'other')
            if code == 61:
                gates.append('TOOLCHAIN: Q-NEW-FLG61LEG (approved 2026-10-08) is not implemented in the batch-16 gen_tee (tools/_out/puckfix_20261009/tools_new: no SPEC 1703 / 1704) nor in wave2_impl_20261009 - design only (wave2_shapes_20261008/FLG61LEG/DESIGN.md)')
            else:
                gates.append('TOOLCHAIN: Q-NEW-FLG62LEG approved 2026-10-18 (A17) but no generator has a code 62 flange leg, and the FLANGE62 code 62 head end (FLG62_MEAS) is not merged into the batch-16 tools')
            for g in R['other_gates'].split(';'):
                g = g.strip()
                if g.startswith('Q-NEW-F3-CODE62LEN'):
                    notes.append('Q-NEW-F3-CODE62LEN answered 2026-10-18 (A13): catalog values as printed (ASSUMED), no STEP confirmation of printed values')
                if g.startswith('Q-NEW-F3-1803PILOT'):
                    notes.append('Q-NEW-F3-1803PILOT: moot with Brennan = yes (OPEN_QUESTIONS "Closed"), but this size is not in the catalog (see above)')
                if g.startswith('Q-NEW-F2A-NOTINCAT'):
                    notes.append('Q-NEW-F2A-NOTINCAT: moot with Brennan = yes (OPEN_QUESTIONS "Closed"), but the catalog has no forged -08 1703 -> own STEP only')
            notes.append('Q-TRANSFER-FLG61-O: closed by Q-TRANSFER-ENDLIB = yes (head OD by flange size from the 1700 master)' if code == 61 else '')
        # ---------------------------------------------------------------- beaded hose-barb elbows
        elif t in HOSE_COLS:
            hd, ed = ds[0], ds[1]
            assert tC is not None, pn
            mp = HOSE_COLS[t]
            for ccol, tcol in mp.items():
                ev = {'L LGTH': ' (mapping: drawing TFA p28-p30 L = hose leg; Brennan 4501-08-08-FG STEP hose hex arm 0.555 + stub 1.500 = 2.055 = catalog L, HOSEB DESIGN s3; Air-Way L4/L5 = crossing to the hose tip, equal to Brennan on 4601-06-06 / -32-32 NWO-FG only - Air-Way\'s hose legs differ on most sizes, evidence/pairing.txt)',
                      'N LGTH': ' (pairing: Air-Way plain 4501 N = catalog N on 20 of 28 sizes, 4501-32-32-FG equal; 4503 N2 14 of 19)',
                      'M LGTH': ' (pairing: Air-Way NWO-FG M7 / M8 = catalog M on 7 of 7 / 5 of 6 to rounding)',
                      'C HEX': ' (pairing: Air-Way C4HEX = catalog C HEX on 26 of 27 (4601) / 21 of 21 (4603) sizes)',
                      'Y FLATS': ' (drawing: Y = forged body flats; Air-Way prints the non-forged body hex on most rows, equal on 4501-32-32-FG / 4601 3 of 7 NWO-FG)',
                      'A2 DRILL': ' (pairing: Air-Way D1 = catalog A2 DRILL on 27 of 28 sizes (1 rounding), 4501-32-32-FG equal; 4503 14 of 19)',
                      'B2 DRILL': ' (pairing: Air-Way D = catalog B2 DRILL on 27 of 27 (4601) / 21 of 21 (4603) sizes)',
                      'B1 DRILL': ' (B1 = the hose end drill: differs from the Air-Way single bore on 26 of 28 4501 sizes; Brennan 4501-08-08-FG STEP hose drill 0.406 = catalog B1, pipe drill 0.531 = A2, HOSEB DESIGN s3)'}[ccol]
                add(t, pn, tcol, num(rC['values'][ccol]), catsrc(tC, rC, ccol, 'TFA') + ' - ' + MEANING[ccol] + ev, 'brennan-dim')
            b, bs = bead(hd)
            if b is None:
                missing.append('AA@DRIVE'); notes.append(bs)
            else:
                add(t, pn, 'AA@DRIVE', b, bs, 'endlib')
            tube = rC['values']['B1 TUBE O.D.']
            assert num(tube) == num(tube_frac(hd)), (pn, tube)
            add(t, pn, '$PRP@TubeOD', tube_frac(hd), f'hose size of dash -{hd:02d} = Brennan TFA "B1 TUBE O.D." {tube} (stub OD of the barb)', 'other')
            if t in ('4501', '4503'):
                v, s = pipe_od(ed); add(t, pn, 'THDMAJ_A@DRIVE', v, s, 'endlib')
                txt, ref = thread_text(t, ed, 1)
                add(t, pn, '$PRP@Threads', txt, f'thread callout of the -{ed:02d} male NPTF end in the table\'s format (row {ref}); = Brennan TFA "A2 NPTF" {rC["values"]["A2 NPTF"]}', 'other')
            else:
                v, s = orb_tm(ed); add(t, pn, 'THDMAJ_B@DRIVE', v, s, 'endlib')
                txt, ref = thread_text(t, ed, 1)
                add(t, pn, '$PRP@Threads', txt, f'thread callout of the -{ed:02d} adjustable ORB end in the table\'s format (row {ref}); = Brennan TFA "B2 AORB" {rC["values"]["B2 AORB"]}', 'other')
            # hose-leg bore vs stub
            dcol = [c for c in mp.values() if c in ('D@BORE', 'D1@BORE')][0]
            dthr = num(rC['values']['A2 DRILL' if t in ('4501', '4503') else 'B2 DRILL'])
            dh = num(rC['values']['B1 DRILL'])
            stub = num(tube_frac(hd))
            if stub - dthr < 0.03 and t in ('4601', '4603'):
                notes.append(f'under Q-HOSEB-BORE (a) the {dthr:g} bore through the {stub:g} hose stub leaves {(stub - dthr) / 2:.3f} wall: the HOSEB generator BLOCKS this row (hose stub wall < 0.015) - Brennan drills the hose leg {dh:g}')
                gates.append('BLOCKS under Q-HOSEB-BORE (a): own bore wider than the hose stub allows (see notes)')
            elif stub - dthr < 0.03:
                notes.append(f'one-bore generators: the threaded-end drill {dthr:g} through the {stub:g} hose stub leaves {(stub - dthr) / 2:.3f} wall - this row needs the hose drill DH@BORE {dh:g} on the hose leg (hose_bore option reading DH@BORE)')
            gates.append('TOOLCHAIN: Q-NEW-HOSEB (approved 2026-10-08) is implemented only in the unmerged wave2_impl_20261009 tools, not in the batch-16 gen_tee (no SPEC ' + t + ')')
            if t in ('4501', '4503'):
                gates.append('TOOLCHAIN (hose bore): Q-HOSEB-BORE for 4501 answered 2026-10-18 (A15 "look at Brennan"): the catalog prints the hose drill B1 -> DH@BORE (4503 = the same male-pipe case); the HOSEB generator needs a hose_bore option that reads DH@BORE (today only \'own\' / \'=4404\')')
            else:
                notes.append(f'Q-HOSEB-BORE answered 2026-10-18 (A15): (a) the row\'s own single bore {dthr:g} through both legs; Brennan hose drill B1 {dh:g} only noted')
            notes.append('STEP gate dropped: the catalog prints every body size and the family shape sample is in 09_STEP_FILES (HOSEB DESIGN s3) - the own STEP is not needed for values')
            comment.append(f'Brennan hose drill {dh:g} (DH@BORE) vs threaded-end drill {dthr:g} ({dcol})' if t in ('4501', '4503') else
                           f'Brennan hose drill {dh:g} not modelled: Q-HOSEB-BORE (a) = one bore {dthr:g} ({dcol}) through both legs')
        else:
            raise SystemExit(t)
        ready = 'no' if (gates or missing or conflicts) else 'yes'
        if not missing and not conflicts and gates and all(g.startswith('TOOLCHAIN') for g in gates):
            notes.insert(0, 'DATA COMPLETE: every driving cell sourced; held only by the generator (TOOLCHAIN) - ready once the capability is merged')
        brennan = any(c['kind'] == 'brennan-dim' and c['config'] == pn for c in cells)
        com = '2026-10-18 Brennan = yes (staged, not shipped): ' + ('filled from Brennan ($PRP@DimSource BRENNAN; to undo clear the BRENNAN cells and regenerate)' if brennan else 'end-library cells only, no Brennan size') + ('; ' + '; '.join(comment) if comment else '')
        add(t, pn, '$COMMENT', com, 'note appended to the row comment', 'other')
        rowsout.append(dict(table=t, config=pn, ready=ready, missing_columns=' ; '.join(missing), gates=' ; '.join(gates),
                            conflicts=' ; '.join(conflicts), notes=' ; '.join(n for n in notes if n)))
    with open(os.path.join(W, 'cells.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, ['table', 'config', 'column', 'value', 'source', 'kind']); w.writeheader(); w.writerows(cells)
    with open(os.path.join(W, 'rows.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, ['table', 'config', 'ready', 'missing_columns', 'gates', 'conflicts', 'notes']); w.writeheader(); w.writerows(rowsout)
    k = collections.Counter(c['kind'] for c in cells)
    print(len(cells), 'cells', dict(k), '|', len(rowsout), 'rows, ready', sum(r['ready'] == 'yes' for r in rowsout))
    for r in rowsout:
        print(r['table'], r['config'], r['ready'], '| missing:', r['missing_columns'], '| conflicts:', r['conflicts'][:80])


if __name__ == '__main__':
    main()
