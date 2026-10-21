"""review_pack.py (2026-10-21): turn map_columns.py's output into a workbook for the user's review, and count which not-ready parts the
proposed Brennan cells would make complete.
"Required" columns of a review table = numeric columns filled on >= 90 % of its rows that are already shipped (04_DESIGN_TABLES SW2
configurations); a not-ready BUILD row is COMPLETE when every required column is filled after the proposals. (A row can still need a
generator SPEC, a decision or a shape file - see the PC's BLOCKED_NEEDS.md; "complete" only means no dimension is missing.)
Writes BRENNAN_CELLS_for_review.xlsx (Summary, Proposals, Complete rows, Differences, Column map) next to this script."""
import os, re, csv, glob, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

H = os.path.dirname(os.path.abspath(__file__)); PC = os.path.dirname(H)
ROOT = os.path.normpath(os.path.join(PC, '..', '..', '..'))
J = lambda *a: os.path.join(ROOT, *a)
norm = lambda s: re.sub(r'\s+', '-', str(s).strip()).upper()
rd = lambda n: list(csv.DictReader(open(os.path.join(H, n), encoding='utf-8')))
props, diffs, cmap = rd('proposals.csv'), rd('differences.csv'), rd('column_map.csv')
NOT_READY = {norm(l) for l in open(J('tools', '_out', 'sources_20261021', 'not_ready_1022.txt')) if l.strip()}

shipped = set()
for p in glob.glob(J('04_DESIGN_TABLES', '*_SW2.xlsx')):
    ws = openpyxl.load_workbook(p, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True)); h = rows[0]
    for r in rows[1:]:
        if r and r[0]:
            d = dict(zip(h, r)); shipped.add(norm(d.get('$PRP@OMiSKU') or d.get('$PRP@PartNo') or r[0]))

isnum = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)
# columns every built row takes from the thread size / end library (Q-TRANSFER-ENDLIB = yes, A5), not from a dimension source
STD = re.compile(r'^(THDMAJ|TAPDRILL|ORB_|JIC_|SEAT|RELIEF|NPT|NPSM|NOSE|FSM_|FSW|NWO_)', re.I)
fill = collections.defaultdict(dict)
for r in props:
    fill[(r['root'], r['part'])][r['column']] = r
complete, partial = [], []
paths = [p for p in glob.glob(J('04_DESIGN_TABLES', 'DesignTable_*.xlsx')) if not p.endswith(('_SW2.xlsx', '_FINAL.xlsx', '_v2.xlsx'))]
paths += glob.glob(J('tools', '_out', 'brennan_yes_20261018', 'stage_orfs', '04_DESIGN_TABLES', 'DesignTable_FS*.xlsx'))
for path in paths:
    root = os.path.basename(path)[12:-5]
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:6]) if r and '$PRP@PartNo' in r), None)
    if hi is None:
        continue
    head = rows[hi]; numcols = [c for c in head[1:] if c and not str(c).startswith('$')]
    data = [(norm(dict(zip(head, r)).get('$PRP@OMiSKU') or dict(zip(head, r)).get('$PRP@PartNo') or r[0]), dict(zip(head, r)))
            for r in rows[hi + 1:] if r and r[0]]
    ship = [d for p, d in data if p in shipped]
    if not ship:
        continue
    req = [c for c in numcols if sum(isnum(d.get(c)) for d in ship) >= 0.9 * len(ship)]
    for p, d in data:
        if p not in NOT_READY or str(d.get('$PRP@ScopeAction') or '').upper() != 'BUILD':
            continue
        f = fill.get((root, p), {})
        if not f:
            continue
        miss = [c for c in req if not isnum(d.get(c)) and c not in f]
        lib = [c for c in miss if STD.match(c)]; miss = [c for c in miss if not STD.match(c)]
        (complete if not miss else partial).append(dict(root=root, part=p, cells_from_brennan=len(f), still_missing=', '.join(miss),
                                                        from_end_library=', '.join(lib)))

wb = openpyxl.Workbook(); B = Font(bold=True); HF = PatternFill('solid', fgColor='E4E7EB')


def sheet(title, rows_, note):
    ws = wb.create_sheet(title); ws['A1'] = note; ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[1].height = 45; ws.merge_cells('A1:H1')
    if not rows_:
        return
    for j, k in enumerate(rows_[0], 1):
        c = ws.cell(row=2, column=j, value=k); c.font = B; c.fill = HF
    for i, r in enumerate(rows_, 3):
        for j, k in enumerate(rows_[0], 1):
            v = r[k]
            try:
                v = float(v) if k in ('value', 'table', 'brennan') else v
            except ValueError:
                pass
            ws.cell(row=i, column=j, value=v)
    for j in range(1, len(rows_[0]) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 18
    ws.freeze_panes = 'A3'; ws.auto_filter.ref = f'A2:{openpyxl.utils.get_column_letter(len(rows_[0]))}{len(rows_) + 2}'


s = wb.active; s.title = 'Summary'
lines = ['Brennan CAD values proposed for empty cells (2026-10-21) - NOT applied. For review.',
         'Rule (user 2026-10-21): Air-Way first, then OMi\'s file / answers, then Brennan. Only EMPTY cells of not-ready BUILD rows are proposed; '
         'every cell would be marked $PRP@DimSource BRENNAN.',
         'Each column is mapped to a Brennan code on rows whose values are known (shipped / filled rows): "evidence" = rows that agree / rows '
         'with both values (a 2-decimal Air-Way value agrees with its 3-decimal Brennan rounding).', '',
         f'Proposed cells: {len(props)} in {len(set((r["root"], r["part"]) for r in props))} not-ready parts',
         f'Parts with no required dimension missing afterwards: {len(complete)} (sheet "Complete rows"; thread / port / seat columns come from the '
         'end library by thread size, as on every built row - listed in from_end_library)',
         f'Parts still missing a required column: {len(partial)} (same sheet, still_missing)',
         f'Filled cells where Brennan differs from the table (NOT changed - Air-Way stays): {len(diffs)} (sheet "Differences")',
         f'Columns mapped: {sum(r["status"] == "MAPPED" for r in cmap)} of {len(cmap)} (sheet "Column map")']
for i, l in enumerate(lines, 1):
    s.cell(row=i, column=1, value=l)
s.column_dimensions['A'].width = 140
sheet('Proposals', props, 'One row per proposed cell. value = Brennan CAD record value (inches). code = Brennan record code. evidence = calibration rows.')
sheet('Complete rows', complete + partial, 'still_missing empty = every required dimension present after the Brennan cells (a generator SPEC, '
      'decision or shape file may still be needed - PC BLOCKED_NEEDS.md).')
sheet('Differences', diffs, 'Filled table cell vs Brennan record value, beyond rounding. Not changed (Air-Way first). For information / OMi.')
sheet('Column map', cmap, 'MAPPED = agreed on >= 3 rows and >= 90 % of the rows with both values. tie / weak / no code agrees = not used.')
wb.save(os.path.join(H, 'BRENNAN_CELLS_for_review.xlsx'))
print(f'{len(props)} cells, {len(complete)} parts complete, {len(partial)} still missing something, {len(diffs)} differences')
print('complete by root:', collections.Counter(r['root'] for r in complete).most_common())
