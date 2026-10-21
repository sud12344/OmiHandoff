"""stage_tables.py (group D, 2026-10-21): the bent-tube review tables for gen_bent.py, from the Brennan CAD records (data/records.csv,
scripts/records.py) decoded on the family STEPs (data/dump_*.txt), Air-Way s115 (1701-61, 4 sizes - Air-Way FIRST there) and the Brennan
catalogs (cross-check). Writes stage/DesignTable_1704T.xlsx (1704 + 1804 -Tube), DesignTable_1703T.xlsx (1703 + 1803 -Tube),
DesignTable_6701T.xlsx (6701 + 6701-L), stage/cells.csv (every written cell with its source), stage/rows.csv (status / holds) and
stage/checks.txt (record vs catalog / library cross-checks).

Mapping (decoded on BRENNAN_1704-08-08 / 1804-08-08 / 1703-08-08 / 1803-12-12 / 6701-08-08, data/dump_*.txt, README s2):
  90 deg   M1 = L1 (JIC tip -> leg-2 axis), L1 = L2 (JIC axis -> flange face)
  45 deg   M1 = L1 - L2, L1 = L2 x sqrt2   (record L1 / L2 = horizontal / vertical from the JIC tip to the flange face centre)
  CHEX = AF, HEXL = JT, O = AFL, D = DRILLFL, TUBEOD = TOD, HUBD = DFL, HUBL = FFL
  RB  90 deg = 0.8 x EXT (EXT = min(L1 - GT - JT, L2 - FFL) on all 35 records; R 1.0152 = 0.8 x 1.269 on the 1704 / 1804-08-08 STEPs)
      45 deg = TOD (R 0.618 / 0.800 = TOD on the 1703-08-08 / 1803-12-12 STEPs)
  6701     M1 = M, L1 = L - F - J + LN (crossing -> nut face; L = crossing -> nose tip), NOSETIP = LN - F - J, CHEX = C1, HEXL = THK,
           NUTHEX = C2, D = DIA6, TUBEOD = DIA, RB = R
"""
import csv, json, math, os, re
import openpyxl

H = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.abspath(os.path.join(H, '..', '..', '..'))
OUT = os.path.join(H, 'stage')
BASE = os.path.join(ROOT, 'tools', '_out', 'brennancad_20261021', 'stage', '04_DESIGN_TABLES')
CAT = os.path.join(ROOT, 'tools', '_out', 'brennan_catalog_20261001', 'catalog_verified.json')
S115 = os.path.join(ROOT, 'tools', '_out', 'airway_20261014', 'series_115_pass1.txt')
S115B = os.path.join(ROOT, 'tools', '_out', 'airway_20261014', 'series_115_pass2.txt')
S61 = os.path.join(ROOT, 'tools', '_out', 'airway_20261014', 'series_61_pass1.txt')
TODAY = '2026-10-21'
REC_SRC = 'Brennan CAD record {pn} (PARTcommunity reversemap, tools/_out/partcommunity_20261021/cache.json)'
JIC_GT = {4: 0.55, 6: 0.556, 8: 0.657, 12: 0.864, 16: 0.911, 20: 0.958, 24: 1.083, 32: 1.333}   # gen_straight2 end_JICM endlen (2403 table)
FLG61_GO = {8: 0.5, 12: 0.625, 16: 0.78, 20: 0.875, 24: 1.0575, 32: 1.245}                     # gen_straight2 FLG61_MEAS (groove OD / 2)
FLG61_TH = {8: 0.265, 12: 0.265, 16: 0.315, 20: 0.315, 24: 0.315, 32: 0.375}
FLG62 = {8: (1.250, 0.305, 1.000), 12: (1.625, 0.345, 1.250), 16: (1.875, 0.375, 1.560), 20: (2.125, 0.405, 1.750),
         24: (2.500, 0.495, 2.115), 32: (3.125, 0.495, 2.490)}                                  # gen_bent FLG62_HEAD: O, TH, groove OD
HDR_FL = ['M1@DRIVE', 'L1@DRIVE', 'CHEX@HEX', 'HEXL@HEX', 'O@DRIVE', 'D@BORE', 'TUBEOD@TUBE', 'RB@PATH', 'HUBD@HEAD', 'HUBL@HEAD',
          '$PRP@FlangeCode', '$PRP@TubeOD', '$PRP@PartNo', '$PRP@DataStatus', '$PRP@AxialSource', '$COMMENT', '$PRP@ScopeAction',
          '$PRP@OMiSKU', '$PRP@DimSource']
HDR_SW = ['M1@DRIVE', 'L1@DRIVE', 'CHEX@HEX', 'HEXL@HEX', 'NUTHEX@NUTHEXSK', 'NOSETIP@SWSK', 'D@BORE', 'TUBEOD@TUBE', 'RB@PATH',
          '$PRP@TubeOD', '$PRP@PartNo', '$PRP@DataStatus', '$PRP@AxialSource', '$COMMENT', '$PRP@ScopeAction', '$PRP@OMiSKU',
          '$PRP@DimSource']
# data holds: the record contradicts itself or the catalog - cells are staged, the row stays MISSING until the user decides
HOLD = {
    '1804-16-24-Tube': 'Brennan record TOD 0.800 with DRILLFL 0.740 = a 0.030 tube wall (every other -16 JIC row: tube 1.043 / bore 0.75, '
                       'incl. 1704-16-24); not built - needs BRENNAN_1804-16-24.stp or the user\'s call (Q-GD-1804-16-24)',
    '1704-20-12-Tube': 'Brennan record TOD 1.339 (tube) is wider than its hub DFL 1.260 (-12 code 61 head): the tube cannot enter the hub '
                       'as recorded; not built - needs BRENNAN_1704-20-12.stp or the user\'s call (Q-GD-1704-20-12)',
    '6701-20-20': 'M conflict between two Brennan sources: CAD record M 3.260 vs TFA catalog p38 M 2.750 (L 3.750 in both; the -16 has M 3.10) '
                  '- never picked silently; Q-GD-6701-20 (the record value 3.26 fits the size series)',
    '6701-16-16': 'no Brennan CAD record (PARTcommunity: none) and no -16 6701 STEP: tube OD / bore / bend radius / hex length / nut '
                  'position have no source (catalog p38 gives only M 3.10 / L 3.15 / C1 1.375 / C2 1.500) - needs BRENNAN_6701-16-16.stp',
    '6701-L-16-16': 'no Brennan CAD record (PARTcommunity: none) and no 6701-L STEP: tube OD / bore / bend radius / hex length / nut '
                    'position have no source (catalog p38 6701-L gives only L 4.870 / M 3.125 / C1 1.375 / C2 1.500) - needs '
                    'BRENNAN_6701-L-16-16.stp',
}


def f(x):
    return float(x)


def frac(s):
    s = str(s).strip().replace('-', ' ')
    t = 0.0
    for p in s.split():
        if '/' in p:
            a, b = p.split('/'); t += int(a) / int(b)
        else:
            t += float(p)
    return t


def catalog():
    d = json.load(open(CAT))
    out = {}
    for t in d['tables']:
        for r in t['rows']:
            out[r['pn']] = (t['table_id'], t['series'], r['values'])
    return out


def airway115():
    rows = {}
    for p in (S115, S115B):
        for l in open(p):
            if l.startswith('1701-61'):
                c = [x.strip() for x in l.split('|')]
                rows.setdefault(c[0], []).append(c)
    out = {}
    for k, v in rows.items():
        assert all(x == v[0] for x in v), f'Air-Way s115 passes differ on {k}'
        c = v[0]
        a, b = c[0].split('-')[2:4]
        out[f'1704-{int(a):02d}-{int(b):02d}-Tube'] = dict(aw=c[0], N4=f(c[3]), M3=f(c[4]), C=f(c[5]), D=f(c[6]), O=f(c[7]), tube=c[1])
    return out, len(rows)


def airway61():
    o = {}
    for l in open(S61):
        if l.startswith('1700-61'):
            c = [x.strip() for x in l.split('|')]
            o.setdefault(int(c[0].split('-')[3]), set()).add(f(c[6]))
    return o


def base_rows():
    rows = {}
    for t in ('1703', '1704', '1803', '1804', '6701'):
        ws = openpyxl.load_workbook(os.path.join(BASE, f'DesignTable_{t}.xlsx'), data_only=True).worksheets[0]
        R = list(ws.iter_rows(values_only=True)); h = R[1]
        for x in R[2:]:
            if x[0] and ('TUBE' in str(x[0]).upper() or t == '6701'):
                rows[x[0]] = dict(zip(h, x))
    return rows


def main():
    os.makedirs(OUT, exist_ok=True)
    recs = {r['part']: r for r in csv.DictReader(open(os.path.join(H, 'data', 'records.csv')))}
    cat = catalog()
    aw, naw = airway115()
    aw61 = airway61()
    base = base_rows()
    assert len(base) == 68 and set(base) == set(recs), (len(base), set(base) ^ set(recs))
    cells, rowinfo, chk = [], [], []
    tables = {'1704T': [], '1703T': [], '6701T': []}

    def cell(tbl, pn, col, val, src):
        cells.append(dict(table=tbl, config=pn, column=col, value=val, source=src))
        return val

    for pn in sorted(base):
        b = base[pn]; r = recs[pn]; fam = pn.split('-')[0]
        has = r['match'] in ('exact', 'alias')
        bpn = r['brennan_pn'] or pn.replace('-Tube', '')
        rs = REC_SRC.format(pn=bpn)
        catpn = bpn if bpn in cat else pn.replace('-Tube', '')
        cv = cat.get(catpn, (None, None, {}))[2]
        row = {'$PRP@PartNo': pn, '$PRP@ScopeAction': b['$PRP@ScopeAction'], '$PRP@OMiSKU': b['$PRP@OMiSKU'], '$PRP@DimSource': 'BRENNAN'}
        used, notes = [], []
        if fam in ('1703', '1704', '1803', '1804'):
            tbl = '1704T' if fam in ('1704', '1804') else '1703T'
            bend90 = tbl == '1704T'
            code = '61' if fam.startswith('17') else '62'
            a, fl = int(pn.split('-')[1]), int(pn.split('-')[2])
            L1, L2, AF, JT, GT = f(r['L1']), f(r['L2']), f(r['AF']), f(r['JT']), f(r['GT'])
            AFL, EFL, BFL, DFL, FFL, DR, TOD = (f(r[k]) for k in ('AFL', 'EFL', 'BFL', 'DFL', 'FFL', 'DRILLFL', 'TOD'))
            # ---- cross-checks: record vs catalog (L1 / L2 / A/F), JIC end vs library, head vs library / catalog / Air-Way s61
            if cv:
                for k_, rv in (('L1 (in.)', L1), ('L2 (in.)', L2), ('A/F (in.)', AF)):
                    c_ = frac(cv[k_]); ok = abs(c_ - rv) <= 0.0051
                    chk.append(f'{pn:18} record {k_[:3]:3} {rv:<7} catalog {cat[catpn][0]} {cv[k_]:<7} {"ok" if ok else "DIFFERS"}')
            else:
                chk.append(f'{pn:18} no Brennan catalog row {catpn}')
            chk.append(f'{pn:18} JIC -{a:02d} end length GT {GT} vs library {JIC_GT[a]} {"ok" if abs(GT - JIC_GT[a]) < 0.0015 else "DIFFERS"}')
            if code == '61':
                okh = abs(BFL / 2 - FLG61_GO[fl]) < 0.0026 and abs(EFL - FLG61_TH[fl]) < 0.0015      # groove OD: record 2.12 / 2.495 vs measured 2.115 / 2.49 = rounding
                a61 = sorted(aw61.get(fl, []))
                chk.append(f'{pn:18} code 61 -{fl:02d} head record AFL {AFL} EFL {EFL} BFL {BFL} vs FLG61_MEAS TH {FLG61_TH[fl]} groove OD '
                           f'{2 * FLG61_GO[fl]} {"ok" if okh else "DIFFERS"}; Air-Way s61 O {a61} '
                           f'{"ok" if any(abs(AFL - x) <= 0.0051 for x in a61) else "DIFFERS"}')
            else:
                O62, T62, G62 = FLG62[fl]
                okh = abs(AFL - O62) < 0.0015 and abs(EFL - T62) < 0.0015 and abs(BFL - G62) < 0.0015
                chk.append(f'{pn:18} code 62 -{fl:02d} head record AFL {AFL} EFL {EFL} BFL {BFL} vs FLG62_HEAD (3300-62 / measured) '
                           f'{O62} / {T62} / {G62} {"ok" if okh else "DIFFERS"}')
            A = aw.get(pn)
            if bend90:
                M1v, L1v = (A['M3'], A['N4']) if A else (L1, L2)
            else:
                M1v, L1v = round(L1 - L2, 4), round(L2 * math.sqrt(2), 4)
            if A:
                awsrc = f'Air-Way s115 {A["aw"]} (07_RAW_EXTRACTS/airway_passes/1701_p1 / _p2, tools/_out/airway_20261014/series_115_pass1/2.txt; 2 passes identical)'
                cell(tbl, pn, 'M1@DRIVE', M1v, f'{awsrc}: M3 {A["M3"]} = JIC tip -> leg-2 axis (= Brennan L1 reading; Brennan {L1}) [Air-Way first]')
                cell(tbl, pn, 'L1@DRIVE', L1v, f'{awsrc}: N4 {A["N4"]} = JIC axis -> flange face (Brennan L2 {L2}) [Air-Way first]')
                cell(tbl, pn, 'CHEX@HEX', A['C'], f'{awsrc}: C {A["C"]} = end-piece hex (Brennan A/F {AF}) [Air-Way first]')
                cell(tbl, pn, 'O@DRIVE', A['O'], f'{awsrc}: O {A["O"]} = flange head OD (Brennan AFL {AFL}) [Air-Way first]')
                cell(tbl, pn, 'D@BORE', A['D'], f'{awsrc}: D {A["D"]} = bore (Brennan DRILLFL {DR}) [Air-Way first]')
                notes.append(f'Air-Way s115 {A["aw"]} first for M1 / L1 / CHEX / O / D (user 2026-10-20 rule; settles the Air-Way-vs-Brennan '
                             f'differences: Brennan L1 {L1} / L2 {L2} / A/F {AF} / AFL {AFL} / DRILLFL {DR})')
                if any(abs(x - y) > 0.0101 for x, y in ((A['M3'], L1), (A['N4'], L2), (A['D'], DR))):
                    notes.append('Air-Way and Brennan differ beyond rounding (M3 / N4 / D) - built to Air-Way')
                used_aw = ['M1', 'L1', 'CHEX', 'O', 'D']
            else:
                if bend90:
                    cell(tbl, pn, 'M1@DRIVE', M1v, f'{rs}: L1 {L1} = JIC tip -> leg-2 axis (decoded on BRENNAN_1704-08-08.stp: tip x -0.657, leg-2 axis x 1.663)')
                    cell(tbl, pn, 'L1@DRIVE', L1v, f'{rs}: L2 {L2} = JIC axis -> flange face (BRENNAN_1704-08-08.stp: face y -1.810)')
                else:
                    cell(tbl, pn, 'M1@DRIVE', M1v, f'{rs}: L1 {L1} - L2 {L2} = JIC tip -> crossing of the leg axes (L1 / L2 = horizontal / vertical '
                                                   f'JIC tip -> flange face centre, decoded on BRENNAN_1703-08-08.stp: face centre (2.063, -0.910) from the tip -0.657)')
                    cell(tbl, pn, 'L1@DRIVE', L1v, f'{rs}: L2 {L2} x sqrt2 = crossing -> flange face along leg 2 (45 deg)')
                cell(tbl, pn, 'CHEX@HEX', AF, f'{rs}: AF {AF} = end-piece hex across flats (STEP planes +-{AF / 2:.4f})')
                cell(tbl, pn, 'O@DRIVE', AFL, f'{rs}: AFL {AFL} = flange head OD (STEP cylinder); = Brennan catalog 3300-{code} A' +
                     (f', Air-Way s61 O {sorted(aw61.get(fl, []))}' if code == '61' else ''))
                cell(tbl, pn, 'D@BORE', DR, f'{rs}: DRILLFL {DR} = through bore (STEP: JIC end, tube and head bore one cylinder)')
                used_aw = []
            cell(tbl, pn, 'HEXL@HEX', JT, f'{rs}: JT {JT} = end-piece hex length (BRENNAN_1704-08-08.stp: hex x 0 ... 0.394)')
            cell(tbl, pn, 'TUBEOD@TUBE', TOD, f'{rs}: TOD {TOD} = bent tube OD (STEP: tube cylinder and bend torus tube dia)')
            cell(tbl, pn, 'HUBD@HEAD', DFL, f'{rs}: DFL {DFL} = flange-head braze hub dia (STEP cylinder behind the head)')
            cell(tbl, pn, 'HUBL@HEAD', FFL, f'{rs}: FFL {FFL} = flange face -> hub back face (STEP planes)')
            if bend90:
                ext = min(M1v - GT - JT, L1v - FFL)
                RB = round(0.8 * ext, 4)
                if not A:
                    assert abs(ext - f(r['EXT'])) < 0.0011, (pn, ext, r['EXT'])
                cell(tbl, pn, 'RB@PATH', RB, f'0.8 x EXT {ext:.4f} (EXT = min(M1 - GT {GT} - JT {JT}, L1 - FFL {FFL}){" = record EXT " + r["EXT"] if not A else " on the Air-Way lengths"}): '
                                         'Brennan\'s 90 deg bend rule - R 1.0152 = 0.8 x 1.269 on BRENNAN_1704-08-08 / 1804-08-08.stp; EXT = min(...) on all 35 records')
            else:
                RB = TOD
                cell(tbl, pn, 'RB@PATH', RB, f'{rs}: TOD {TOD} - Brennan\'s 45 deg bend radius = tube OD (R 0.618 / 0.800 = TOD on BRENNAN_1703-08-08 / 1803-12-12.stp)')
            row.update({'M1@DRIVE': M1v, 'L1@DRIVE': L1v, 'CHEX@HEX': A['C'] if A else AF, 'HEXL@HEX': JT, 'O@DRIVE': A['O'] if A else AFL,
                        'D@BORE': A['D'] if A else DR, 'TUBEOD@TUBE': TOD, 'RB@PATH': RB, 'HUBD@HEAD': DFL, 'HUBL@HEAD': FFL,
                        '$PRP@FlangeCode': code, '$PRP@TubeOD': str(f(r['OD']))})
            bcols = [c for c in ('M1', 'L1', 'CHEX', 'HEXL', 'O', 'D', 'TUBEOD', 'RB', 'HUBD', 'HUBL') if c not in used_aw]
            fam_step = {'1704': '1704-08-08', '1804': '1804-08-08', '1703': '1703-08-08', '1803': '1803-12-12'}[fam]
            ax = (f'[BRENNAN-CAD {TODAY}: {" ".join(bcols)}] ' + (f'[AIRWAY s115 {A["aw"]}: {" ".join(used_aw)}] ' if A else '') +
                  f'Brennan CAD record {bpn} decoded on BRENNAN_{fam_step}.stp (tools/_out/groupD_20261021 README s2); '
                  f'{"90 deg: M1 = L1, L1 = L2" if bend90 else "45 deg: M1 = L1 - L2, L1 = L2 x sqrt2"}; RB = '
                  f'{"0.8 x EXT (Brennan 90 deg bend rule, 1 STEP size)" if bend90 else "TOD (Brennan 45 deg bend rule, 2 STEP sizes)"}; '
                  f'catalog {cat[catpn][0] if cv else "-"} L1 / L2 / A/F = record | {b["$PRP@AxialSource"]}')
            com = (f'Brennan record {bpn}: L1 {L1} L2 {L2} AF {AF} JT {JT} GT {GT} AFL {AFL} EFL {EFL} BFL {BFL} DFL {DFL} FFL {FFL} DRILLFL {DR} '
                   f'TOD {TOD}' + (f' EXT {r["EXT"]}' if bend90 else '') + ' (codes: README s2). ' + ' '.join(notes))
        else:   # 6701 / 6701-L
            tbl = '6701T'
            if not has:
                c1, c2 = frac(cv['C1 HEX']), frac(cv['C2 HEX'])
                cell(tbl, pn, 'M1@DRIVE', frac(cv['M LGTH']), f'Brennan TFA catalog {cat[catpn][0]} {catpn}: M {cv["M LGTH"]} = JIC tip -> swivel-leg axis')
                cell(tbl, pn, 'CHEX@HEX', c1, f'Brennan TFA catalog {cat[catpn][0]} {catpn}: C1 HEX {cv["C1 HEX"]}')
                cell(tbl, pn, 'NUTHEX@NUTHEXSK', c2, f'Brennan TFA catalog {cat[catpn][0]} {catpn}: C2 HEX {cv["C2 HEX"]}')
                row.update({'M1@DRIVE': frac(cv['M LGTH']), 'CHEX@HEX': c1, 'NUTHEX@NUTHEXSK': c2, '$PRP@TubeOD': str(frac(cv['TUBE O.D.']))})
                ax = f'[BRENNAN-CATALOG {TODAY}: M1 CHEX NUTHEX] Brennan TFA catalog {cat[catpn][0]} {catpn} | {b["$PRP@AxialSource"]}'
                com = f'catalog {catpn}: M {cv["M LGTH"]} L {cv["L LGTH"]} C1 {cv["C1 HEX"]} C2 {cv["C2 HEX"]}. HELD: {HOLD[pn]}'
            else:
                M, L, F, J, LN, C1, C2, THK, DIA, D6, R, GL1 = (f(r[k]) for k in ('M', 'L', 'F', 'J', 'LN', 'C1', 'C2', 'THK', 'DIA', 'DIA6', 'R', 'L1'))
                a = int(pn.split('-')[1])
                for k_, rv in (('M LGTH', M), ('L LGTH', L), ('C1 HEX', C1), ('C2 HEX', C2)):
                    c_ = frac(cv[k_]); ok = abs(c_ - rv) <= 0.0051
                    chk.append(f'{pn:18} record {k_[:2]} {rv:<6} catalog {cat[catpn][0]} {cv[k_]:<6} {"ok" if ok else "DIFFERS"}')
                chk.append(f'{pn:18} JIC -{a:02d} end length L1 {GL1} vs library {JIC_GT[a]} {"ok" if abs(GL1 - JIC_GT[a]) < 0.0025 else "DIFFERS"}')
                nf = round(L - F - J + LN, 4); nt = round(LN - F - J, 4)
                if pn in HOLD:
                    cell(tbl, pn, 'M1@DRIVE', None, f'HELD: record M {M} vs catalog M {cv["M LGTH"]}')
                    row['M1@DRIVE'] = None
                else:
                    cell(tbl, pn, 'M1@DRIVE', M, f'{rs}: M {M} = JIC tip -> swivel-leg axis (BRENNAN_6701-08-08.stp: tip x -1.770, leg-2 axis x 0) = catalog')
                    row['M1@DRIVE'] = M
                cell(tbl, pn, 'L1@DRIVE', nf, f'{rs}: L {L} - F {F} - J {J} + LN {LN} = crossing -> swivel nut face (STEP 6701-08-08: L 1.520 = nose tip, '
                                              'F = nose above the body shoulder, J = shoulder -> nut back, LN = nut length; nut face y 1.928)')
                cell(tbl, pn, 'NOSETIP@SWSK', nt, f'{rs}: LN {LN} - F {F} - J {J} = nut face -> nose tip (STEP 6701-08-08: 1.928 - 1.520 = 0.408)')
                cell(tbl, pn, 'CHEX@HEX', C1, f'{rs}: C1 {C1} = end-piece hex (STEP planes +-0.406 at -08) = catalog C1')
                cell(tbl, pn, 'HEXL@HEX', THK, f'{rs}: THK {THK} = end-piece hex length (STEP 6701-08-08: hex x -1.113 ... -0.916)')
                cell(tbl, pn, 'NUTHEX@NUTHEXSK', C2, f'{rs}: C2 {C2} = swivel nut hex (= catalog C2 = SWIV_MEAS AF)')
                cell(tbl, pn, 'D@BORE', D6, f'{rs}: DIA6 {D6} = through bore (STEP: tube / bend torus inner dia 0.374)')
                cell(tbl, pn, 'TUBEOD@TUBE', DIA, f'{rs}: DIA {DIA} = bent tube OD (STEP: bend torus tube dia 0.549)')
                cell(tbl, pn, 'RB@PATH', R, f'{rs}: R {R} = bend radius (STEP: torus R 0.827 at -08)')
                row.update({'L1@DRIVE': nf, 'NOSETIP@SWSK': nt, 'CHEX@HEX': C1, 'HEXL@HEX': THK, 'NUTHEX@NUTHEXSK': C2, 'D@BORE': D6,
                            'TUBEOD@TUBE': DIA, 'RB@PATH': R, '$PRP@TubeOD': str(frac(r['OD']))})
                ax = (f'[BRENNAN-CAD {TODAY}: M1 L1 CHEX HEXL NUTHEX NOSETIP D TUBEOD RB] Brennan CAD record {bpn} decoded on '
                      f'BRENNAN_6701-08-08.stp (tools/_out/groupD_20261021 README s2): M1 = M, L1 = L - F - J + LN (nut face), NOSETIP = LN - F - J, '
                      f'RB = R; catalog {cat[catpn][0]} M / L / C1 / C2 = record | {b["$PRP@AxialSource"]}')
                com = (f'Brennan record {bpn}: M {M} L {L} F {F} J {J} LN {LN} C1 {C1} C2 {C2} THK {THK} DIA {DIA} DIA6 {D6} R {R}. '
                       f'Swivel shape = the straights\' SWIV_MEAS (Q-TRANSFER-ENDLIB); NOSETIP from the record.')
                if pn in HOLD:
                    com += f' HELD: {HOLD[pn]}'
        if pn in HOLD:
            status = 'MISSING'
            if fam not in ('6701',) or has:
                com = com + ('' if 'HELD' in com else f' HELD: {HOLD[pn]}')
        else:
            status = 'ASSUMED'
        row.update({'$PRP@DataStatus': status, '$PRP@AxialSource': ax, '$COMMENT': com})
        tables[tbl].append(row)
        rowinfo.append(dict(table=tbl, config=pn, status=status, airway='yes' if fam == '1704' and pn in aw else '',
                            record=r['match'], hold=HOLD.get(pn, '')))
    assert naw == 4 and all(k in base for k in aw), aw
    for tbl, rows in tables.items():
        hdr = HDR_SW if tbl == '6701T' else HDR_FL
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
        n_as = sum(1 for r in rows if r['$PRP@DataStatus'] == 'ASSUMED')
        ws.append([f'OMi {tbl} bent-tube review table (gen_bent.py) | group D stage {TODAY} | {len(rows)} OMi rows, {n_as} ASSUMED | Brennan CAD '
                   'records decoded on the family STEPs + Air-Way s115 first | NOT SHIPPED'])
        ws.append([None] + hdr)
        for r in rows:
            ws.append([r['$PRP@PartNo']] + [r.get(k) for k in hdr])
        wb.save(os.path.join(OUT, f'DesignTable_{tbl}.xlsx'))
        print(f'{tbl}: {len(rows)} rows, {n_as} ASSUMED, {len(rows) - n_as} MISSING')
    with open(os.path.join(OUT, 'cells.csv'), 'w', newline='') as fh:
        w = csv.DictWriter(fh, ['table', 'config', 'column', 'value', 'source']); w.writeheader(); w.writerows(cells)
    with open(os.path.join(OUT, 'rows.csv'), 'w', newline='') as fh:
        w = csv.DictWriter(fh, ['table', 'config', 'status', 'airway', 'record', 'hold']); w.writeheader(); w.writerows(rowinfo)
    open(os.path.join(OUT, 'checks.txt'), 'w').write('\n'.join(chk) + '\n')
    bad = [c for c in chk if 'DIFFERS' in c or 'no Brennan catalog' in c]
    print(f'{len(cells)} cells; checks {len(chk)} lines, {len(bad)} DIFFERS / no catalog row:')
    for c in bad:
        print('  ', c)


if __name__ == '__main__':
    main()
