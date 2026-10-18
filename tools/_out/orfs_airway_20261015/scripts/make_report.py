"""make_report.py (2026-10-15): the readable outputs of the ORFS / compression / test-gauge intake.
  00_DOCS/ORFS_AirWay_Brennan_2026-10-15.xlsx   Summary, Parts, Conflicts, Forged, Compression + test gauge (formatted)
  (the Markdown page 00_DOCS/ORFS_Compression_TestGauge_2026-10-15.md is written by hand from the same numbers)
python -I make_report.py"""
import collections, json, os, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import load
from compare_orfs import PAIRS, NO_AIRWAY

H = load.HERE
D = json.load(open(os.path.join(H, 'compare_orfs.json'), encoding='utf-8'))
S = json.load(open(os.path.join(H, 'review_tables_summary.json'), encoding='utf-8'))
F = json.load(open(os.path.join(H, 'forged_info.json'), encoding='utf-8'))
O = load.omi_file()
OUT = os.path.join(load.ROOT, '00_DOCS', 'ORFS_AirWay_Brennan_2026-10-15.xlsx')
BIG = Font(size=12)
HEAD = Font(size=12, bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='305496')
COLORS = {'VERIFIED': 'C6EFCE', 'ASSUMED': 'FFEB9C', 'MISSING': 'FFC7CE'}


def sheet(wb, title, header, rows, widths, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BIG
            c.alignment = Alignment(wrap_text=True, vertical='top')
        if status_col is not None:
            st = str(row[status_col].value or '').split()[0] if row[status_col].value else ''
            if st in COLORS:
                row[status_col].fill = PatternFill('solid', fgColor=COLORS[st])
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 34
    return ws


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    by = collections.defaultdict(list)
    for r in D:
        by[r['series']].append(r)
    # ---- Summary
    rows = []
    for x in S:
        s = x['series']
        rs = by.get(s, [])
        small = sum(1 for r in rs if r.get('conflict_size') == 'small')
        large = sum(1 for r in rs if r.get('conflict_size') == 'large')
        noaw = sum(1 for r in rs if r['status'] == 'MISSING' and not r.get('conflict_size'))
        desc = O[rs[0]['omi_pn']].get('Series Description') if rs else ''
        note = ((PAIRS.get(s) or {}).get('note') or NO_AIRWAY.get(s) or
                ('forged (-FG): VERIFIED only where Air-Way\'s non-forged row = Brennan\'s forged row on every compared cell (user 2026-10-15); '
                 'the rest MISSING' if any(r['reason'].startswith('forged') for r in rs) else ''))
        rows.append([s, (desc or '').strip(), x['airway'] or 'none', x['build'], x['VERIFIED'], x['ASSUMED'], small, large, noaw, x['drop'], note])
    fs = [(s, rs) for s, rs in sorted(by.items()) if s not in {x['series'] for x in S}]
    for s, rs in fs:
        aw = sorted({r['airway_nonforged'][0].split('-')[0] for r in rs if r.get('airway_nonforged')})
        rows.append([s, (O[rs[0]['omi_pn']].get('Series Description') or '').strip(), ('plain ' + ', '.join(aw)) if aw else 'none', len(rs), 0, 0, 0, 0,
                     len(rs), 0, f'forged (-FG): held by the -FG rule; {sum(bool(r.get("airway_nonforged")) for r in rs)} have a non-forged Air-Way row'])
    sheet(wb, 'Summary', ['OMi series', 'What it is (OMi)', 'Air-Way series', 'OMi parts (parts file)', 'VERIFIED (Air-Way = Brennan)',
                          'ASSUMED (Air-Way; incl. the next two columns)', 'of which Brennan differs <= 0.02 in', 'of which Brennan differs > 0.02 in', 'MISSING (no Air-Way row / forged)', 'Air-Way rows OMi does not list (DROP)', 'Notes'],
          rows, [13, 34, 16, 12, 13, 12, 11, 11, 12, 13, 70])
    # ---- Parts
    rows = [[r['omi_pn'], r['series'], r.get('description'), r['status'] + (f' (Brennan differs, {r["conflict_size"]})' if r.get('conflict_size') else ''), r['reason'],
             '; '.join(r.get('airway_rows') or r.get('airway_nonforged') or []), '; '.join(r.get('brennan_rows') or []), r.get('thread_note') or '']
            for r in D]
    sheet(wb, 'Parts', ['OMi part', 'Series', 'Description (OMi)', 'Status', 'Why', 'Air-Way row', 'Brennan catalog row', 'Thread note'],
          rows, [22, 11, 36, 22, 80, 26, 30, 40], status_col=3)
    # ---- Conflicts
    rows = []
    for r in sorted((r for r in D if r.get('conflict_size')), key=lambda r: -r['conflict_max_in']):
        cells = [c for c in r.get('cells', []) if c['result'] == 'different']
        rows.append([r['omi_pn'], r['conflict_size'], r['conflict_max_in'], '\n'.join(f'{c["meaning"]}: Air-Way {c["airway"]}  /  Brennan {c["brennan"]}' for c in cells),
                     f'{r.get("airway_series")} / Brennan p{r.get("brennan_page")}'])
    sheet(wb, 'Brennan differs', ['OMi part', 'Size of the difference', 'Largest gap (in)', 'Cells that differ (built to Air-Way)', 'Sources'], rows, [22, 9, 12, 70, 26])
    # ---- Forged (information only)
    rows = []
    for bs, x in F['series'].items():
        rows.append([bs, x['airway'], x['sizes_in_both'], x['rows_all_equal'], '; '.join(f'{k}: {v["equal"]}/{v["equal"] + v["different"]}' for k, v in x['columns'].items())])
    sheet(wb, 'Forged (info)', ['Brennan forged series', 'Air-Way plain series', 'Sizes in both', 'Sizes where every compared cell is equal',
                                'Agreement per column (equal / compared)'], rows, [20, 14, 10, 14, 100])
    # ---- Compression + test gauge
    rows = []
    for pn, r in sorted(O.items()):
        if pn.startswith('FS'):
            continue
        rows.append([pn, str(r.get('Series')), r.get('Part Description'), 'MISSING',
                     'no Air-Way source (Air-Way has no compression / bite-type series; its test-point series 2101 / 2103 are metric test points, a different part); '
                     'Brennan prints only V0318 / V0319 (flareless nut / sleeve) - and Brennan alone may not be used'])
    sheet(wb, 'Compression + test gauge', ['OMi part', 'Series', 'Description (OMi)', 'Status', 'Why'], rows, [24, 12, 44, 12, 90], status_col=3)
    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
