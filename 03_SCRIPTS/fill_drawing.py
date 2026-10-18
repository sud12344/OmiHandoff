#!/usr/bin/env python3
"""
Fill the OMi tabulated fitting drawing — one drawing per GeometryRoot.

Professional fitting drawings are TABULATED: one sheet carries the whole
series in a dimension table rather than one sheet per part. This fills that
table from the Parts sheet.

  python fill_drawing.py --root 2404 --dry-run
  python fill_drawing.py --root 2404
  python fill_drawing.py --all-roots
"""
import argparse, datetime as dt, sys
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

HERE = Path(__file__).resolve().parent
WB   = HERE / "OMi_Fitting_Library_MasterData_v1.xlsx"
TPL  = HERE / "OMi_DrawingTemplate_Pro_A4.pdf"
ROWS = 18

def fmt(v):
    if v is None: return ""
    if isinstance(v, float): return f"{v:g}"
    return str(v)

def load_rows():
    ws = load_workbook(WB, data_only=True)["Parts"]
    hdr = [ws.cell(2,c).value for c in range(1, ws.max_column+1)]
    out = []
    for r in range(3, ws.max_row+1):
        rec = {str(h): ws.cell(r,i).value for i,h in enumerate(hdr,1) if h}
        if rec.get("PartNumber_OMi"): out.append(rec)
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root"); ap.add_argument("--all-roots", action="store_true")
    ap.add_argument("--out", type=Path, default=HERE/"drawings")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--revision", default="A")
    ap.add_argument("--include-unverified", action="store_true",
                    help="include MISSING-data parts in the table (marked with *)")
    a = ap.parse_args()

    rows = load_rows()
    roots = sorted({r["GeometryRoot"] for r in rows}) if a.all_roots else [a.root]
    if not roots or roots == [None]:
        print("Specify --root or --all-roots", file=sys.stderr); return 2
    a.out.mkdir(parents=True, exist_ok=True)

    for root in roots:
        sel = [r for r in rows if str(r.get("GeometryRoot")) == str(root)]
        # Carbon-steel scope gate: never emit a part that is not in the OMi catalog.
        dropped = [r for r in sel if str(r.get("ScopeAction", "BUILD")).upper() == "DROP"]
        sel = [r for r in sel if str(r.get("ScopeAction", "BUILD")).upper() != "DROP"]
        if dropped:
            print(f"      {len(dropped)} row(s) excluded: ScopeAction=DROP (not in OMi catalog)")
        if not a.include_unverified:
            sel = [r for r in sel
                   if str(r.get("DataStatus","")).upper() in ("VERIFIED","ASSUMED")]
        if not sel:
            print(f"SKIP  {root}: no releasable rows"); continue

        sheets = [sel[i:i+ROWS] for i in range(0, len(sel), ROWS)]
        for sn, chunk in enumerate(sheets, 1):
            vals = {
                "PartNumber_OMi":  f"{root} SERIES",
                "Description":     fmt(sel[0].get("Description","")).split(",")[0] + " — TABULATED",
                "Material":        fmt(sel[0].get("Material")) or "SEE NOTE 4",
                "Finish":          fmt(sel[0].get("Finish")) or "SEE NOTE 5",
                "GeometryRoot":    str(root),
                "EndA_ThreadStd":  fmt(sel[0].get("EndA_ThreadStd")),
                "EndB_ThreadStd":  fmt(sel[0].get("EndB_ThreadStd")),
                "MaterialSpec":    fmt(sel[0].get("MaterialSpec")),
                "FinishSpec":      fmt(sel[0].get("FinishSpec")),
                "IssueDate":       dt.date.today().isoformat(),
                "Revision":        a.revision,
                "Scale":           "NTS",
                "Sheet":           f"{sn} OF {len(sheets)}",
                "DataStatus":      "MIXED" if len({r.get("DataStatus") for r in chunk})>1
                                   else fmt(chunk[0].get("DataStatus")),
                "StdCheck_Result": fmt(chunk[0].get("StdCheck_Result")),
                "Notes": ("PARTS MARKED * ARE NOT STANDARDS-VERIFIED AND ARE NOT RELEASED "
                          "FOR PRODUCTION." if any(str(r.get("DataStatus","")).upper()!="VERIFIED"
                                                   for r in chunk) else ""),
            }
            for i, r in enumerate(chunk, 1):
                star = "" if str(r.get("DataStatus","")).upper()=="VERIFIED" else " *"
                vals[f"TBL{i:02d}_PN"]   = fmt(r.get("PartNumber_OMi")) + star
                vals[f"TBL{i:02d}_OD"]   = fmt(r.get("TubeOD_Frac"))
                vals[f"TBL{i:02d}_THDB"] = fmt(r.get("EndB_Thread"))
                vals[f"TBL{i:02d}_THDA"] = fmt(r.get("EndA_Thread"))
                vals[f"TBL{i:02d}_L"]    = fmt(r.get("D_L"))
                vals[f"TBL{i:02d}_C"]    = fmt(r.get("D_C"))
                vals[f"TBL{i:02d}_HEX"]  = fmt(r.get("D_HEXB"))

            out = a.out / f"{root}_SERIES_DWG{'' if len(sheets)==1 else f'_SH{sn}'}.pdf"
            if a.dry_run:
                print(f"DRY   {root} sheet {sn}/{len(sheets)}: {len(chunk)} rows -> {out.name}")
                continue
            rd = PdfReader(str(TPL)); w = PdfWriter(); w.append(rd)
            w.set_need_appearances_writer(True)
            for p in w.pages: w.update_page_form_field_values(p, vals)
            with open(out, "wb") as fh: w.write(fh)
            print(f"OK    {root} sheet {sn}/{len(sheets)}: {len(chunk)} rows -> {out.name}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
